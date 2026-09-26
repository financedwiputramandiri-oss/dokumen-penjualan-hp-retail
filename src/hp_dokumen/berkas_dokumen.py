"""Pembuat berkas dokumen — dipakai bersama oleh perintah manual dan bot penyapu.

Dipisahkan ke sini supaya dokumen yang dibuat manusia lewat `jalankan.py buat`
dan dokumen yang dibuat sendiri oleh bot penyapu benar-benar keluar dari satu
kode yang sama. Kalau dua-duanya punya kode sendiri, cepat atau lambat hasilnya
akan berbeda tanpa ada yang sadar.
"""
from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook

from .dokumen.faktur_pajak import buat_faktur_pajak
from .dokumen.invoice import buat_invoice
from .dokumen.proforma import buat_proforma
from .dokumen.mtn import buat_invoice_mtn, buat_surat_jalan_mtn
from .dokumen import rumus as rms
from .dokumen.surat_jalan import buat_packing_list, buat_surat_jalan
from .pdf import ke_pdf, libreoffice_ada

NOMOR_KOSONG = "________"


def nama_aman(teks: str) -> str:
    """Ubah nama tab jadi nama berkas yang aman."""
    teks = re.sub(r"[^\w\s\-()&]", "", teks)
    return re.sub(r"\s+", "_", teks.strip())


def buat_berkas(
    cfg,
    order,
    keputusan,
    folder: Path,
    nomor: str = NOMOR_KOSONG,
    *,
    master: dict | None = None,
    pdf: bool = False,
    sertakan_packing_list: bool = True,
    cetak=None,
) -> list[Path]:
    """Tulis dokumen satu PO ke `folder`. Mengembalikan daftar berkas yang jadi.

    Urutan dokumen mengikuti penjelasan Yosua 12 September 2026: Invoice,
    Surat Jalan, dan Faktur Pajak adalah yang diminta. Packing List ikut dibuat
    selama `sertakan_packing_list` masih True, karena belum dipastikan apakah
    masih dipakai.

    TIAP DOKUMEN BERKAS SENDIRI (permintaan Yosua 26 September 2026):
    *"MULAI SEKARANG DAN SETERUSNYA BUATLAH FILE PROFORMA INVOICE, INVOICE
    DAN SURAT JALAN SECARA TERPISAH"*. Ini MEMBATALKAN penggabungan Invoice +
    Surat Jalan yang diminta 20 September 2026 — jangan digabung lagi.

    Lembar MASTER HARGA ikut HANYA di berkas Invoice, sebab hanya rumus
    faktur yang menunjuk ke sana (VLOOKUP harga dan tarif PPN). Surat Jalan
    cuma memakai `=SUM(kolom ukuran)` yang tidak menyentuh master, jadi
    menyertakannya di situ hanya akan membocorkan daftar harga tanpa guna.

    Lembar itu bukan pelengkap: karena harga retail berubah sepanjang tahun
    (Milo Set 49.400 -> 61.000 pada Agustus 2026), tiap faktur jadi membawa
    bukti harganya sendiri.

    `master` boleh None — dokumennya tetap terbit, hanya tanpa lembar MASTER
    HARGA dan tanpa rumus. Dokumen yang tidak terbit jauh lebih merugikan
    daripada dokumen tanpa rumus.
    """
    cust = cfg.customer.cari(order.nama_tab)
    pt = cfg.perusahaan.untuk(cust)
    folder.mkdir(parents=True, exist_ok=True)
    aman = nama_aman(order.nama_tab)

    # CV. Mutiara Timur Nusantara memakai TATA LETAK SENDIRI, bukan sekadar
    # kop yang berbeda - dibuktikan dari berkas asli FA 0010526 BABY FAME (MTN)
    # yang dikirim Yosua 19 September 2026. Lihat dokumen/mtn.py.
    pakai_mtn = (pt.kode or "").strip().upper() == "MTN"

    pakai_rumus = bool(master)

    tugas = [
        ("FAKTUR_PAJAK",
         lambda ws: buat_faktur_pajak(ws, order, keputusan, cust, pt, cfg.pengaturan, nomor)),
    ]
    if sertakan_packing_list:
        tugas.append(
            ("PACKING_LIST", lambda ws: buat_packing_list(ws, order, cust, pt, nomor))
        )

    # Proforma HANYA untuk customer yang fakturnya dipecah per ukuran
    # (Haritsa & Katamama). Permintaan Yosua 17 September 2026; customer lain
    # tidak pernah memintanya, jadi jangan diterbitkan untuk semua.
    if cust and cust.pecah_per_ukuran:
        tugas.append(
            ("PROFORMA",
             lambda ws: buat_proforma(ws, order, keputusan, cust, pt,
                                      cfg.pengaturan, nomor))
        )

    # Surat Jalan berkas SENDIRI, tanpa MASTER HARGA — rumusnya tidak
    # menunjuk ke sana.
    tugas.insert(0, ("SURAT_JALAN", lambda ws: (
        buat_surat_jalan_mtn(ws, order, cust, pt, nomor, pakai_rumus=pakai_rumus)
        if pakai_mtn else
        buat_surat_jalan(ws, order, cust, pt, nomor, pakai_rumus=pakai_rumus)
    )))

    dibuat: list[Path] = []

    # ---- berkas INVOICE (+ lembar MASTER HARGA kalau ada) ---------------
    wb = Workbook()
    ws_inv = wb.active
    ws_inv.title = rms.TAB_INVOICE
    if pakai_mtn:
        buat_invoice_mtn(ws_inv, order, keputusan, cust, pt, cfg.pengaturan,
                         nomor, pakai_rumus=pakai_rumus)
    else:
        buat_invoice(ws_inv, order, keputusan, cust, pt, cfg.pengaturan, nomor,
                     pakai_rumus=pakai_rumus)
    # buat_invoice* menyetel ws.title sendiri; dikembalikan supaya nama
    # lembarnya sama untuk DPM maupun MTN, sebab rumus menunjuk ke nama itu.
    ws_inv.title = rms.TAB_INVOICE

    if master:
        tarif = cfg.pengaturan.tarif_ppn if pt.kenakan_ppn else 0.0
        rms.tulis_master(wb.create_sheet(rms.TAB_MASTER), master, pt, tarif,
                         order.tanggal_po)
    berkas_utama = folder / f"INVOICE_{aman}.xlsx"
    wb.save(berkas_utama)
    dibuat.append(berkas_utama)

    # PDF dokumen ini dibuat dari SALINAN TANPA lembar MASTER HARGA.
    # PDF-lah yang dikirim ke customer, dan master harga memuat SELURUH harga
    # retail Happy Pumpkin — 195 artikel, sepuluh halaman. Ikut tercetak
    # berarti membocorkan daftar harga seluruh produk ke satu customer.
    # Lembarnya tetap ada di berkas Excel-nya, yang dipakai di dalam kantor.
    berkas_pdf_sumber = berkas_utama
    if master:
        # Lembarnya DISEMBUNYIKAN, bukan dihapus. Versi pertama menghapusnya
        # dan seluruh rumus faktur langsung rusak: Subtotal jadi nol dan
        # DPP/PPN jadi `#NAME?`, karena rumusnya menunjuk ke lembar yang
        # sudah tidak ada. Ketahuan dari PDF hasil render, bukan dari kode.
        # LibreOffice tidak mencetak lembar tersembunyi, tapi rumus yang
        # menunjuk ke sana tetap terhitung.
        wb[rms.TAB_MASTER].sheet_state = "hidden"
        berkas_pdf_sumber = folder / f"_tanpa_master_{aman}.xlsx"
        wb.save(berkas_pdf_sumber)

    for nama, fungsi in tugas:
        wb = Workbook()
        ws = wb.active
        ws.title = nama.replace("_", " ").title()[:31]
        fungsi(ws)
        p = folder / f"{nama}_{aman}.xlsx"
        wb.save(p)
        dibuat.append(p)

    if pdf:
        if not libreoffice_ada():
            if cetak:
                cetak("  (PDF dilewati: LibreOffice tidak terpasang di komputer ini.")
                cetak("   Di Ubuntu/Debian: sudo apt install libreoffice-calc)")
        else:
            for p in list(dibuat):
                sumber = berkas_pdf_sumber if p == berkas_utama else p
                hasil_pdf, pesan = ke_pdf(sumber, folder)
                if hasil_pdf and sumber != p:
                    # Namanya dikembalikan supaya PDF dan Excel-nya sepasang.
                    tujuan = folder / f"{p.stem}.pdf"
                    hasil_pdf.replace(tujuan)
                    hasil_pdf = tujuan
                if hasil_pdf:
                    dibuat.append(hasil_pdf)
                elif cetak:
                    cetak(f"  (PDF {p.name} dilewati: {pesan})")

    if berkas_pdf_sumber != berkas_utama:
        berkas_pdf_sumber.unlink(missing_ok=True)
    return dibuat
