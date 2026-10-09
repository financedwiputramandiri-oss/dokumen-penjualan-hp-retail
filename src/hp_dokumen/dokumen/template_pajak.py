"""Lembar TEMPLATE PAJAK per PO — bentuknya dari berkas Yosua sendiri.

Yosua mengirim `Template_Baby_Wise_Surabaya.xlsx` (6 Oktober 2026):

> *"template tersebut diambil dari invoice, jadi setiap invoice yang anda
> buat, buatkanlah template nya seperti itu pastikan jumlahnya berupakan
> angka bulat sesuai dengan invoice"*

Berkasnya dibongkar sel per sel, BUKAN ditiru dari tampilannya. Isinya
ternyata `PO 31 Agustus - Baby Wise (Surabaya)`: jumlah kolom G
Rp22.179.900 dan kolom M Rp15.758.819 cocok persis dengan angka yang sudah
tercatat di CLAUDE.md bagian 7 dan 12. Jadi template ini sekaligus
membenarkan pembacaan order sheet oleh program.

Susunan yang ditiru
-------------------
Baris 1 keterangan:  I1 NPWP pembeli | K1 tanggal | L1 nomor | M1 nama pembeli
Baris 2 judul kolom, baris 3 dan seterusnya data, baris terakhir jumlah
(berlatar KUNING).

    A NO                              H TOTAL EXCL. PPN    = Qty x F
    B ART                             I DISC               = Qty x F x tarif
    C (nama barang, tanpa judul)      J DPP                = H - I
    D QTY                             K DPP NILAI LAINNYA  = J x 11/12
    E HARGA JUAL INC. PPN             L PPN 12%            = K x 12%
    F HARGA JUAL EXC. PPN             M DPP + PPN          = J + L
      (Pembulatan)         = E/111%
    G HARGA JUAL INC. PPN  = Qty x E

PPN 12% atas DPP Nilai Lain — dan hasilnya tetap 11%
----------------------------------------------------
`K = J x 11/12` lalu `L = K x 12%` memberi `L = J x 0,11`. Jadi tarif
efektifnya tetap 11%, hanya cara menuliskannya mengikuti aturan DJP sejak
2025: PPN 12% dikenakan atas DPP Nilai Lain sebesar 11/12 DPP. Ini menjawab
"Tarif PPN sementara 11%, perlu dicek" di bagian 9 — angkanya benar, yang
perlu diperbaiki cara penyajiannya. JANGAN diubah jadi `L = J x 11%`
walaupun hasilnya sama: faktur pajak harus memperlihatkan mekanismenya.

Kolom DISC: potongan BERUNTUN, bukan dijumlahkan
------------------------------------------------
Rumus di berkas Yosua:

    =D3*((27.5/100*F3)+(2/100*(F3-(27.5/100*F3))))

yaitu `Qty x F x (0,275 + 0,02 x (1 - 0,275))` = `Qty x F x 28,95%`. Itu
potongan 27,5% LALU 2%, bukan 29,5% — aturan yang sama dengan invoice
(bagian 22). Kedua tarifnya di sana ditulis mati karena template itu untuk
satu customer; di sini keduanya diambil dari order yang sedang diproses.

"Angka bulat" dipenuhi lewat FORMAT, bukan ROUND
-------------------------------------------------
Diukur pada data template itu sendiri:

    rumus asli (tanpa ROUND)     DPP+PPN 15.758.818,95   selisih Rp0,00
    hanya F dibulatkan                   15.758.841,12   selisih Rp22,17
    F + semua sel dibulatkan             15.758.843,00   selisih Rp24,05
    semua kecuali F dibulatkan           15.758.819,89   selisih Rp0,94

Hanya rumus asli yang tepat sama dengan total invoice. Karena itu TIDAK ADA
ROUND di lembar ini; yang membuat angkanya terlihat bulat adalah format
`#,##0` pada kolom M dan baris jumlah — persis seperti berkas Yosua sendiri.
Pelajaran yang sama dengan bagian 43: bulatkan TAMPILANnya, jangan nilainya.
"""
from __future__ import annotations

from datetime import date

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from . import gaya
from .invoice import susun_baris
from ..model import KeputusanNett, Order
from ..nilai_bersih import persen_diskon_efektif, tarif_tambahan_tertulis

NAMA_LEMBAR = "TEMPLATE"

BARIS_JUDUL = 2
BARIS_DATA = 3

# Lebar kolom A..M, dibaca dari berkas Yosua.
LEBAR = {1: 4.86, 2: 9.71, 3: 34.29, 4: 5.86, 5: 11.57, 6: 12.0, 7: 13.0,
         8: 14.71, 9: 13.57, 10: 15.57, 11: 15.14, 12: 13.86, 13: 13.14}

JUDUL = {
    1: "NO", 2: "ART", 3: "", 4: "QTY",
    5: "HARGA JUAL INC. PPN",
    6: "HARGA JUAL EXC. PPN (Pembulatan)",
    7: "HARGA JUAL INC. PPN",
    8: "TOTAL EXCL. PPN", 9: "DISC", 10: "DPP",
    11: "DPP NILAI LAINNYA", 12: "PPN 12%", 13: "DPP + PPN",
}

# Format akuntansi dari berkas Yosua. Dua desimal untuk kolom hitungan,
# bulat untuk kolom yang memang dilihat orang sebagai rupiah utuh.
FMT_BULAT = r'_-* #,##0_-;\-* #,##0_-;_-* "-"??_-;_-@_-'
FMT_DESIMAL = r'_(* #,##0.00_);_(* \(#,##0.00\);_(* "-"??_);_(@_)'
KUNING = PatternFill("solid", fgColor="FFFF00")

# Kolom yang dijumlahkan di baris terakhir (G..M), sama seperti berkas Yosua.
KOLOM_JUMLAH = range(7, 14)


def _tarif_potongan(order: Order, keputusan: KeputusanNett,
                    kotor: float, nett: float) -> tuple[float, float]:
    """(tarif dasar, tarif tambahan) — dipakai menyusun rumus DISC.

    Dipisah dua supaya rumusnya berbentuk sama dengan berkas Yosua, yang
    memperlihatkan potongan beruntun apa adanya. Kalau ordernya TOP, tarif
    tambahannya nol dan rumusnya menyusut sendiri jadi satu potongan.
    """
    efektif = persen_diskon_efektif(kotor, nett)
    tambahan = 0.0
    for k in order.kolom_nett():
        if k.kunci == keputusan.kolom and k.jenis in ("CBD", "COD"):
            tambahan = tarif_tambahan_tertulis(k.judul) or 0.0
            break
    if tambahan and tambahan < 1.0:
        # Dipulihkan dengan MEMBAGI, bukan mengurangi — potongannya
        # beruntun (bagian 22). Mengurangi memberi angka yang meleset.
        dasar = 1.0 - (1.0 - efektif) / (1.0 - tambahan)
    else:
        dasar, tambahan = efektif, 0.0
    return dasar, tambahan


def _angka(x: float) -> str:
    """Tanpa nol yang tidak perlu: 2.0 -> "2", 27.5 -> "27.5".

    Berkas Yosua menulis `2/100`, bukan `2.0/100`. Rumusnya memang sama
    saja bagi Excel, tapi lembar ini dibaca orang.
    """
    return f"{round(x, 6):.6f}".rstrip("0").rstrip(".") or "0"


def _rumus_disc(r: int, dasar: float, tambahan: float) -> str:
    """Bentuknya sengaja sama persis dengan berkas Yosua."""
    d = _angka(dasar * 100)
    if not tambahan:
        return f"=D{r}*({d}/100*F{r})"
    t = _angka(tambahan * 100)
    return f"=D{r}*(({d}/100*F{r})+({t}/100*(F{r}-({d}/100*F{r}))))"


def _kode(teks: str):
    """Kode yang seluruhnya angka ditulis sebagai ANGKA, seperti berkas Yosua.

    Di sana `42052` tersimpan sebagai bilangan sedangkan
    `42032.S (Top/Atasan)` tetap teks.
    """
    t = str(teks or "").strip()
    try:
        return int(t) if t.isdigit() else float(t)
    except ValueError:
        return teks


def buat_template_pajak(ws: Worksheet, order: Order, keputusan: KeputusanNett,
                        customer, perusahaan, pengaturan, nomor: str,
                        tanggal: date | None = None) -> dict:
    """Tulis lembar TEMPLATE untuk satu PO."""
    ws.title = NAMA_LEMBAR
    baris = [b for b in susun_baris(
        order, keputusan,
        pecah_per_ukuran=bool(customer and customer.pecah_per_ukuran),
        akhiran_y=pengaturan.akhiran_y_untuk_angka) if b.qty]

    kotor = sum(b.kotor for b in baris)
    nett = sum(b.nett for b in baris)
    dasar, tambahan = _tarif_potongan(order, keputusan, kotor, nett)

    # ---- baris 1: keterangan pembeli -----------------------------------
    # Dikosongkan kalau belum diketahui. Permintaan Yosua 20 September 2026:
    # jangan menulis tulisan penampung, biar ia sendiri yang mengisi.
    tgl = tanggal or order.tanggal_po or date.today()
    ws.cell(1, 9, (customer.npwp or None) if customer else None)
    sel = ws.cell(1, 11, tgl)
    sel.number_format = "d-mmm"
    for kolom, isi in ((12, nomor or None),
                       (13, (customer.nama_di_dokumen or None) if customer else None)):
        s = ws.cell(1, kolom, isi)
        s.font = Font(bold=True, size=12)

    # ---- baris 2: judul kolom ------------------------------------------
    for kolom, teks in JUDUL.items():
        s = ws.cell(BARIS_JUDUL, kolom, teks or None)
        s.font = Font(bold=kolom >= 4, size=11)
        s.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=True)

    # ---- data ------------------------------------------------------------
    r = BARIS_DATA
    for i, b in enumerate(baris, start=1):
        ws.cell(r, 1, i).number_format = FMT_BULAT
        ws.cell(r, 2, _kode(b.kode))
        ws.cell(r, 3, b.deskripsi.upper())
        ws.cell(r, 4, b.qty)
        ws.cell(r, 5, b.harga).number_format = FMT_BULAT
        # TANPA ROUND, walaupun judulnya berbunyi "(Pembulatan)" — lihat
        # pengukurannya di kepala berkas ini.
        ws.cell(r, 6, f"=E{r}/111%").number_format = FMT_DESIMAL
        ws.cell(r, 7, f"=D{r}*E{r}").number_format = FMT_BULAT
        ws.cell(r, 8, f"=D{r}*F{r}").number_format = FMT_DESIMAL
        ws.cell(r, 9, _rumus_disc(r, dasar, tambahan)).number_format = FMT_DESIMAL
        ws.cell(r, 10, f"=H{r}-I{r}").number_format = FMT_DESIMAL
        ws.cell(r, 11, f"=J{r}*11/12").number_format = FMT_DESIMAL
        ws.cell(r, 12, f"=K{r}*12%").number_format = FMT_DESIMAL
        # Kolom M dibulatkan TAMPILANnya — ini "angka bulat" yang diminta.
        ws.cell(r, 13, f"=J{r}+L{r}").number_format = FMT_BULAT
        r += 1
    akhir = r - 1

    # ---- baris jumlah, berlatar kuning ---------------------------------
    for kolom in KOLOM_JUMLAH:
        huruf = gaya.huruf(kolom)
        s = ws.cell(r, kolom, f"=SUM({huruf}{BARIS_DATA}:{huruf}{akhir})")
        s.number_format = FMT_BULAT
        s.fill = KUNING

    for kolom, lebar in LEBAR.items():
        ws.column_dimensions[gaya.huruf(kolom)].width = lebar
    ws.freeze_panes = f"A{BARIS_DATA}"

    return {
        "baris": len(baris),
        "qty": sum(b.qty for b in baris),
        "kotor": kotor,
        "nett": nett,
        "tarif_dasar": dasar,
        "tarif_tambahan": tambahan,
        "baris_jumlah": r,
    }
