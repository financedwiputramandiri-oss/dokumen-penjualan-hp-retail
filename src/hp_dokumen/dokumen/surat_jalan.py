"""Surat Jalan dan Packing List.

Tata letaknya disalin dari berkas asli CV Dwi Putra Mandiri di Drive
(0250726 KATAMAMA TAPOS, 0310726 MAE BEBE), BUKAN dikarang:

    kop dengan ruang logo -> SURAT JALAN -> kalimat "Diterima dengan baik..."
    -> BRAND -> tabel per blok -> tanda tangan

Satu tabel per blok, bertumpuk dalam satu lembar. Terbukti dari berkas asli:
Katamama hanya punya satu tabel (semua ukurannya sejenis), Mae Bebe punya dua
karena sistem ukurannya berbeda. Judul ukuran tiap tabel diambil dari blok
asalnya sendiri.

DESKRIPSI BARANG memakai tiga kolom yang digabung (C:E), dan judulnya berbahasa
Indonesia: DESKRIPSI BARANG, WARNA, Qty/PCS. Versi sebelumnya memakai PRODUCT
NAME/COLOUR/TOTAL dan itu tidak sesuai faktur asli.
"""
from __future__ import annotations

from datetime import date
from typing import Optional

from openpyxl.worksheet.worksheet import Worksheet

from ..model import Order
from . import gaya

KOL_NO = 1
KOL_KODE = 2
KOL_DESK = 3          # digabung C:E
KOL_WARNA = 6
KOL_UKURAN_MULAI = 7  # G dan seterusnya

BARIS_KOP = 1
BARIS_JUDUL_DOK = 8
BARIS_KALIMAT = 9
BARIS_MEREK = 10
BARIS_TABEL = 12

LEBAR = {1: 5.0, 2: 14.0, 3: 5.6, 4: 5.6, 5: 21.3, 6: 12.0}
KALIMAT = "Diterima dengan baik barang-barang tersebut dibawah ini :"


def _kolom_terakhir(order: Order, tambahan: int) -> int:
    """Kolom paling kanan yang dipakai, mengikuti blok dengan ukuran terbanyak."""
    paling = max((len(_label_terpakai(b)) for b in order.blok), default=1)
    return KOL_UKURAN_MULAI + paling + tambahan  # + kolom Qty (dan kolom gudang)


def _posisi_terpakai(blok) -> list[int]:
    """Posisi kolom ukuran yang benar-benar dipakai blok ini.

    Kolom ukuran ke-9 di order sheet tidak pernah terisi (sisa rancangan lama,
    judulnya cuma angka '9'), dan sebagian blok menyisakan kolom kosong di
    tengah. Kalau semuanya dicetak, Surat Jalan penuh kolom hampa dan tidak
    seperti berkas asli. Aturannya: cetak sampai ukuran terakhir yang ada
    isinya, dan lewati yang judulnya kosong.
    """
    label = list(blok.label_ukuran)
    terakhir = -1
    for b in blok.baris:
        for i, q in enumerate(b.qty_per_ukuran):
            if q:
                terakhir = max(terakhir, i)
    if terakhir < 0:
        return []
    return [i for i in range(terakhir + 1)
            if i < len(label) and str(label[i] or "").strip()]


def _label_terpakai(blok) -> list[str]:
    label = list(blok.label_ukuran)
    return [str(label[i]).strip() for i in _posisi_terpakai(blok)]


def _tulis_tabel(ws: Worksheet, blok, baris: int, kolom_gudang: list[str]) -> int:
    """Satu tabel untuk satu blok. Mengembalikan baris kosong sesudahnya."""
    posisi = _posisi_terpakai(blok)
    label = _label_terpakai(blok)
    n = len(label)
    kol_qty = KOL_UKURAN_MULAI + n
    j = baris

    # judul dua tingkat
    for kolom, teks in ((KOL_NO, "No."), (KOL_KODE, "ARTICLE CODE"), (KOL_WARNA, "WARNA")):
        ws.merge_cells(start_row=j, start_column=kolom, end_row=j + 1, end_column=kolom)
        gaya.sel_judul(ws, j, kolom, teks)
    ws.merge_cells(start_row=j, start_column=KOL_DESK, end_row=j + 1, end_column=KOL_WARNA - 1)
    gaya.sel_judul(ws, j, KOL_DESK, "DESKRIPSI BARANG")
    for i, teks in enumerate(label):
        gaya.sel_judul(ws, j, KOL_UKURAN_MULAI + i, teks)
        gaya.sel_judul(ws, j + 1, KOL_UKURAN_MULAI + i, None)
    gaya.sel_judul(ws, j, kol_qty, "Qty")
    gaya.sel_judul(ws, j + 1, kol_qty, "PCS")
    for i, teks in enumerate(kolom_gudang):
        ws.merge_cells(start_row=j, start_column=kol_qty + 1 + i,
                       end_row=j + 1, end_column=kol_qty + 1 + i)
        gaya.sel_judul(ws, j, kol_qty + 1 + i, teks)

    r = j + 2
    urut = 0
    for b in blok.baris:
        if b.qty <= 0:
            continue
        urut += 1
        gaya.sel_isi(ws, r, KOL_NO, urut, rata="center")
        gaya.sel_isi(ws, r, KOL_KODE, b.kode)
        ws.merge_cells(start_row=r, start_column=KOL_DESK, end_row=r, end_column=KOL_WARNA - 1)
        gaya.sel_isi(ws, r, KOL_DESK, b.nama)
        gaya.sel_isi(ws, r, KOL_WARNA, b.warna)
        for kolom_ke, i in enumerate(posisi):
            q = b.qty_per_ukuran[i] if i < len(b.qty_per_ukuran) else 0
            gaya.sel_isi(ws, r, KOL_UKURAN_MULAI + kolom_ke, q or None, rata="center")
        gaya.sel_isi(ws, r, kol_qty, b.qty, rata="center", tebal=True)
        for i in range(len(kolom_gudang)):
            gaya.sel_isi(ws, r, kol_qty + 1 + i, None)
        r += 1

    gaya.beri_garis(ws, j, KOL_NO, r - 1, kol_qty + len(kolom_gudang))
    return r + 1


def _bangun(ws: Worksheet, order: Order, customer, perusahaan, nomor: str,
            tanggal_dokumen: Optional[date], judul: str,
            kolom_gudang: list[str]) -> None:
    tanggal = tanggal_dokumen or order.tanggal_po or date.today()
    kolom_akhir = _kolom_terakhir(order, len(kolom_gudang))
    for kolom, lebar in LEBAR.items():
        ws.column_dimensions[gaya.huruf(kolom)].width = lebar

    gaya.kop_dpm(
        ws, perusahaan,
        nama_customer=(customer.nama_di_dokumen if customer else "") or order.customer_kunci,
        alamat_customer=(customer.alamat if customer else ""),
        tanggal_dokumen=tanggal,
        baris_mulai=BARIS_KOP,
        kolom_kanan=KOL_UKURAN_MULAI,
    )

    ws.merge_cells(start_row=BARIS_JUDUL_DOK, start_column=1,
                   end_row=BARIS_JUDUL_DOK, end_column=kolom_akhir)
    gaya.judul(ws, BARIS_JUDUL_DOK, 1, judul, ukuran=14)

    gaya.sel_isi(ws, BARIS_KALIMAT, 1, KALIMAT)
    gaya.sel_isi(ws, BARIS_KALIMAT, KOL_UKURAN_MULAI + 1, f"No. {nomor}", tebal=True)
    gaya.sel_isi(ws, BARIS_KALIMAT, kolom_akhir, order.qty, rata="center", tebal=True)
    gaya.sel_isi(ws, BARIS_MEREK, 1, "BRAND : HAPPY PUMPKIN", tebal=True)

    r = BARIS_TABEL
    for blok in order.blok:
        if any(b.qty > 0 for b in blok.baris):
            r = _tulis_tabel(ws, blok, r, kolom_gudang)

    r += 1
    gaya.blok_tanda_tangan(ws, r, [2, 5, 10], ["Pengirim :", "Penerima : ", "Mengetahui :"])
    gaya.siapkan_cetak(ws, kolom_akhir, landscape=True)


def buat_surat_jalan(ws, order, customer, perusahaan, nomor: str, tanggal_dokumen=None) -> None:
    _bangun(ws, order, customer, perusahaan, nomor, tanggal_dokumen, "SURAT JALAN", [])


def buat_packing_list(ws, order, customer, perusahaan, nomor: str, tanggal_dokumen=None) -> None:
    """Sama seperti Surat Jalan, ditambah dua kolom yang diisi gudang."""
    _bangun(ws, order, customer, perusahaan, nomor, tanggal_dokumen, "PACKING LIST",
            ["JUMLAH DIKIRIM", "NO. KOLI"])
