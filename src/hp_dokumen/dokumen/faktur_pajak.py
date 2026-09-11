"""Faktur Pajak — lembar data siap ketik ke Coretax.

Bukan untuk dicetak sebagai faktur. Isinya data yang tinggal disalin:
nama & NPWP pembeli, alamat, nomor referensi, tanggal, DPP, PPN, total,
ditambah rincian per artikel.

Harga di order sheet SUDAH termasuk PPN, jadi DPP dihitung mundur:
    DPP = Total / (1 + tarif)
    PPN = Total - DPP
"""
from __future__ import annotations

from datetime import date
from typing import Optional

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from ..model import KeputusanNett, Order
from . import gaya
from .invoice import susun_baris

KOLOM_TERAKHIR = 6


def _label(ws, r, teks, nilai, *, bold=False, format_angka=None, catatan=""):
    a = ws.cell(r, 1, teks)
    a.font = Font(name=gaya.FONT, size=10, bold=True)
    b = ws.cell(r, 2, nilai)
    b.font = Font(name=gaya.FONT, size=10, bold=bold)
    if format_angka:
        b.number_format = format_angka
    else:
        b.alignment = Alignment(horizontal="left")
    if catatan:
        c = ws.cell(r, 4, catatan)
        c.font = Font(name=gaya.FONT, size=8, italic=True)
    return r + 1


def buat_faktur_pajak(
    ws: Worksheet,
    order: Order,
    keputusan: KeputusanNett,
    customer,
    perusahaan,
    pengaturan,
    nomor_referensi: str,
    *,
    tanggal_dokumen: Optional[date] = None,
) -> dict:
    tanggal_dokumen = tanggal_dokumen or order.tanggal_po
    tarif = pengaturan.tarif_ppn
    total = keputusan.nett_total
    dpp = total / (1 + tarif) if tarif else total
    ppn = total - dpp

    r = 1
    t = ws.cell(r, 1, "DATA FAKTUR PAJAK — SIAP KETIK KE CORETAX")
    t.font = Font(name=gaya.FONT, size=13, bold=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=KOLOM_TERAKHIR)
    r += 1
    s = ws.cell(r, 1, "Lembar ini BUKAN faktur pajak. Isinya data untuk diketik ke Coretax.")
    s.font = Font(name=gaya.FONT, size=9, italic=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=KOLOM_TERAKHIR)
    r += 2

    r = _label(ws, r, "PENJUAL", "")
    r = _label(ws, r, "Nama", perusahaan.nama)
    r = _label(ws, r, "Alamat", ", ".join(perusahaan.alamat_baris[:2]))
    r += 1

    r = _label(ws, r, "PEMBELI", "")
    r = _label(
        ws, r, "Nama",
        (customer.nama_di_dokumen if customer else "") or "(BELUM DIISI)",
        catatan="" if (customer and customer.nama_di_dokumen) else "isi di config/customer.csv",
    )
    r = _label(
        ws, r, "Alamat",
        (customer.alamat if customer else "") or "(BELUM DIISI)",
        catatan="" if (customer and customer.alamat) else "isi di config/customer.csv",
    )
    r = _label(
        ws, r, "NPWP",
        (customer.npwp if customer else "") or "(BELUM DIISI)",
        catatan="" if (customer and customer.npwp) else "NPWP tidak ada di Drive, harus diisi manual",
    )
    r += 1

    r = _label(ws, r, "TRANSAKSI", "")
    r = _label(ws, r, "Nomor referensi", nomor_referensi)
    r = _label(ws, r, "Tanggal", gaya.tanggal_indonesia(tanggal_dokumen))
    r = _label(ws, r, "PO / tab sumber", order.nama_tab)
    r = _label(ws, r, "Cara bayar", keputusan.cara_bayar, catatan=f"nett dari {keputusan.kolom_sumber}")
    r = _label(ws, r, "Jumlah barang (pcs)", order.qty, format_angka=gaya.ANGKA)
    r += 1

    r = _label(ws, r, "NILAI", "")
    r = _label(ws, r, "Total (termasuk PPN)", total, bold=True, format_angka=gaya.RUPIAH)
    r = _label(ws, r, "DPP (Dasar Pengenaan Pajak)", dpp, bold=True, format_angka=gaya.RUPIAH_DESIMAL)
    r = _label(ws, r, f"PPN {tarif:.0%}", ppn, bold=True, format_angka=gaya.RUPIAH_DESIMAL)
    r += 1

    p = ws.cell(
        r, 1,
        f"PERHATIAN: tarif PPN {tarif:.0%} "
        + ("sudah dikonfirmasi." if pengaturan.ppn_dikonfirmasi
           else "MASIH SEMENTARA dan BELUM dikonfirmasi. Cek aturan yang berlaku sebelum lapor."),
    )
    p.font = Font(name=gaya.FONT, size=9, bold=not pengaturan.ppn_dikonfirmasi)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=KOLOM_TERAKHIR)
    r += 3

    j = ws.cell(r, 1, "RINCIAN PER ARTIKEL")
    j.font = Font(name=gaya.FONT, size=11, bold=True)
    r += 1

    pecah = bool(customer and customer.pecah_per_ukuran)
    baris_inv = susun_baris(
        order, keputusan,
        pecah_per_ukuran=pecah,
        akhiran_y=pengaturan.akhiran_y_untuk_angka,
    )
    judul = ["No.", "ARTICLE CODE", "NAMA BARANG", "Qty", "Harga Satuan", "Jumlah (nett)"]
    gaya.baris_judul_tabel(ws, r, judul)
    awal = r
    r += 1
    for no, x in enumerate(baris_inv, start=1):
        ws.cell(r, 1, no).alignment = Alignment(horizontal="center")
        ws.cell(r, 2, x.kode)
        ws.cell(r, 3, x.deskripsi)
        c = ws.cell(r, 4, x.qty); c.alignment = Alignment(horizontal="center"); c.number_format = gaya.ANGKA
        c = ws.cell(r, 5, x.harga); c.number_format = gaya.RUPIAH
        c = ws.cell(r, 6, x.nett); c.number_format = gaya.RUPIAH
        for cc in range(1, KOLOM_TERAKHIR + 1):
            ws.cell(r, cc).font = Font(name=gaya.FONT, size=9)
        r += 1
    # baris total rincian
    sel = ws.cell(r, 3, "TOTAL")
    sel.font = Font(name=gaya.FONT, size=9, bold=True)
    sel.alignment = Alignment(horizontal="right")
    c = ws.cell(r, 4, sum(x.qty for x in baris_inv))
    c.font = Font(name=gaya.FONT, size=9, bold=True)
    c.alignment = Alignment(horizontal="center")
    c.number_format = gaya.ANGKA
    c = ws.cell(r, 6, sum(x.nett for x in baris_inv))
    c.font = Font(name=gaya.FONT, size=9, bold=True)
    c.number_format = gaya.RUPIAH
    gaya.beri_garis(ws, awal, 1, r, KOLOM_TERAKHIR)

    gaya.atur_lebar(ws, {1: 5, 2: 20, 3: 48, 4: 10, 5: 15, 6: 17})
    gaya.siapkan_cetak(ws, KOLOM_TERAKHIR, landscape=False)

    return {"dpp": dpp, "ppn": ppn, "total": total, "tarif": tarif}
