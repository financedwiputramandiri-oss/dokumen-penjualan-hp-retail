"""Rumus Excel di dalam dokumen, dan lembar MASTER HARGA yang dirujuknya.

Permintaan Yosua 20 September 2026: Invoice dan Surat Jalan digabung dalam
SATU berkas Excel di lembar berbeda, ditambah lembar berisi master harga
sebagaimana adanya saat dokumen itu dibuat, dan perhitungan invoice ditulis
sebagai RUMUS supaya bisa dilihat dan ditelusuri.

Satu hal yang SENGAJA tidak dijadikan rumus: kolom "Nilai Diskon" per baris.
------------------------------------------------------------------------
Godaannya besar, sebab `qty x harga x persen` terlihat paling jelas. Tapi
nilai bersih di order sheet TIDAK dihitung ulang dari persentase — ia dibaca
apa adanya dari kolom nilai bersih (Aturan 2, CLAUDE.md bagian 7). Diuji pada
dua PO sungguhan 20 September 2026:

    PO 17 September - Satu Sama (Veteran)  22% + 1.5%  selisih Rp1
    PO 17 September - Baby Fame            25% + 1.5%  selisih Rp2

Kecil, tapi artinya total faktur tidak lagi sama persis dengan baris TOTAL
order sheet — padahal kecocokan sampai rupiah terakhir itulah yang dijaga
seluruh program ini, dan yang membuat pencocokan berani MENOLAK menerbitkan
dokumen. Jadi Nilai Diskon tetap berupa ANGKA dari order sheet, sedangkan
semua yang bisa diturunkan darinya ditulis sebagai rumus:

    Harga satuan   VLOOKUP ke lembar MASTER HARGA
    Deskripsi      VLOOKUP ke lembar MASTER HARGA
    Jumlah         = Qty x Harga - Nilai Diskon
    Subtotal       = SUMPRODUCT(Qty, Harga)
    Diskon         = SUM(Nilai Diskon)
    Total          = Subtotal - Diskon
    DPP / PPN      diturunkan dari Total memakai tarif di lembar MASTER HARGA
    Qty Surat Jalan = SUM(kolom ukuran)
"""
from __future__ import annotations

from datetime import date

from . import gaya

TAB_INVOICE = "INVOICE"
TAB_SURAT_JALAN = "SURAT JALAN"
TAB_MASTER = "MASTER HARGA"

# Letak di lembar MASTER HARGA. Tarif PPN ditaruh di sel tersendiri supaya
# rumus DPP/PPN di faktur merujuk ke satu tempat, bukan menanam angka 0,11
# di belasan sel.
BARIS_TARIF = 3
KOLOM_TARIF = 3            # C3
BARIS_JUDUL_MASTER = 5
BARIS_DATA_MASTER = 6

_M = f"'{TAB_MASTER}'"
SEL_TARIF = f"{_M}!$C${BARIS_TARIF}"


def vlookup(kolom_kode: str, baris: int, kolom_hasil: int, bawaan: str) -> str:
    """VLOOKUP ke lembar MASTER HARGA, aman kalau kodenya tidak ketemu.

    Dibungkus IFERROR dengan sengaja: kode yang tidak terdaftar tidak boleh
    membuat faktur penuh `#N/A` di depan customer. Kalau itu terjadi,
    pencocokan sudah lebih dulu memperingatkannya di layar.
    """
    return (f"=IFERROR(VLOOKUP(${kolom_kode}{baris},{_M}!$A:$C,"
            f"{kolom_hasil},FALSE),{bawaan})")


def harga(baris: int, kolom_kode: str = "B") -> str:
    return vlookup(kolom_kode, baris, 3, "0")


def nama_barang(baris: int, kolom_kode: str = "B") -> str:
    return vlookup(kolom_kode, baris, 2, '""')


def jumlah_baris(baris: int, kolom_qty: str, kolom_harga: str,
                 kolom_diskon: str | None) -> str:
    """Jumlah per baris = qty x harga, dikurangi nilai diskon kalau ada."""
    pokok = f"{kolom_qty}{baris}*{kolom_harga}{baris}"
    return f"={pokok}-{kolom_diskon}{baris}" if kolom_diskon else f"={pokok}"


def subtotal(awal: int, akhir: int, kolom_qty: str, kolom_harga: str) -> str:
    return (f"=SUMPRODUCT(${kolom_qty}${awal}:${kolom_qty}${akhir},"
            f"${kolom_harga}${awal}:${kolom_harga}${akhir})")


def jumlahkan(awal: int, akhir: int, kolom: str) -> str:
    return f"=SUM(${kolom}${awal}:${kolom}${akhir})"


def selisih(kolom: str, baris_a: int, baris_b: int) -> str:
    return f"={kolom}{baris_a}-{kolom}{baris_b}"


def tambah(kolom: str, baris_a: int, baris_b: int) -> str:
    return f"={kolom}{baris_a}+{kolom}{baris_b}"


def dpp(kolom: str, baris_total: int) -> str:
    """DPP dihitung MUNDUR — harga di order sheet sudah termasuk PPN."""
    return f"=ROUND({kolom}{baris_total}/(1+{SEL_TARIF}),0)"


def qty_surat_jalan(baris: int, kolom_awal: str, kolom_akhir: str) -> str:
    return f"=SUM({kolom_awal}{baris}:{kolom_akhir}{baris})"


def tulis_master(ws, master: dict, perusahaan, tarif_ppn: float,
                 tanggal: date | None = None) -> None:
    """Isi lembar MASTER HARGA: harga yang BERLAKU saat dokumen ini dibuat.

    Gunanya bukan cuma untuk VLOOKUP. Harga retail berubah sepanjang tahun
    (CLAUDE.md bagian 18: Milo Set 49.400 -> 61.000 pada Agustus 2026), dan
    faktur lama tidak boleh ikut berubah hanya karena master hari ini
    berbeda. Dengan salinan ini, tiap faktur membawa buktinya sendiri.
    """
    from openpyxl.styles import Font

    ws.cell(1, 1, "MASTER HARGA - salinan tab 'Harga Retail' order sheet").font = (
        Font(bold=True, size=12)
    )
    ws.cell(2, 1, f"Diambil saat dokumen ini dibuat: "
                  f"{gaya.tanggal_indonesia(tanggal or date.today())}")
    ws.cell(BARIS_TARIF, 1, "Tarif PPN yang dipakai faktur ini")
    sel = ws.cell(BARIS_TARIF, KOLOM_TARIF, tarif_ppn)
    sel.number_format = "0%"
    ws.cell(BARIS_TARIF, KOLOM_TARIF + 1,
            f"diproses lewat {getattr(perusahaan, 'nama', '')}"
            + ("" if tarif_ppn else " — tidak mengenakan PPN"))

    for i, judul in enumerate(("KODE ARTIKEL", "NAMA BARANG",
                               "HARGA (SUDAH TERMASUK PPN)"), start=1):
        gaya.sel_judul(ws, BARIS_JUDUL_MASTER, i, judul)

    r = BARIS_DATA_MASTER
    for kode in sorted(master):
        nama, nilai = master[kode]
        ws.cell(r, 1, kode)
        ws.cell(r, 2, nama)
        sel = ws.cell(r, 3, nilai)
        sel.number_format = gaya.FORMAT_RP
        r += 1

    gaya.beri_garis(ws, BARIS_JUDUL_MASTER, 1, max(r - 1, BARIS_JUDUL_MASTER), 3)
    for huruf, lebar in (("A", 26.0), ("B", 38.0), ("C", 24.0), ("D", 34.0)):
        ws.column_dimensions[huruf].width = lebar
    ws.freeze_panes = f"A{BARIS_DATA_MASTER}"
