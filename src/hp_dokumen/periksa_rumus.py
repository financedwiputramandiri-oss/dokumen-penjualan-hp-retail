"""Peringatan untuk rumus order sheet yang salah atau tidak tepat.

Permintaan Yosua 20 Oktober 2026, sesudah audit bagian 49:
*"berikan warning ketika di order sheet ada formula yang tidak sesuai dan
tidak tepat"*.

Audit itu menemukan cacat rumus yang TIDAK PERNAH terlihat dari angkanya —
`=SUM(N3:S3)` yang berhenti sebelum kolom ukuran terakhir tampak wajar
sampai Sales mengisi kolom yang terlewat, lalu qty-nya hilang diam-diam
(Buchi Kids September 2026: 12 pcs, Rp890.400). Modul ini memunculkan
peringatannya SAAT dokumen dibuat, bukan hanya saat audit dijalankan.

Dua tingkat, dan pemisahan ini yang membuat peringatannya layak dibaca:

  GENTING  angkanya SUDAH salah sekarang
  POTENSI  rumusnya salah tapi kolom/baris yang terlewat masih kosong

Tanpa pemisahan itu tiap tab memunculkan ratusan baris peringatan dan
orang berhenti membacanya — pelajaran alarm palsu CLAUDE.md bagian 14.
"""
from __future__ import annotations

import re
from typing import Optional

from openpyxl.utils import get_column_letter

GALAT_EXCEL = ("#REF!", "#N/A", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!")

_SEL = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_RENTANG = re.compile(r"(\$?[A-Z]{1,3}\$?\d+)\s*:\s*(\$?[A-Z]{1,3}\$?\d+)")
_TUNGGAL = re.compile(r"(?<![:\w$])\$?([A-Z]{1,3})\$?(\d+)(?![:\d])")


def nomor_kolom(huruf: str) -> int:
    n = 0
    for c in huruf:
        n = n * 26 + (ord(c) - 64)
    return n


def rentang_rumus(rumus: str) -> list[tuple[int, int, int, int]]:
    """Daftar (kol1, bar1, kol2, bar2) dari tiap rentang di dalam rumus."""
    out = []
    for a, b in _RENTANG.findall(rumus):
        ma, mb = _SEL.match(a), _SEL.match(b)
        if ma and mb:
            out.append((nomor_kolom(ma.group(1)), int(ma.group(2)),
                        nomor_kolom(mb.group(1)), int(mb.group(2))))
    return out


def baris_tercakup(rumus: str, kolom: int) -> set[int]:
    """Baris mana saja di `kolom` yang ikut dijumlah rumus ini.

    Rumus bisa menyebut sel SATUAN, bukan cuma rentang — contoh nyata
    `=sum(W3,W7,W11:W16)`. Yang hanya membaca rentang akan melaporkan W3
    dan W7 "tidak ikut dijumlah" padahal ikut.
    """
    ikut: set[int] = set()
    for c1, r1, c2, r2 in rentang_rumus(rumus):
        if c1 == c2 == kolom:
            ikut |= set(range(r1, r2 + 1))
    for huruf, nomor in _TUNGGAL.findall(rumus):
        if nomor_kolom(huruf) == kolom:
            ikut.add(int(nomor))
    return ikut


def _angka(v) -> float:
    return float(v) if isinstance(v, (int, float)) else 0.0


def _rumus(ws, r: int, c: int) -> Optional[str]:
    v = ws.cell(r, c).value
    return v if isinstance(v, str) and v.startswith("=") else None


def periksa(ws_nilai, ws_rumus, blok_list, maks_baris: int = 2000) -> list[str]:
    """Peringatan rumus untuk satu tab.

    `ws_nilai` lembar berisi HASIL hitung, `ws_rumus` lembar berisi RUMUS-nya.
    Keduanya boleh worksheet openpyxl atau LembarNilai dari Sheets API.
    Kalau `ws_rumus` None, hanya pemeriksaan yang tidak butuh rumus yang
    dijalankan.
    """
    pesan: list[str] = []
    if not blok_list:
        return pesan
    tata0 = blok_list[0].tata
    if tata0 is None or getattr(tata0, "hanya_qty", False):
        return pesan
    kol_t = tata0.kol_ato_total
    if not kol_t:
        return pesan

    baris_data = sorted(b.baris_sheet for blk in blok_list for b in blk.baris)
    if not baris_data:
        return pesan

    # ---- 1. galat Excel di kolom nilai bersih --------------------------
    kolom_nett = [k.kolom for k in (tata0.nett or [])]
    data_set = set(baris_data)
    galat_data, galat_lain = [], []
    for c in kolom_nett:
        for r in range(min(baris_data), min(maks_baris, max(baris_data) + 40) + 1):
            v = ws_nilai.cell(r, c).value
            if isinstance(v, str) and v.strip() in GALAT_EXCEL:
                (galat_data if r in data_set else galat_lain).append(
                    f"{get_column_letter(c)}{r}")
    if galat_data:
        pesan.append(
            "RUMUS SALAH (GENTING): kolom nilai bersih berisi galat di BARIS "
            "DATA — " + ", ".join(galat_data[:6])
            + (f" dan {len(galat_data)-6} sel lain" if len(galat_data) > 6 else "")
            + ". Nilai bersih baris itu terbaca NOL, jadi cara bayar bisa "
              "salah pilih. Minta Sales membetulkan rumusnya."
        )
    if galat_lain:
        pesan.append(
            "RUMUS SALAH: baris TOTAL kolom nilai bersih berisi galat — "
            + ", ".join(galat_lain[:6])
            + (f" dan {len(galat_lain)-6} sel lain" if len(galat_lain) > 6 else "")
            + ". Dokumen tetap benar (nilai bersih dibaca per baris), tapi "
              "total di order sheet tidak bisa dipakai untuk mencocokkan."
        )

    if ws_rumus is None:
        return pesan

    # ---- 2. rumus qty per baris menjangkau seluruh kolom ukuran --------
    pendek_berdampak: list[str] = []
    pendek_kosong: list[str] = []
    qty_hilang = 0.0
    berhenti_di: set[int] = set()
    for blk in blok_list:
        tata = blk.tata or tata0
        for b in blk.baris:
            r = b.baris_sheet
            f = _rumus(ws_rumus, r, tata.kol_ato_total or kol_t)
            if not f:
                continue
            for c1, r1, c2, r2 in rentang_rumus(f):
                if r1 != r or r2 != r or c1 != tata.ato_mulai:
                    continue
                if c2 >= tata.ato_selesai:
                    continue
                berhenti_di.add(c2)
                hilang = sum(_angka(ws_nilai.cell(r, c).value)
                             for c in range(c2 + 1, tata.ato_selesai + 1))
                if hilang:
                    qty_hilang += hilang
                    pendek_berdampak.append(
                        f"{get_column_letter(tata.kol_ato_total or kol_t)}{r}")
                else:
                    pendek_kosong.append(str(r))
    if pendek_berdampak:
        pesan.append(
            f"RUMUS SALAH (GENTING): rumus qty di {len(pendek_berdampak)} baris "
            f"berhenti sebelum kolom ukuran terakhir, "
            f"{qty_hilang:g} pcs TIDAK ikut terhitung — "
            + ", ".join(pendek_berdampak[:6])
            + (f" dan {len(pendek_berdampak)-6} baris lain"
               if len(pendek_berdampak) > 6 else "")
            + f". Seharusnya dijumlah sampai kolom "
              f"{get_column_letter(tata0.ato_selesai)}."
        )
    elif pendek_kosong:
        henti = ", ".join(get_column_letter(c) for c in sorted(berhenti_di))
        pesan.append(
            f"RUMUS KURANG TEPAT: rumus qty {len(pendek_kosong)} baris berhenti "
            f"di kolom {henti}, bukan {get_column_letter(tata0.ato_selesai)}. "
            "Belum berdampak karena kolom sisanya masih kosong, tapi qty akan "
            "HILANG DIAM-DIAM begitu Sales mengisinya."
        )

    # ---- 3. baris TOTAL menjangkau seluruh baris data ------------------
    # Satu tab bisa punya beberapa baris TOTAL (satu per blok, lalu satu
    # TOTAL seluruh PO). Yang diperiksa yang TERAKHIR, dan yang MENJUMLAH
    # MENURUN diutamakan — yang mendatar memakai aturan berbeda.
    menurun, mendatar = [], []
    for r in range(min(baris_data), min(maks_baris, max(baris_data) + 60) + 1):
        if r in data_set:
            continue
        f = _rumus(ws_rumus, r, kol_t)
        if not f or not f.upper().startswith("=SUM"):
            continue
        rg = rentang_rumus(f)
        if any(c1 == c2 == kol_t for c1, _, c2, _ in rg):
            menurun.append(r)
        elif any(r1 == r2 == r for _, r1, _, r2 in rg):
            mendatar.append(r)

    def periksa_menurun(baris_sel: int, kolom: int):
        f = _rumus(ws_rumus, baris_sel, kolom)
        if not f:
            return
        ikut = baris_tercakup(f, kolom)
        luput = [r for r in baris_data
                 if r not in ikut and _angka(ws_nilai.cell(r, kolom).value)]
        if luput:
            hilang = sum(_angka(ws_nilai.cell(r, kolom).value) for r in luput)
            pesan.append(
                f"RUMUS SALAH (GENTING): baris TOTAL "
                f"{get_column_letter(kolom)}{baris_sel} tidak menjumlah "
                f"{len(luput)} baris data (baris {luput[0]}-{luput[-1]}), "
                f"{hilang:g} pcs terlewat. Rumusnya: {f}"
            )

    if menurun:
        periksa_menurun(menurun[-1], kol_t)
    elif mendatar:
        bt = mendatar[-1]
        f = _rumus(ws_rumus, bt, kol_t) or ""
        for c1, r1, c2, r2 in rentang_rumus(f):
            if r1 == r2 == bt and c1 == tata0.ato_mulai and c2 < tata0.ato_selesai:
                hilang = sum(_angka(ws_nilai.cell(bt, c).value)
                             for c in range(c2 + 1, tata0.ato_selesai + 1))
                if hilang:
                    pesan.append(
                        f"RUMUS SALAH (GENTING): baris TOTAL "
                        f"{get_column_letter(kol_t)}{bt} berhenti di kolom "
                        f"{get_column_letter(c2)}, {hilang:g} pcs terlewat."
                    )
        for c in range(tata0.ato_mulai, tata0.ato_selesai + 1):
            periksa_menurun(bt, c)
    return pesan
