"""Aturan 1 — pindai SELURUH isi tiap tab, tanpa kecuali.

Satu tab PO bisa berisi beberapa tabel bertumpuk. Tiap tabel punya baris judul
`ARTICLE CODE` sendiri dan sistem ukurannya sendiri.

Versi lama pernah hanya menangkap blok pertama tiap tab dan kehilangan
Rp147.829.100. Modul ini memindai sampai baris terakhir.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Optional

import openpyxl

from .konfigurasi import DaftarCustomer, pecah_nama_tab
from .model import JUMLAH_KOLOM_UKURAN, Baris, Blok, Order, TotalSheet
from .ukuran import rapikan_label

# Nomor kolom (1 = A)
KOL_ARTICLE = 1      # A
KOL_NAMA = 2         # B
KOL_WARNA = 3        # C
KOL_ORI_MULAI = 4    # D..L  ORIGINAL PO (9 kolom ukuran)
KOL_ORI_TOTAL = 13   # M
KOL_ATO_MULAI = 14   # N..V  AVAILABLE TO ORDER (9 kolom ukuran) <- dasar dokumen
KOL_ATO_TOTAL = 23   # W
KOL_HARGA = 24       # X   PRICE W/ VAT
KOL_ORI_VALUE = 25   # Y
KOL_ATO_VALUE = 26   # Z   nilai sebelum diskon
KOL_LOSSES = 27      # AA
KOL_DISC = 28        # AB
KOL_TOTAL_VALUE = 29  # AC  nett TOP/Tempo
KOL_CBD = 30         # AD  nett CBD
KOL_COD = 31         # AE  nett COD
KOL_NOTE = 32        # AF

TAB_BUKAN_PO = {"harga retail"}


def angka(nilai) -> float:
    """Baca sel jadi angka. Sel kosong, teks, atau #REF! dianggap 0."""
    if nilai is None:
        return 0.0
    if isinstance(nilai, bool):
        return 0.0
    if isinstance(nilai, (int, float)):
        return float(nilai)
    teks = str(nilai).strip()
    if not teks or teks.startswith("#"):
        return 0.0
    teks = teks.replace("Rp", "").replace(" ", "")
    # format Indonesia: 1.234.567,89
    if "," in teks and "." in teks:
        teks = teks.replace(".", "").replace(",", ".")
    elif "," in teks:
        teks = teks.replace(",", ".")
    try:
        return float(teks)
    except ValueError:
        return 0.0


def _angka_atau_none(nilai) -> Optional[float]:
    """Untuk kolom CBD/COD: bedakan 'sel kosong' dari 'terisi nol'.

    Hanya sel yang benar-benar berisi bilangan dihitung sebagai terisi.
    Teks, rumus error, dan sel kosong dianggap BELUM terisi.
    """
    if isinstance(nilai, bool) or nilai is None:
        return None
    if isinstance(nilai, (int, float)):
        return float(nilai)
    return None


def _teks(nilai) -> str:
    if nilai is None:
        return ""
    if isinstance(nilai, float) and nilai.is_integer():
        return str(int(nilai))
    return str(nilai).strip()


def _baris_judul(ws) -> list[int]:
    """Cari SEMUA baris judul di kolom A. Tiap sel berisi 'ARTICLE' = blok baru."""
    hasil = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, KOL_ARTICLE).value
        if v is not None and "ARTICLE" in str(v).upper():
            hasil.append(r)
    return hasil


def pindai_tab(ws, kolom_qty_mulai: int = KOL_ATO_MULAI) -> tuple[list[Blok], TotalSheet, list[str]]:
    """Pindai satu tab jadi daftar blok + baris total milik sheet itu sendiri."""
    peringatan: list[str] = []
    judul = _baris_judul(ws)
    blok_list: list[Blok] = []
    baris_akhir_data = 0

    for nomor, h in enumerate(judul, start=1):
        label = [
            rapikan_label(ws.cell(h + 1, KOL_ORI_MULAI + i).value)
            for i in range(JUMLAH_KOLOM_UKURAN)
        ]
        blok = Blok(nomor=nomor, baris_judul=h, label_ukuran=label)

        r = h + 2
        while r <= ws.max_row:
            kode = ws.cell(r, KOL_ARTICLE).value
            if kode is None or str(kode).strip() == "":
                break
            qty = [
                int(round(angka(ws.cell(r, kolom_qty_mulai + i).value)))
                for i in range(JUMLAH_KOLOM_UKURAN)
            ]
            if sum(qty) > 0:
                blok.baris.append(
                    Baris(
                        baris_sheet=r,
                        kode=_teks(kode),
                        nama=_teks(ws.cell(r, KOL_NAMA).value),
                        warna=_teks(ws.cell(r, KOL_WARNA).value),
                        qty_per_ukuran=qty,
                        harga=angka(ws.cell(r, KOL_HARGA).value),
                        nilai_kotor=angka(ws.cell(r, KOL_ATO_VALUE).value),
                        total_value=angka(ws.cell(r, KOL_TOTAL_VALUE).value),
                        disc_cbd=_angka_atau_none(ws.cell(r, KOL_CBD).value),
                        disc_cod=_angka_atau_none(ws.cell(r, KOL_COD).value),
                        disc_persen=angka(ws.cell(r, KOL_DISC).value),
                        catatan=_teks(ws.cell(r, KOL_NOTE).value) if ws.max_column >= KOL_NOTE else "",
                    )
                )
            r += 1

        baris_akhir_data = max(baris_akhir_data, r)
        if blok.baris:
            blok_list.append(blok)
        # blok tanpa baris berisi tetap dicatat sebagai peringatan, bukan didiamkan
        elif h <= ws.max_row:
            peringatan.append(f"Blok di baris {h} tidak punya baris dengan qty > 0, dilewati.")

        # periksa label ganda dalam satu blok (bisa membuat invoice per ukuran salah gabung)
        terpakai = [x for x in label if x]
        ganda = {x for x in terpakai if terpakai.count(x) > 1}
        if ganda:
            peringatan.append(
                f"Blok baris {h}: label ukuran kembar {sorted(ganda)}. "
                "Periksa baris judul di order sheet."
            )

    # ---- baris TOTAL milik order sheet -------------------------------------
    # Terletak tepat di bawah blok terakhir. Order sheet hanya menaruh satu
    # baris total untuk seluruh tab, bukan per blok.
    total = TotalSheet(None, None, None, None)
    for r in range(baris_akhir_data, min(baris_akhir_data + 4, ws.max_row) + 1):
        q = angka(ws.cell(r, KOL_ATO_TOTAL).value)
        g = angka(ws.cell(r, KOL_ATO_VALUE).value)
        if q > 0 or g > 0:
            total = TotalSheet(
                baris_sheet=r,
                qty=int(round(q)),
                nilai_kotor=g,
                total_value=angka(ws.cell(r, KOL_TOTAL_VALUE).value),
            )
            break
    if total.baris_sheet is None:
        peringatan.append(
            "Baris TOTAL milik order sheet tidak ditemukan di tab ini, "
            "jadi hasil tidak bisa dicocokkan otomatis."
        )
    return blok_list, total, peringatan


def _tanggal(nama_tab: str, tahun_bawaan: int) -> Optional[date]:
    tg, _ = pecah_nama_tab(nama_tab)
    if not tg:
        return None
    hari, bulan, tahun = tg
    try:
        return date(tahun or tahun_bawaan, bulan, hari)
    except ValueError:
        return None


def baca_order_sheet(
    berkas: Path,
    daftar_customer: DaftarCustomer,
    tahun_bawaan: int = 2026,
) -> list[Order]:
    """Baca berkas order sheet (.xlsx) jadi daftar Order, satu per tab PO."""
    wb = openpyxl.load_workbook(berkas, data_only=True, read_only=False)
    hasil: list[Order] = []
    for ws in wb.worksheets:
        nama = ws.title
        if nama.strip().lower() in TAB_BUKAN_PO:
            continue
        # Tab 'Packing List ...' tidak punya blok ORIGINAL PO, qty ada di D..L
        packing = nama.strip().lower().startswith("packing list")
        kolom_qty = KOL_ORI_MULAI if packing else KOL_ATO_MULAI

        blok, total, peringatan = pindai_tab(ws, kolom_qty_mulai=kolom_qty)
        if not blok:
            continue
        baris_judul_pertama = blok[0].baris_judul
        judul_cbd = _teks(ws.cell(baris_judul_pertama, KOL_CBD).value).replace("\n", " ")
        judul_cod = _teks(ws.cell(baris_judul_pertama, KOL_COD).value).replace("\n", " ")
        cust = daftar_customer.cari(nama)
        if cust is None:
            peringatan.append(
                f"Tab '{nama}' belum ada padanannya di config/customer.csv. "
                "Tambahkan barisnya supaya nama & alamat customer bisa dicetak."
            )
        order = Order(
            nama_tab=nama,
            customer_kunci=cust.kunci if cust else "",
            tanggal_po=_tanggal(nama, tahun_bawaan),
            blok=blok,
            total_sheet=total,
            peringatan=peringatan,
            judul_cbd=judul_cbd,
            judul_cod=judul_cod,
        )
        hasil.append(order)
    return hasil


def baca_master_harga(berkas: Path) -> dict[str, tuple[str, float]]:
    """Baca tab 'Harga Retail' jadi {kode: (nama_barang, harga)}."""
    wb = openpyxl.load_workbook(berkas, data_only=True)
    if "Harga Retail" not in wb.sheetnames:
        return {}
    ws = wb["Harga Retail"]
    master: dict[str, tuple[str, float]] = {}
    for r in range(1, ws.max_row + 1):
        kode = _teks(ws.cell(r, 1).value)
        if not kode or kode.upper() in ("COLUMN 1", "ARTICLE CODE"):
            continue
        master[kode] = (_teks(ws.cell(r, 2).value), angka(ws.cell(r, 3).value))
    return master
