"""Laporan pencocokan dan rekap bulanan.

Semua laporan berbentuk tabel dan selalu disertai ringkasan, sesuai permintaan
Yosua. Bahasa Indonesia, tanpa istilah teknis.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .dokumen import gaya
from .rekonsiliasi import HasilRekonsiliasi


def _rp(x: float) -> str:
    return f"Rp{x:,.0f}".replace(",", ".")


def _ribu(x: int) -> str:
    return f"{x:,}".replace(",", ".")


# ------------------------------------------------------------ ke layar
def cetak_ringkas(hasil: list[HasilRekonsiliasi]) -> None:
    print()
    print("RINGKASAN PENCOCOKAN ORDER SHEET")
    print("=" * 118)
    print(
        f"{'TAB / PO':<34}{'BLOK':>5}{'BARIS':>7}{'QTY':>8}"
        f"{'SEBELUM DISKON':>18}{'BAYAR':>7}{'NILAI BERSIH':>18}{'HASIL':>9}"
    )
    print("-" * 118)
    tq = 0
    tg = tn = 0.0
    tb = tblok = 0
    for h in hasil:
        o, k = h.order, h.keputusan
        tq += o.qty
        tg += o.nilai_kotor
        tn += k.nett_total
        tb += o.jumlah_baris
        tblok += len(o.blok)
        print(
            f"{o.nama_tab[:33]:<34}{len(o.blok):>5}{o.jumlah_baris:>7}{_ribu(o.qty):>8}"
            f"{_rp(o.nilai_kotor):>18}{k.cara_bayar:>7}{_rp(k.nett_total):>18}"
            f"{('COCOK' if h.lolos else 'GAGAL'):>9}"
        )
    print("-" * 118)
    print(
        f"{'TOTAL ' + str(len(hasil)) + ' PO':<34}{tblok:>5}{tb:>7}{_ribu(tq):>8}"
        f"{_rp(tg):>18}{'':>7}{_rp(tn):>18}"
    )
    print("=" * 118)


def cetak_rinci(h: HasilRekonsiliasi) -> None:
    o, k = h.order, h.keputusan
    print()
    print(f"--- {o.nama_tab} ---")
    print(f"  Cara bayar   : {k.cara_bayar}  (nilai bersih diambil dari {k.kolom_sumber})")
    print(f"  Alasan       : {k.alasan}")
    print(f"  Kolom CBD    : terisi {k.terisi_cbd} dari {k.jumlah_baris} baris")
    print(f"  Kolom COD    : terisi {k.terisi_cod} dari {k.jumlah_baris} baris")
    print(f"  Blok/kategori: {len(o.blok)}")
    for b in o.blok:
        label = [b.label_ukuran[i] for i in b.kolom_terpakai()]
        print(f"     blok {b.nomor} (baris judul {b.baris_judul}): "
              f"{len(b.baris)} baris, {b.qty} pcs, ukuran {label}")
    print()
    print(f"  {'PEMERIKSAAN':<42}{'HASIL HITUNG':>20}{'ANGKA ORDER SHEET':>22}{'':>8}")
    print("  " + "-" * 92)
    for p in h.periksa:
        tanda = "COCOK" if p.cocok else "BEDA"
        print(f"  {p.nama:<42}{p.nilai_hitung:>20}{p.nilai_sheet:>22}{tanda:>8}")
        if p.keterangan:
            print(f"     ({p.keterangan})")
    if h.peringatan:
        print()
        print("  PERLU DIPERHATIKAN:")
        for w in h.peringatan:
            print(f"   - {w}")


# ------------------------------------------------------------ ke Excel
def tulis_laporan_pencocokan(hasil: list[HasilRekonsiliasi], berkas: Path,
                             peringatan_umum: list[str]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"

    r = 1
    t = ws.cell(r, 1, "LAPORAN PENCOCOKAN ORDER SHEET")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    r += 2

    judul = ["TAB / PO", "Customer", "Blok", "Baris", "Qty (pcs)", "Sebelum diskon",
             "Cara bayar", "Kolom sumber nett", "Nilai bersih", "Kolom CBD terisi", "Hasil"]
    gaya.baris_judul_tabel(ws, r, judul)
    awal = r
    r += 1
    tq = 0; tg = tn = 0.0; tb = 0
    for h in hasil:
        o, k = h.order, h.keputusan
        ws.cell(r, 1, o.nama_tab)
        ws.cell(r, 2, o.customer_kunci or "(belum terdaftar)")
        ws.cell(r, 3, len(o.blok)).alignment = Alignment(horizontal="center")
        ws.cell(r, 4, o.jumlah_baris).alignment = Alignment(horizontal="center")
        c = ws.cell(r, 5, o.qty); c.number_format = gaya.ANGKA
        c = ws.cell(r, 6, o.nilai_kotor); c.number_format = gaya.RUPIAH
        ws.cell(r, 7, k.cara_bayar).alignment = Alignment(horizontal="center")
        ws.cell(r, 8, k.kolom_sumber)
        c = ws.cell(r, 9, k.nett_total); c.number_format = gaya.RUPIAH
        ws.cell(r, 10, f"{k.terisi_cbd} / {k.jumlah_baris}").alignment = Alignment(horizontal="center")
        c = ws.cell(r, 11, "COCOK" if h.lolos else "GAGAL")
        c.font = Font(name=gaya.FONT, size=9, bold=not h.lolos)
        c.alignment = Alignment(horizontal="center")
        for cc in range(1, 12):
            if ws.cell(r, cc).font.bold is not True:
                ws.cell(r, cc).font = Font(name=gaya.FONT, size=9)
        tq += o.qty; tg += o.nilai_kotor; tn += k.nett_total; tb += o.jumlah_baris
        r += 1
    sel = ws.cell(r, 1, f"TOTAL {len(hasil)} PO"); sel.font = Font(name=gaya.FONT, size=9, bold=True)
    ws.cell(r, 4, tb).font = Font(name=gaya.FONT, size=9, bold=True)
    c = ws.cell(r, 5, tq); c.number_format = gaya.ANGKA; c.font = Font(name=gaya.FONT, size=9, bold=True)
    c = ws.cell(r, 6, tg); c.number_format = gaya.RUPIAH; c.font = Font(name=gaya.FONT, size=9, bold=True)
    c = ws.cell(r, 9, tn); c.number_format = gaya.RUPIAH; c.font = Font(name=gaya.FONT, size=9, bold=True)
    gaya.beri_garis(ws, awal, 1, r, 11)
    gaya.atur_lebar(ws, {1: 34, 2: 20, 3: 6, 4: 7, 5: 10, 6: 18, 7: 10, 8: 20, 9: 18, 10: 14, 11: 9})

    # ---- lembar pemeriksaan rinci
    ws2 = wb.create_sheet("Pemeriksaan Rinci")
    r = 1
    t = ws2.cell(r, 1, "PEMERIKSAAN RINCI PER PO")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    r += 2
    for h in hasil:
        s = ws2.cell(r, 1, h.order.nama_tab)
        s.font = Font(name=gaya.FONT, size=11, bold=True)
        r += 1
        gaya.baris_judul_tabel(ws2, r, ["Pemeriksaan", "Hasil hitung", "Angka order sheet", "Hasil", "Keterangan"])
        awal = r
        r += 1
        for p in h.periksa:
            ws2.cell(r, 1, p.nama)
            ws2.cell(r, 2, p.nilai_hitung)
            ws2.cell(r, 3, p.nilai_sheet)
            c = ws2.cell(r, 4, "COCOK" if p.cocok else "BEDA")
            c.font = Font(name=gaya.FONT, size=9, bold=not p.cocok)
            ws2.cell(r, 5, p.keterangan)
            for cc in range(1, 6):
                if ws2.cell(r, cc).font.bold is not True:
                    ws2.cell(r, cc).font = Font(name=gaya.FONT, size=9)
            r += 1
        gaya.beri_garis(ws2, awal, 1, r - 1, 5)
        r += 1
    gaya.atur_lebar(ws2, {1: 42, 2: 22, 3: 24, 4: 9, 5: 70})

    # ---- lembar hal yang perlu diperhatikan
    ws3 = wb.create_sheet("Perlu Diperhatikan")
    r = 1
    t = ws3.cell(r, 1, "HAL YANG PERLU DIPERHATIKAN")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    r += 2
    gaya.baris_judul_tabel(ws3, r, ["PO / Bagian", "Catatan"])
    awal = r
    r += 1
    for w in peringatan_umum:
        ws3.cell(r, 1, "PENGATURAN UMUM").font = Font(name=gaya.FONT, size=9, bold=True)
        c = ws3.cell(r, 2, w); c.font = Font(name=gaya.FONT, size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[r].height = 30
        r += 1
    for h in hasil:
        for w in h.peringatan:
            ws3.cell(r, 1, h.order.nama_tab).font = Font(name=gaya.FONT, size=9)
            c = ws3.cell(r, 2, w); c.font = Font(name=gaya.FONT, size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws3.row_dimensions[r].height = 30
            r += 1
    if r == awal + 1:
        ws3.cell(r, 1, "-")
        ws3.cell(r, 2, "Tidak ada catatan.")
        r += 1
    gaya.beri_garis(ws3, awal, 1, r - 1, 2)
    gaya.atur_lebar(ws3, {1: 34, 2: 110})

    berkas.parent.mkdir(parents=True, exist_ok=True)
    wb.save(berkas)
    return berkas


def tulis_rekap_penjualan(hasil: list[HasilRekonsiliasi], berkas: Path, tarif_ppn: float) -> Path:
    """Rekap penjualan sebulan — bahan laporan keuangan dan pajak."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Penjualan"
    r = 1
    t = ws.cell(r, 1, "REKAP PENJUALAN PER PO")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    r += 1
    s = ws.cell(r, 1, f"DPP dan PPN dihitung mundur dari nilai bersih, tarif {tarif_ppn:.0%}.")
    s.font = Font(name=gaya.FONT, size=9, italic=True)
    r += 2

    judul = ["Tanggal PO", "PO / Tab", "Customer", "Qty (pcs)", "Sebelum diskon",
             "Diskon", "Nilai bersih", "DPP", f"PPN {tarif_ppn:.0%}", "Cara bayar"]
    gaya.baris_judul_tabel(ws, r, judul)
    awal = r
    r += 1
    tq = 0
    tg = tn = tdpp = tppn = 0.0
    for h in sorted(hasil, key=lambda x: (x.order.tanggal_po is None, x.order.tanggal_po)):
        o, k = h.order, h.keputusan
        nett = k.nett_total
        dpp = nett / (1 + tarif_ppn) if tarif_ppn else nett
        ppn = nett - dpp
        ws.cell(r, 1, gaya.tanggal_indonesia(o.tanggal_po))
        ws.cell(r, 2, o.nama_tab)
        ws.cell(r, 3, o.customer_kunci or "(belum terdaftar)")
        c = ws.cell(r, 4, o.qty); c.number_format = gaya.ANGKA
        c = ws.cell(r, 5, o.nilai_kotor); c.number_format = gaya.RUPIAH
        c = ws.cell(r, 6, o.nilai_kotor - nett); c.number_format = gaya.RUPIAH
        c = ws.cell(r, 7, nett); c.number_format = gaya.RUPIAH
        c = ws.cell(r, 8, dpp); c.number_format = gaya.RUPIAH
        c = ws.cell(r, 9, ppn); c.number_format = gaya.RUPIAH
        ws.cell(r, 10, k.cara_bayar).alignment = Alignment(horizontal="center")
        for cc in range(1, 11):
            ws.cell(r, cc).font = Font(name=gaya.FONT, size=9)
        tq += o.qty; tg += o.nilai_kotor; tn += nett; tdpp += dpp; tppn += ppn
        r += 1
    sel = ws.cell(r, 2, "TOTAL"); sel.font = Font(name=gaya.FONT, size=10, bold=True)
    for kolom, nilai, fmt in ((4, tq, gaya.ANGKA), (5, tg, gaya.RUPIAH), (6, tg - tn, gaya.RUPIAH),
                              (7, tn, gaya.RUPIAH), (8, tdpp, gaya.RUPIAH), (9, tppn, gaya.RUPIAH)):
        c = ws.cell(r, kolom, nilai); c.number_format = fmt
        c.font = Font(name=gaya.FONT, size=10, bold=True)
    gaya.beri_garis(ws, awal, 1, r, 10)
    gaya.atur_lebar(ws, {1: 16, 2: 34, 3: 20, 4: 10, 5: 18, 6: 16, 7: 18, 8: 18, 9: 16, 10: 11})

    berkas.parent.mkdir(parents=True, exist_ok=True)
    wb.save(berkas)
    return berkas
