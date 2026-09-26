"""Berkas Excel database customer: master, riwayat, dan daftar perlu diperiksa."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .db_customer import Customer, dugaan_nama_sama
from .dokumen import gaya

PERSEN = "0.0%"


def _judul_lembar(ws, teks: str, sub: str = "") -> int:
    t = ws.cell(1, 1, teks)
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    r = 2
    if sub:
        s = ws.cell(2, 1, sub)
        s.font = Font(name=gaya.FONT, size=9, italic=True)
        r = 3
    return r + 1


def _isi(ws, r, kolom, nilai, *, fmt=None, bold=False, tengah=False):
    c = ws.cell(r, kolom, nilai)
    c.font = Font(name=gaya.FONT, size=9, bold=bold)
    if fmt:
        c.number_format = fmt
    if tengah:
        c.alignment = Alignment(horizontal="center")
    return c


def tulis(daftar: list[Customer], berkas: Path, jejak_gabung: list[str],
          order_sheet_terbaca: list[str]) -> Path:
    wb = Workbook()

    # ------------------------------------------------ 1. MASTER_CUSTOMER
    ws = wb.active
    ws.title = "MASTER_CUSTOMER"
    r = _judul_lembar(
        ws, "MASTER CUSTOMER — dari seluruh order sheet 2025 & 2026",
        "Kolom berjudul (ISI MANUAL) dikosongkan karena tidak ada di order sheet. "
        "Kolom lainnya dihitung dari order sheet, jangan diketik ulang.",
    )
    judul = [
        "NAMA CUSTOMER", "JML PO", "PERIODE AKTIF", "PO TERAKHIR",
        "DISKON TERAKHIR", "DISKON PERNAH DIPAKAI", "DISKON BERUBAH?",
        "CARA BAYAR TERAKHIR", "CARA BAYAR UTAMA", "CARA BAYAR BERUBAH?",
        "PERNAH PAKAI KOLOM PPN", "TOTAL QTY", "TOTAL SEBELUM DISKON", "TOTAL NETT",
        "PERUSAHAAN PEMROSES (ISI MANUAL)", "TERMIN HARI (ISI MANUAL)",
        "NAMA DI DOKUMEN (ISI MANUAL)", "ALAMAT (ISI MANUAL)", "NPWP (ISI MANUAL)",
        "FORMAT INVOICE (ISI MANUAL)", "EJAAN LAIN DI ORDER SHEET",
    ]
    gaya.baris_judul_tabel(ws, r, judul)
    awal = r
    r += 1
    for c in daftar:
        _isi(ws, r, 1, c.nama_tampil, bold=True)
        _isi(ws, r, 2, c.jumlah_po, tengah=True)
        _isi(ws, r, 3, c.periode_aktif)
        _isi(ws, r, 4, c.terakhir.periode)
        _isi(ws, r, 5, c.diskon_terakhir, fmt=PERSEN, bold=True)
        _isi(ws, r, 6, ", ".join(f"{d:.1%}" for d in c.diskon_dipakai()))
        _isi(ws, r, 7, "YA" if c.diskon_berubah else "tetap", tengah=True)
        _isi(ws, r, 8, c.cara_bayar_terakhir, tengah=True, bold=True)
        _isi(ws, r, 9, c.cara_bayar_utama(), tengah=True)
        _isi(ws, r, 10, "YA" if c.cara_bayar_berubah else "tetap", tengah=True)
        _isi(ws, r, 11, "YA" if c.pernah_ppn else "", tengah=True)
        _isi(ws, r, 12, c.total_qty, fmt=gaya.ANGKA)
        _isi(ws, r, 13, c.total_kotor, fmt=gaya.RUPIAH)
        _isi(ws, r, 14, c.total_nett, fmt=gaya.RUPIAH)
        for k in (15, 16, 17, 18, 19, 20):
            _isi(ws, r, k, "")
        lain = sorted(x for x in c.ejaan if x != c.nama_tampil)
        _isi(ws, r, 21, ", ".join(lain))
        r += 1
    gaya.beri_garis(ws, awal, 1, r - 1, len(judul))
    gaya.atur_lebar(ws, {1: 32, 2: 8, 3: 24, 4: 16, 5: 15, 6: 26, 7: 14, 8: 16,
                         9: 16, 10: 16, 11: 18, 12: 11, 13: 20, 14: 20, 15: 28,
                         16: 18, 17: 26, 18: 30, 19: 18, 20: 22, 21: 44})
    ws.freeze_panes = ws.cell(awal + 1, 2)

    # ------------------------------------------------ 2. RIWAYAT_PO
    ws2 = wb.create_sheet("RIWAYAT_PO")
    r = _judul_lembar(
        ws2, "RIWAYAT SEMUA PO",
        "Satu baris satu PO. Dipakai untuk menelusuri dari mana angka di "
        "MASTER_CUSTOMER berasal.",
    )
    judul2 = ["TAHUN", "BULAN", "PERIODE", "ORDER SHEET", "NAMA TAB",
              "CUSTOMER (BAKU)", "BARIS", "QTY", "SEBELUM DISKON", "NETT",
              "DISKON NYATA", "DISKON DI KOLOM DISC", "CARA BAYAR",
              "KOLOM SUMBER NETT"]
    gaya.baris_judul_tabel(ws2, r, judul2)
    awal = r
    r += 1
    for c in daftar:
        for x in c.urut():
            _isi(ws2, r, 1, x.tahun, tengah=True)
            _isi(ws2, r, 2, x.bulan, tengah=True)
            _isi(ws2, r, 3, x.periode)
            _isi(ws2, r, 4, x.sumber)
            _isi(ws2, r, 5, x.tab)
            _isi(ws2, r, 6, c.nama_tampil)
            _isi(ws2, r, 7, x.baris, tengah=True)
            _isi(ws2, r, 8, x.qty, fmt=gaya.ANGKA)
            _isi(ws2, r, 9, x.kotor, fmt=gaya.RUPIAH)
            _isi(ws2, r, 10, x.nett, fmt=gaya.RUPIAH)
            _isi(ws2, r, 11, x.diskon_nyata, fmt=PERSEN)
            _isi(ws2, r, 12, ", ".join(f"{d:.1%}" for d in x.diskon_ditulis))
            _isi(ws2, r, 13, x.cara_bayar, tengah=True)
            _isi(ws2, r, 14, x.kolom_sumber)
            r += 1
    gaya.beri_garis(ws2, awal, 1, r - 1, len(judul2))
    gaya.atur_lebar(ws2, {1: 8, 2: 8, 3: 16, 4: 46, 5: 34, 6: 30, 7: 8, 8: 10,
                          9: 18, 10: 18, 11: 13, 12: 20, 13: 12, 14: 30})
    ws2.freeze_panes = ws2.cell(awal + 1, 1)

    # ------------------------------------------------ 3. PERIKSA_NAMA
    ws3 = wb.create_sheet("PERIKSA_NAMA")
    r = _judul_lembar(
        ws3, "NAMA YANG PERLU DIPERIKSA",
        "Nama-nama ini MIRIP tapi TIDAK digabung otomatis, karena bisa jadi "
        "toko yang sama atau cabang berbeda. Mohon diperiksa: kalau memang sama, "
        "tulis di config/alias_customer.csv.",
    )
    judul3 = ["GRUP", "NAMA DI ORDER SHEET", "JML PO", "PERIODE AKTIF",
              "DISKON TERAKHIR", "CARA BAYAR TERAKHIR", "TOTAL NETT", "SAMA? (ISI YA/TIDAK)"]
    gaya.baris_judul_tabel(ws3, r, judul3)
    awal = r
    r += 1
    grup = dugaan_nama_sama(daftar)
    for i, g in enumerate(grup, start=1):
        for c in g:
            _isi(ws3, r, 1, i, tengah=True)
            _isi(ws3, r, 2, c.nama_tampil)
            _isi(ws3, r, 3, c.jumlah_po, tengah=True)
            _isi(ws3, r, 4, c.periode_aktif)
            _isi(ws3, r, 5, c.diskon_terakhir, fmt=PERSEN)
            _isi(ws3, r, 6, c.cara_bayar_terakhir, tengah=True)
            _isi(ws3, r, 7, c.total_nett, fmt=gaya.RUPIAH)
            _isi(ws3, r, 8, "")
            r += 1
    if not grup:
        _isi(ws3, r, 2, "Tidak ada nama yang perlu diperiksa.")
        r += 1
    gaya.beri_garis(ws3, awal, 1, r - 1, len(judul3))
    gaya.atur_lebar(ws3, {1: 7, 2: 34, 3: 8, 4: 24, 5: 15, 6: 18, 7: 18, 8: 20})

    # ------------------------------------------------ 4. CARA BACA
    ws4 = wb.create_sheet("CARA BACA")
    baris = [
        ("APA ISI BERKAS INI", ""),
        ("", "Database customer yang disusun dari SELURUH order sheet 2025 dan 2026."),
        ("", ""),
        ("MASTER_CUSTOMER", "Satu baris satu customer. Diskon dan cara bayar diambil dari PO terakhirnya."),
        ("RIWAYAT_PO", "Satu baris satu PO. Untuk menelusuri asal angkanya."),
        ("PERIKSA_NAMA", "Nama mirip yang perlu dipastikan sama atau beda."),
        ("", ""),
        ("ARTI KOLOM CARA BAYAR", ""),
        ("TOP", "Nilai bersih diambil dari kolom TOTAL VALUE. Pembayaran tempo."),
        ("CBD", "Kolom DISCOUNT CBD terisi penuh. Bayar di muka, dapat potongan tambahan."),
        ("COD", "Kolom DISCOUNT COD terisi penuh. Bayar saat barang datang."),
        ("PPN", "Kolom DISCOUNT PPN + 11% terisi penuh. Muncul di order sheet 2025."),
        ("", "Kolom PPN ini kemungkinan penanda order yang diproses lewat perusahaan"),
        ("", "yang memungut PPN. Mohon dipastikan, lalu isi kolom PERUSAHAAN PEMROSES."),
        ("", ""),
        ("DISKON NYATA", "Dihitung mundur dari TOTAL VALUE dibagi nilai sebelum diskon."),
        ("", "Bukan diambil dari kolom DISC, karena kolom DISC beberapa kali salah isi."),
        ("", ""),
        ("CATATAN PENGGABUNGAN NAMA", ""),
    ]
    r = 1
    for a, b in baris:
        _isi(ws4, r, 1, a, bold=bool(a) and not b)
        _isi(ws4, r, 2, b)
        r += 1
    for j in jejak_gabung:
        _isi(ws4, r, 2, j)
        r += 1
    r += 1
    _isi(ws4, r, 1, "ORDER SHEET YANG TERBACA", bold=True)
    r += 1
    for s in order_sheet_terbaca:
        _isi(ws4, r, 2, s)
        r += 1
    gaya.atur_lebar(ws4, {1: 30, 2: 96})

    berkas.parent.mkdir(parents=True, exist_ok=True)
    wb.save(berkas)
    return berkas
