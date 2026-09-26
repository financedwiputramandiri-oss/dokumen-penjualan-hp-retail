"""Laporan hasil sapuan, untuk ditaruh ke Google Drive."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from ..dokumen import gaya
from . import rekap as rkp
from .pantau import GENTING, KABAR, PERHATIAN, Perubahan


def _isi(ws, r, c, v, *, bold=False, fmt=None, wrap=False):
    s = ws.cell(r, c, v)
    s.font = Font(name=gaya.FONT, size=9, bold=bold)
    if fmt:
        s.number_format = fmt
    if wrap:
        s.alignment = Alignment(wrap_text=True, vertical="top")
    return s


def _tabel(ws, r: int, judul: list[str], baris: list[list],
           lebar: dict[int, int]) -> int:
    """Satu tabel bergaris. Kolom yang judulnya menyebut rupiah diformat Rp."""
    uang = {c for c, j in enumerate(judul, start=1)
            if j.startswith(("SEBELUM DISKON", "NILAI BERSIH"))}
    gaya.baris_judul_tabel(ws, r, judul)
    awal = r
    r += 1
    for isi in baris or [["(tidak ada)"] + [""] * (len(judul) - 1)]:
        for c, v in enumerate(isi, start=1):
            _isi(ws, r, c, v,
                 fmt=gaya.FORMAT_RP if c in uang and isinstance(v, (int, float)) else None)
        r += 1
    gaya.beri_garis(ws, awal, 1, r - 1, len(judul))
    gaya.atur_lebar(ws, lebar)
    ws.freeze_panes = ws.cell(awal + 1, 1)
    return r


def _lembar_rekap(wb, rekaman: list[dict], waktu: datetime) -> None:
    """Dua lembar: ringkas per order sheet, lalu rinci per PO."""
    ws = wb.create_sheet("REKAP", 1)
    t = ws.cell(1, 1, "REKAP HASIL SAPUAN — PER ORDER SHEET")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    _isi(ws, 2, 1, f"Sapuan {waktu:%d %B %Y, %H:%M}")
    r = _tabel(ws, 4, rkp.JUDUL_PER_SHEET, rkp.baris_per_sheet(rekaman),
               {1: 46, 2: 8, 3: 16, 4: 14, 5: 22, 6: 22, 7: 60})
    r += 1
    _isi(ws, r, 1,
         "Qty dan nilai SENGAJA dikosongkan untuk order sheet yang dokumennya "
         "belum terbit. Angka itu justru yang belum cocok dengan baris TOTAL "
         "order sheet, jadi tidak boleh dipakai sebagai angka resmi.", wrap=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 30

    ws2 = wb.create_sheet("REKAP PER PO", 2)
    t = ws2.cell(1, 1, "REKAP HASIL SAPUAN — PER PO")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    _tabel(ws2, 3, rkp.JUDUL_PER_PO, rkp.baris_per_po(rekaman),
           {1: 42, 2: 34, 3: 22, 4: 13, 5: 10, 6: 20, 7: 12, 8: 20, 9: 14, 10: 52})


def tulis(
    berkas: Path,
    perubahan: list[Perubahan],
    diperiksa: list[tuple[str, str, int]],
    draf_dibuat: list[str],
    masalah: list[str],
    waktu: datetime | None = None,
    dilewati: list[str] | None = None,
    tab_ditarik: int = 0,
    rekaman: list[dict] | None = None,
) -> Path:
    dilewati = dilewati or []
    rekaman = rekaman or []
    waktu = waktu or datetime.now()
    wb = Workbook()

    # ---------------------------------------------------------- RINGKASAN
    ws = wb.active
    ws.title = "RINGKASAN"
    genting = [p for p in perubahan if p.tingkat == GENTING]
    perhatian = [p for p in perubahan if p.tingkat == PERHATIAN]

    t = ws.cell(1, 1, "LAPORAN SAPUAN ORDER SHEET")
    t.font = Font(name=gaya.FONT, size=15, bold=True)
    _isi(ws, 2, 1, f"Waktu sapuan: {waktu:%d %B %Y, %H:%M}")

    r = 4
    if genting:
        p = ws.cell(r, 1, f"ADA {len(genting)} PERUBAHAN PENTING YANG PERLU DIPERIKSA")
        p.font = Font(name=gaya.FONT, size=13, bold=True)
        r += 1
        _isi(ws, r, 1,
             "Perubahan di bawah terjadi pada PO yang ATO-nya SUDAH TERISI. "
             "PO seperti itu seharusnya sudah beku karena barangnya sudah dikirim "
             "dan dokumennya sudah terbit.", wrap=True)
        ws.row_dimensions[r].height = 30
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 2
    else:
        p = ws.cell(r, 1, "TIDAK ADA PERUBAHAN PENTING")
        p.font = Font(name=gaya.FONT, size=13, bold=True)
        r += 1
        _isi(ws, r, 1, "Tidak ada PO lama yang angkanya atau rumusnya berubah.")
        r += 2

    ringkas = [
        ("Order sheet dibaca", len(diperiksa)),
        ("Order sheet dilewati (tidak berubah)", len(dilewati)),
        ("Tab PO diperiksa", sum(n for _, _, n in diperiksa)),
        ("Tab yang benar-benar ditarik isinya", tab_ditarik),
        ("Perubahan penting (GENTING)", len(genting)),
        ("Perubahan wajar (PERHATIAN)", len(perhatian)),
        ("Draf dokumen dibuat", len(draf_dibuat)),
        ("Masalah teknis", len(masalah)),
    ]
    gaya.baris_judul_tabel(ws, r, ["HAL", "JUMLAH"])
    awal = r
    r += 1
    for a, b in ringkas:
        _isi(ws, r, 1, a)
        _isi(ws, r, 2, b, bold=True)
        r += 1
    gaya.beri_garis(ws, awal, 1, r - 1, 2)
    gaya.atur_lebar(ws, {1: 44, 2: 14, 3: 20, 4: 20, 5: 20, 6: 20})

    # ------------------------------------------------------- PERUBAHAN
    ws2 = wb.create_sheet("PERUBAHAN")
    t = ws2.cell(1, 1, "SEMUA PERUBAHAN YANG TERDETEKSI")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    r = 3
    gaya.baris_judul_tabel(
        ws2, r, ["TINGKAT", "ORDER SHEET", "TAB PO", "JENIS PERUBAHAN",
                 "KETERANGAN", "SEBELUMNYA", "SEKARANG"]
    )
    awal = r
    r += 1
    urut = {GENTING: 0, PERHATIAN: 1, KABAR: 2}
    for p in sorted(perubahan, key=lambda x: (urut.get(x.tingkat, 9), x.nama_sheet, x.tab)):
        _isi(ws2, r, 1, p.tingkat, bold=p.genting)
        _isi(ws2, r, 2, p.nama_sheet)
        _isi(ws2, r, 3, p.tab)
        _isi(ws2, r, 4, p.jenis, bold=p.genting)
        _isi(ws2, r, 5, p.keterangan, wrap=True)
        _isi(ws2, r, 6, p.sebelum)
        _isi(ws2, r, 7, p.sesudah)
        ws2.row_dimensions[r].height = 28
        r += 1
    if not perubahan:
        _isi(ws2, r, 1, "-")
        _isi(ws2, r, 5, "Tidak ada perubahan.")
        r += 1
    gaya.beri_garis(ws2, awal, 1, r - 1, 7)
    gaya.atur_lebar(ws2, {1: 12, 2: 46, 3: 34, 4: 26, 5: 62, 6: 20, 7: 20})
    ws2.freeze_panes = ws2.cell(awal + 1, 1)

    # ------------------------------------------------------- YANG DIPERIKSA
    ws3 = wb.create_sheet("YANG DIPERIKSA")
    t = ws3.cell(1, 1, "ORDER SHEET YANG DISAPU")
    t.font = Font(name=gaya.FONT, size=14, bold=True)
    r = 3
    gaya.baris_judul_tabel(ws3, r, ["ORDER SHEET", "ID BERKAS", "JUMLAH TAB PO"])
    awal = r
    r += 1
    for nama, idb, n in diperiksa:
        _isi(ws3, r, 1, nama)
        _isi(ws3, r, 2, idb)
        _isi(ws3, r, 3, n)
        r += 1
    gaya.beri_garis(ws3, awal, 1, r - 1, 3)

    r += 2
    _isi(ws3, r, 1, "DRAF DOKUMEN YANG DIBUAT", bold=True)
    r += 1
    for d in draf_dibuat or ["(tidak ada)"]:
        _isi(ws3, r, 1, d)
        r += 1
    r += 1
    _isi(ws3, r, 1, "ORDER SHEET YANG DILEWATI (tidak berubah sejak sapuan lalu)", bold=True)
    r += 1
    for d in dilewati or ["(tidak ada)"]:
        _isi(ws3, r, 1, d)
        r += 1
    r += 1
    _isi(ws3, r, 1, "MASALAH TEKNIS", bold=True)
    r += 1
    for m in masalah or ["(tidak ada)"]:
        _isi(ws3, r, 1, m, wrap=True)
        r += 1
    gaya.atur_lebar(ws3, {1: 70, 2: 46, 3: 16})

    # ------------------------------------------------------------- REKAP
    # Ditaruh di depan dengan sengaja: inilah yang dicari orang kalau ia
    # membuka laporan sapuan di Drive — berapa dokumen yang jadi, dari PO
    # mana, dan yang belum jadi itu kenapa.
    _lembar_rekap(wb, rekaman, waktu)

    berkas.parent.mkdir(parents=True, exist_ok=True)
    wb.save(berkas)
    return berkas
