"""Tata letak CV. MUTIARA TIMUR NUSANTARA.

Acuannya berkas asli `FA 0010526 BABY FAME (MTN).xlsx` yang dikirim Yosua
19 September 2026 dan dibongkar sel per sel. Yang dikunci di sini adalah
hal-hal yang kalau salah membuat dokumennya keliru DIAM-DIAM: PPN yang
seharusnya tidak ada, tanggal yang terpotong, dan kolom diskon yang kosong.
"""
from pathlib import Path

import pytest
from openpyxl import Workbook

from buat_contoh import buat_contoh
from hp_dokumen.dokumen.mtn import (
    buat_invoice_mtn, buat_surat_jalan_mtn, nomor_mtn, LEBAR_INV,
)
from hp_dokumen.konfigurasi import Konfigurasi, Customer, DaftarCustomer
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet

AKAR = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def cfg():
    return Konfigurasi.muat()


@pytest.fixture(scope="module")
def daftar():
    return DaftarCustomer([
        Customer("Contoh TOP", "PT Contoh Satu", "Jl. Contoh 1", "",
                 "per_artikel", 30, "", "MTN", ""),
    ])


@pytest.fixture(scope="module")
def order(tmp_path_factory, daftar):
    berkas = buat_contoh(tmp_path_factory.mktemp("data") / "contoh.xlsx")
    orders = baca_order_sheet(berkas, daftar, tahun_bawaan=2026)
    return next(o for o in orders if "TOP" in o.nama_tab)


@pytest.fixture(scope="module")
def mtn(cfg):
    return next(p for p in cfg.perusahaan.semua if p.kode.upper() == "MTN")


def _invoice(order, cfg, daftar, mtn):
    ws = Workbook().active
    cust = daftar.cari(order.nama_tab)
    hasil = buat_invoice_mtn(ws, order, tentukan_nett(order, cust), cust, mtn,
                             cfg.pengaturan, "001")
    return ws, hasil


def test_invoice_mtn_tanpa_dpp_dan_ppn(order, cfg, daftar, mtn):
    """MTN tidak mengenakan PPN - penutupnya hanya Subtotal, Value Disc, Total.

    Kalau suatu saat ada yang "melengkapi" dengan baris DPP/PPN, faktur MTN
    jadi menagih pajak yang tidak seharusnya. Itu kesalahan yang mahal dan
    tidak kelihatan dari angka totalnya.
    """
    ws, _ = _invoice(order, cfg, daftar, mtn)
    teks = {str(ws.cell(r, 7).value or "").upper()
            for r in range(1, ws.max_row + 1)}
    assert "SUBTOTAL" in teks
    assert any(t.startswith("VALUE DISC") for t in teks)
    assert "TOTAL" in teks
    assert not any("DPP" in t for t in teks), "MTN tidak boleh punya baris DPP"
    assert not any("PPN" in t for t in teks), "MTN tidak boleh punya baris PPN"


def test_invoice_mtn_kolom_diskon_tidak_kosong(order, cfg, daftar, mtn):
    """Diskon tunggal harus tetap tercetak, contoh `25%`.

    `_persen_tertulis()` mengembalikan None pada teksnya untuk diskon tunggal
    dan angkanya terpisah. Versi pertama memakai teks itu apa adanya, jadi
    kolom Diskon tercetak KOSONG - cacat yang sama pernah terjadi di proforma.
    """
    ws, _ = _invoice(order, cfg, daftar, mtn)
    # baris 18 = sublabel judul tabel, kolom 6 = Diskon
    assert str(ws.cell(18, 6).value or "").strip(), "kolom Diskon kosong"
    assert "%" in str(ws.cell(18, 6).value)


def test_invoice_mtn_tanggal_muat_di_kolomnya(order, cfg, daftar, mtn):
    """Kolom TANGGAL harus cukup lebar untuk bulan berhuruf panjang.

    "17 September 2026" ada 17 huruf. Pada lebar asli 15,71 tercetak terpotong
    jadi "17 Septe" - dan itu baru kelihatan setelah dirender jadi PDF.
    """
    ws, _ = _invoice(order, cfg, daftar, mtn)
    tanggal = str(ws.cell(12, 8).value or "")
    assert tanggal, "tanggal tidak ditulis"
    assert LEBAR_INV[8] >= len("17 September 2026"), (
        f"kolom H cuma {LEBAR_INV[8]}, tanggal terpanjang butuh 17"
    )


def test_invoice_mtn_memakai_identitas_mtn(order, cfg, daftar, mtn):
    """Kop, rekening, dan logo harus milik MTN - bukan DPM."""
    ws, _ = _invoice(order, cfg, daftar, mtn)
    isi = " ".join(str(ws.cell(r, c).value or "")
                   for r in range(1, 40) for c in (1, 3))
    assert "MUTIARA TIMUR NUSANTARA" in isi.upper()
    assert "DWI PUTRA MANDIRI" not in isi.upper(), "identitas DPM bocor ke MTN"


def test_surat_jalan_mtn_susunan_kolomnya_sendiri(order, cfg, daftar, mtn):
    """Deskripsi di kolom C tunggal, WARNA di D - beda dari DPM (C:E dan F)."""
    ws = Workbook().active
    cust = daftar.cari(order.nama_tab)
    buat_surat_jalan_mtn(ws, order, cust, mtn, "001")
    assert str(ws.cell(19, 3).value).upper().startswith("DESKRIPSI")
    assert str(ws.cell(19, 4).value).upper() == "WARNA"
    # kolom C tidak boleh digabung ke D/E seperti Surat Jalan DPM
    digabung = {str(m) for m in ws.merged_cells.ranges}
    assert "C19:E20" not in digabung


def test_nomor_mtn_berbentuk_FA_dan_SJ():
    from datetime import date
    assert nomor_mtn("FA", "001", date(2026, 5, 5)) == "FA-001/05/2026"
    assert nomor_mtn("SJ", "001", date(2026, 5, 5)) == "SJ-001/05/2026"
    # nomor yang belum diisi tetap jadi penanda kosong, bukan ditebak
    assert "_" in nomor_mtn("FA", "________", date(2026, 5, 5))


# ---------------------------------------------- revisi Yosua 20 Sep 2026
def _mtn_inv(order, cfg, daftar, mtn):
    ws = Workbook().active
    cust = daftar.cari(order.nama_tab)
    buat_invoice_mtn(ws, order, tentukan_nett(order, cust), cust, mtn,
                     cfg.pengaturan, "001")
    return ws


def test_label_customer_digabung_supaya_tidak_terpotong(order, cfg, daftar, mtn):
    """Kolom A cuma 3,43 satuan.

    Tanpa digabung A:C, tulisan "CUSTOMER" yang dirata-tengahkan terpotong
    garis kotaknya dan tercetak "STOMER". Tidak kelihatan sama sekali dari
    nilai selnya - hanya muncul di gambar hasil render.
    """
    ws = _mtn_inv(order, cfg, daftar, mtn)
    assert ws.cell(11, 1).value == "CUSTOMER"
    assert "A11:C11" in {str(m) for m in ws.merged_cells.ranges}


def test_blok_customer_satu_kotak_tanpa_garis_antar_baris(order, cfg, daftar, mtn):
    """Revisi Yosua: kotak luar + satu garis di bawah label CUSTOMER saja.

    Berkas asli MTN memberi garis atas-bawah di TIAP baris, tapi Yosua
    memilih bentuk yang lebih bersih ini. Jangan dikembalikan ke kisi-kisi
    hanya karena berkas aslinya begitu.
    """
    ws = _mtn_inv(order, cfg, daftar, mtn)
    assert ws.cell(11, 1).border.bottom.style, "garis di bawah label hilang"
    for r in (13, 14, 15):
        assert not ws.cell(r, 1).border.top.style, (
            f"A{r} masih bergaris atas - blok customer jadi kisi-kisi"
        )
    assert ws.cell(15, 1).border.bottom.style, "kotak tidak ditutup di bawah"


def test_rekening_mulai_satu_baris_di_bawah_subtotal_dan_berkotak(
        order, cfg, daftar, mtn):
    """Sejajar "Value Disc", dengan kotak medium - seperti faktur DPM."""
    ws = _mtn_inv(order, cfg, daftar, mtn)
    baris_sub = next(r for r in range(1, ws.max_row + 1)
                     if str(ws.cell(r, 7).value or "").strip() == "Subtotal")
    assert ws.cell(baris_sub, 1).value is None, "rekening tidak boleh sejajar Subtotal"
    assert "PEMBAYARAN" in str(ws.cell(baris_sub + 1, 1).value or "").upper()
    assert ws.cell(baris_sub + 1, 1).border.top.style == "medium"


def test_hormat_kami_sejajar_baris_rekening_terakhir(order, cfg, daftar, mtn):
    ws = _mtn_inv(order, cfg, daftar, mtn)
    baris_hormat = next(r for r in range(1, ws.max_row + 1)
                        if str(ws.cell(r, 7).value or "") == "Hormat kami,")
    assert "A/C NO" in str(ws.cell(baris_hormat, 1).value or "").upper()


def test_kotak_nomor_surat_jalan_tipis_bukan_tebal(order, cfg, daftar, mtn):
    """Satu-satunya garis TEBAL di dokumen MTN adalah kotak blok rekening."""
    ws = Workbook().active
    cust = daftar.cari(order.nama_tab)
    buat_surat_jalan_mtn(ws, order, cust, mtn, "001")
    kolom_label = next(c for c in range(1, ws.max_column + 1)
                       if str(ws.cell(11, c).value or "") == "SURAT JALAN")
    assert ws.cell(11, kolom_label).border.top.style == "thin"


def test_surat_jalan_mtn_berjudul_di_tengah(order, cfg, daftar, mtn):
    """Judul "SURAT JALAN" di baris 9, ditengahkan selebar tabel.

    TIDAK ada di berkas MTN asli - ini tambahan Yosua 20 September 2026.
    Ukurannya disamakan dengan judul Surat Jalan DPM supaya kedua perusahaan
    memakai ukuran yang sama.
    """
    from hp_dokumen.dokumen.mtn import BARIS_JUDUL_SJ_DOK, HURUF_JUDUL_SJ_DOK
    from hp_dokumen.dokumen.surat_jalan import HURUF_JUDUL_DOK

    ws = Workbook().active
    cust = daftar.cari(order.nama_tab)
    buat_surat_jalan_mtn(ws, order, cust, mtn, "001")
    sel = ws.cell(BARIS_JUDUL_SJ_DOK, 1)
    assert sel.value == "SURAT JALAN"
    assert sel.alignment.horizontal == "center"
    assert sel.font.bold
    assert HURUF_JUDUL_SJ_DOK == HURUF_JUDUL_DOK, "beda ukuran dengan judul DPM"
    lebar = [m for m in ws.merged_cells.ranges
             if m.min_row == BARIS_JUDUL_SJ_DOK and m.min_col == 1]
    assert lebar, "judul tidak digabung, jadi tidak benar-benar di tengah"


def test_invoice_mtn_TIDAK_berjudul(order, cfg, daftar, mtn):
    """Nomornya sudah ada di label FAKTUR; berkas suntingan Yosua pun kosong."""
    ws = _mtn_inv(order, cfg, daftar, mtn)
    from hp_dokumen.dokumen.mtn import BARIS_JUDUL_SJ_DOK

    for c in range(1, 9):
        assert ws.cell(BARIS_JUDUL_SJ_DOK, c).value in (None, ""), (
            "faktur MTN tidak boleh punya judul dokumen"
        )


def test_label_penutup_mtn_muat_di_kolomnya():
    """`Value Disc 25% + 2%` pernah tercetak jadi `Value Disc 25%`.

    Kolom G bawaan 14,29 satuan hanya memuat tulisan pendeknya, jadi diskon
    gabungan CBD/COD terbaca lebih kecil daripada yang benar-benar ditagih.
    Cacat semacam ini tidak kelihatan dari nilai selnya.
    """
    from hp_dokumen.dokumen.mtn import (AWALAN_LABEL_DISKON, HURUF_PER_SATUAN,
                                        LEBAR_INV, _lebar_inv)

    for tulisan in ("25%", "22% + 1.5%", "25% + 2%"):
        lebar = _lebar_inv(tulisan)
        perlu = len(AWALAN_LABEL_DISKON + tulisan) * HURUF_PER_SATUAN
        assert lebar[7] >= perlu, f"label '{tulisan}' tidak muat di kolom G"
        assert lebar[3] >= 22.0, "kolom deskripsi tidak boleh menyusut habis"
        assert round(sum(lebar.values()), 2) == round(sum(LEBAR_INV.values()), 2)


def test_baris_surat_jalan_mtn_meninggi_kalau_deskripsi_melipat():
    """Tinggi asli 20,1 cuma memuat SATU baris; deskripsi panjang terpotong."""
    from hp_dokumen.dokumen.mtn import (LEBAR_SJ_TETAP, TINGGI_DATA_SJ,
                                        KOL_DESK_SJ, _tinggi_baris_sj)

    lebar = LEBAR_SJ_TETAP[KOL_DESK_SJ]
    assert _tinggi_baris_sj("Ziggy Set", lebar) == TINGGI_DATA_SJ
    panjang = "UltraCool Ruffle Sleeve Tee Medium Size"
    assert _tinggi_baris_sj(panjang, lebar) > TINGGI_DATA_SJ
