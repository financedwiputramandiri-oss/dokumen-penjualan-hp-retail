"""Tes bahwa bot membaca tab lewat Sheets API, bukan mengunduh spreadsheet.

Google tidak dipanggil sungguhan. Sambungan ditiru, lalu dicatat panggilan
apa saja yang dilakukan bot.
"""
from pathlib import Path

import openpyxl
import pytest

from buat_contoh import buat_contoh
from hp_dokumen.konfigurasi import Customer, DaftarCustomer
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_buku, baca_order_sheet, master_harga_dari_buku, pindai_tab
from hp_dokumen.sapu.lembar_api import BukuNilai, LembarNilai


def bentuk_api(ws):
    """Ubah worksheet openpyxl jadi bentuk yang dikirim Sheets API."""
    baris = []
    for r in range(1, ws.max_row + 1):
        row = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            row.append("" if v is None else v)
        while row and row[-1] == "":          # API memotong sel kosong di kanan
            row.pop()
        baris.append(row)
    while baris and not baris[-1]:
        baris.pop()
    return baris


def buku_api(berkas: Path) -> BukuNilai:
    wb = openpyxl.load_workbook(berkas, data_only=True)
    return BukuNilai({ws.title: bentuk_api(ws) for ws in wb.worksheets})


@pytest.fixture(scope="module")
def contoh(tmp_path_factory):
    return buat_contoh(tmp_path_factory.mktemp("d") / "contoh.xlsx")


@pytest.fixture(scope="module")
def daftar():
    return DaftarCustomer([
        Customer("Contoh TOP", "PT A", "x", "", "per_artikel", 30, "", "", ""),
        Customer("Contoh CBD", "PT B", "y", "", "per_ukuran", 30, "", "", ""),
    ])


# ------------------------------------------------ pembungkus Sheets API
def test_sel_kosong_beda_dari_angka_nol():
    """Aturan 2 bergantung pada beda 'sel kosong' dan 'terisi nol'."""
    ws = LembarNilai("t", [["a", "", 0, None]])
    assert ws.cell(1, 1).value == "a"
    assert ws.cell(1, 2).value is None     # teks kosong dianggap sel kosong
    assert ws.cell(1, 3).value == 0        # angka nol TETAP nol
    assert ws.cell(1, 4).value is None


def test_baris_pendek_dan_di_luar_jangkauan():
    """Sheets API memotong sel kosong di ujung kanan tiap baris."""
    ws = LembarNilai("t", [["a", "b", "c"], ["d"]])
    assert ws.max_row == 2 and ws.max_column == 3
    assert ws.cell(2, 1).value == "d"
    assert ws.cell(2, 3).value is None
    assert ws.cell(99, 99).value is None


def test_tab_kosong_tidak_bikin_error():
    ws = LembarNilai("kosong", [])
    assert ws.max_row == 0 and ws.max_column == 0
    assert ws.cell(1, 1).value is None


# ------------------------------- hasil lewat API = hasil lewat Excel
def test_hasil_api_identik_dengan_excel(contoh, daftar):
    dari_excel = baca_order_sheet(contoh, daftar, tahun_bawaan=2026)
    dari_api = baca_buku(buku_api(contoh), daftar, tahun_bawaan=2026)

    assert len(dari_api) == len(dari_excel)
    for a, b in zip(dari_excel, dari_api):
        assert a.nama_tab == b.nama_tab
        assert a.jumlah_baris == b.jumlah_baris
        assert a.qty == b.qty
        assert abs(a.nilai_kotor - b.nilai_kotor) < 0.01
        assert [x.label_ukuran for x in a.blok] == [x.label_ukuran for x in b.blok]
        ka = tentukan_nett(a, daftar.cari(a.nama_tab))
        kb = tentukan_nett(b, daftar.cari(b.nama_tab))
        assert ka.cara_bayar == kb.cara_bayar
        assert ka.kolom == kb.kolom
        assert abs(ka.nett_total - kb.nett_total) < 0.01


def test_master_harga_terbaca_lewat_api(contoh):
    from hp_dokumen.pemindai import baca_master_harga

    assert master_harga_dari_buku(buku_api(contoh)) == baca_master_harga(contoh)


def test_blok_bertumpuk_tetap_terbaca_lewat_api(contoh, daftar):
    """Aturan 1 harus tetap berlaku pada data dari API."""
    top = next(o for o in baca_buku(buku_api(contoh), daftar) if "TOP" in o.nama_tab)
    assert len(top.blok) == 2
    assert top.blok[0].label_ukuran[:3] == ["0-3M", "3-6M", "6-12M"]
    assert top.blok[1].label_ukuran[:3] == ["1", "2", "3"]


# ------------------------------------------- sambungan tiruan untuk bot
class SambunganTiruan:
    """Mencatat panggilan yang dilakukan bot, tanpa menyentuh Google."""

    def __init__(self, berkas: Path, id_sheet="SHEET1", diubah="2026-09-11T05:00:00Z"):
        self.email_bot = "bot@contoh.iam.gserviceaccount.com"
        self._buku = buku_api(berkas)
        self._id = id_sheet
        self._diubah = diubah
        self.panggilan: list[str] = []
        self.tab_ditarik: list[str] = []

    # -- Drive
    def isi_folder(self, id_folder):
        from hp_dokumen.sapu.google import BerkasDrive

        self.panggilan.append("isi_folder")
        return [BerkasDrive(id=self._id, nama="Order Sheet Contoh",
                            mime="application/vnd.google-apps.spreadsheet",
                            diubah=self._diubah)]

    def unduh_sebagai_xlsx(self, *a, **k):      # tidak boleh dipanggil lagi
        self.panggilan.append("unduh_sebagai_xlsx")
        raise AssertionError("bot tidak boleh mengunduh spreadsheet")

    def buat_folder_kalau_belum_ada(self, *a, **k):
        return None

    def unggah_laporan(self, *a, **k):
        return None

    # -- Sheets
    def daftar_tab(self, id_sheet):
        self.panggilan.append("daftar_tab")
        return [{"judul": j, "id": i, "baris": 0, "kolom": 0}
                for i, j in enumerate(self._buku.sheetnames)]

    def nama_tab(self, id_sheet):
        return self._buku.sheetnames

    def ambil_tab(self, id_sheet, judul_tab, rumus=False, per_permintaan=60):
        self.panggilan.append("ambil_tab_rumus" if rumus else "ambil_tab")
        self.tab_ditarik.extend(judul_tab)
        return {j: [] for j in judul_tab} if rumus else {
            j: bentuk_api(self._buku[j]) for j in judul_tab if j in self._buku
        }

    def buku_dari_tab(self, id_sheet, judul_tab, rumus=False):
        return BukuNilai(self.ambil_tab(id_sheet, judul_tab, rumus=rumus))


@pytest.fixture
def pengaturan_bot(tmp_path):
    from hp_dokumen.sapu.bot import Pengaturan

    return Pengaturan(
        folder=[{"nama": "Uji", "id": "FOLDER1"}],
        folder_laporan_id="", nama_folder_laporan="X",
        sheet_otomatisasi_id="",
        buat_draf=True, pantau_perubahan=True, pantau_rumus=False,
        hanya_hari=0,
        berkas_kredensial=tmp_path / "k.json",
        berkas_kondisi=tmp_path / "kondisi.json",
        lewati_yang_tidak_berubah=True,
    )


def _jalankan(monkeypatch, pengaturan, sambungan, cfg):
    import hp_dokumen.sapu.bot as botmod

    monkeypatch.setattr(botmod, "Sambungan", lambda _: sambungan)
    return botmod.sapu(pengaturan, cfg, cetak=lambda *a, **k: None)


@pytest.fixture
def konfigurasi(daftar):
    from hp_dokumen.konfigurasi import Konfigurasi, Pengaturan as PSet, DaftarPerusahaan
    from hp_dokumen.konfigurasi import Perusahaan

    pt = Perusahaan("DPM", "CV A", "CV. A", True, "1", ["Jl. A"], "", ["R"], "Jakarta")
    return Konfigurasi(
        perusahaan=DaftarPerusahaan([pt], "DPM", "Jakarta"),
        customer=daftar, pengaturan=PSet(),
    )


def test_bot_membaca_tab_tanpa_mengunduh(monkeypatch, pengaturan_bot, contoh, konfigurasi, tmp_path):
    s = SambunganTiruan(contoh)
    hasil = _jalankan(monkeypatch, pengaturan_bot, s, konfigurasi)

    assert "unduh_sebagai_xlsx" not in s.panggilan, "bot masih mengunduh spreadsheet"
    assert "daftar_tab" in s.panggilan
    assert "ambil_tab" in s.panggilan
    assert len(hasil.diperiksa) == 1
    assert hasil.tab_ditarik > 0


def test_bot_hanya_menarik_tab_po_dan_harga(monkeypatch, pengaturan_bot, contoh, konfigurasi):
    """Tab lain di spreadsheet tidak ikut ditarik."""
    s = SambunganTiruan(contoh)
    _jalankan(monkeypatch, pengaturan_bot, s, konfigurasi)
    for judul in s.tab_ditarik:
        atas = judul.upper().strip()
        assert atas.startswith(("PO ", "(DELIVERY", "PACKING LIST")) or judul == "Harga Retail"


def test_sapuan_kedua_melewati_sheet_yang_tidak_berubah(
    monkeypatch, pengaturan_bot, contoh, konfigurasi
):
    """Inti penghematan: waktu ubah sama -> tidak ditarik sama sekali."""
    s1 = SambunganTiruan(contoh)
    h1 = _jalankan(monkeypatch, pengaturan_bot, s1, konfigurasi)
    assert h1.dilewati == []
    assert h1.tab_ditarik > 0

    s2 = SambunganTiruan(contoh)           # waktu ubah SAMA
    h2 = _jalankan(monkeypatch, pengaturan_bot, s2, konfigurasi)
    assert h2.dilewati == ["Order Sheet Contoh"]
    assert h2.tab_ditarik == 0, "tidak boleh ada tab yang ditarik"
    assert s2.tab_ditarik == []
    assert "ambil_tab" not in s2.panggilan
    assert "daftar_tab" not in s2.panggilan


def test_sheet_yang_berubah_tetap_dibaca_ulang(
    monkeypatch, pengaturan_bot, contoh, konfigurasi
):
    s1 = SambunganTiruan(contoh)
    _jalankan(monkeypatch, pengaturan_bot, s1, konfigurasi)

    s2 = SambunganTiruan(contoh, diubah="2026-09-12T09:00:00Z")   # waktu ubah BARU
    h2 = _jalankan(monkeypatch, pengaturan_bot, s2, konfigurasi)
    assert h2.dilewati == []
    assert h2.tab_ditarik > 0
    assert "ambil_tab" in s2.panggilan


def test_bisa_dipaksa_membaca_semuanya(monkeypatch, pengaturan_bot, contoh, konfigurasi):
    s1 = SambunganTiruan(contoh)
    _jalankan(monkeypatch, pengaturan_bot, s1, konfigurasi)

    pengaturan_bot.lewati_yang_tidak_berubah = False
    s2 = SambunganTiruan(contoh)
    h2 = _jalankan(monkeypatch, pengaturan_bot, s2, konfigurasi)
    assert h2.dilewati == []
    assert h2.tab_ditarik > 0
