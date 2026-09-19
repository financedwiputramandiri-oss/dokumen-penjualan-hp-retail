"""Tab 'Harga Retail' adalah satu-satunya sumber harga dan nama barang.

Permintaan Yosua 14 September 2026: "untuk bagian harga sesuaikan dengan tab
master harga yang ada di setiap tab yang ada di order sheet".

Diperiksa lebih dulu pada seluruh 8.845 baris order sheet 2026: harga di baris
PO memang SUDAH sama persis dengan master di semua baris. Jadi aturan ini tidak
mengubah angka mana pun hari ini — gunanya menangkap kalau suatu saat Sales
mengetik harga sendiri di baris PO.
"""
import openpyxl
import pytest

from hp_dokumen.konfigurasi import Konfigurasi
from hp_dokumen.pemindai import baca_master_harga, baca_order_sheet, cari_di_master
from buat_contoh import buat_contoh


@pytest.fixture
def berkas(tmp_path):
    return buat_contoh(tmp_path / "contoh.xlsx")


def test_kode_dicari_tanpa_peduli_huruf_besar_kecil():
    """`41065 (bottom/Celana)` harus ketemu walau master menulis `Bottom`.

    Tiga kode nyata di order sheet Januari dan Februari 2026 dulu dianggap
    tidak terdaftar hanya karena beda huruf besar-kecil.
    """
    master = {"41065 (Bottom/Celana)": ("Nilo Straight Denim Pants", 44200.0)}
    assert cari_di_master(master, "41065 (bottom/Celana)") is not None
    assert cari_di_master(master, "41065 (BOTTOM/CELANA)") is not None
    assert cari_di_master(master, " 41065 (Bottom/Celana) ") is not None
    assert cari_di_master(master, "41065 (Bottom/Celana)")[1] == 44200.0
    assert cari_di_master(master, "kode lain") is None
    assert cari_di_master({}, "apa pun") is None


def test_harga_diambil_dari_master_bukan_dari_baris_po(berkas, tmp_path):
    """Harga yang diketik beda di baris PO harus ditimpa oleh master."""
    cfg = Konfigurasi.muat()
    master = baca_master_harga(berkas)
    assert master, "berkas contoh harus punya tab Harga Retail"

    # rusak harga satu baris PO supaya berbeda dari master
    wb = openpyxl.load_workbook(berkas)
    ws = next(w for w in wb.worksheets if w.title.lower() not in ("harga retail",))
    kolom_harga = next(
        c for c in range(1, ws.max_column + 1)
        for r in range(1, 12)
        if "PRICE" in str(ws.cell(r, c).value or "").upper()
    )
    baris_data = next(
        r for r in range(1, ws.max_row + 1)
        if str(ws.cell(r, kolom_harga).value or "").replace(".", "").isdigit()
    )
    kode_rusak = None
    for r in range(baris_data, ws.max_row + 1):
        if isinstance(ws.cell(r, kolom_harga).value, (int, float)):
            kode_rusak = str(ws.cell(r, 1).value)
            ws.cell(r, kolom_harga).value = 999999
            break
    assert kode_rusak, "tidak menemukan baris harga untuk dirusak"
    rusak = tmp_path / "rusak.xlsx"
    wb.save(rusak)

    orders = baca_order_sheet(rusak, cfg.customer, tahun_bawaan=2026)
    for o in orders:
        for blok in o.blok:
            for b in blok.baris:
                m = cari_di_master(master, b.kode)
                if m and m[1] > 0:
                    assert abs(b.harga - m[1]) < 0.5, (
                        f"{b.kode}: harga {b.harga} tidak mengikuti master {m[1]}"
                    )
    assert any("HARGA DISESUAIKAN" in p for o in orders for p in o.peringatan), (
        "penyesuaian harga harus DILAPORKAN, tidak boleh diam-diam"
    )


def test_nama_barang_juga_dari_master(berkas):
    """Nama barang di dokumen ikut master, bukan ketikan di baris PO."""
    cfg = Konfigurasi.muat()
    master = baca_master_harga(berkas)
    for o in baca_order_sheet(berkas, cfg.customer, tahun_bawaan=2026):
        for blok in o.blok:
            for b in blok.baris:
                m = cari_di_master(master, b.kode)
                if m and m[0]:
                    assert b.nama == m[0], f"{b.kode}: nama tidak ikut master"
