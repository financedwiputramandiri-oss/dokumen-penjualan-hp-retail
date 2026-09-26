"""Tes format Coretax.

Patokannya template resmi DJP `Template_v.1.6.1.xlsx`: nama lembar, letak
judul, letak data, penutup END, dan aturan pengisian di lembar 'Keterangan'.
Kalau tes ini gagal, berkasnya kemungkinan besar ditolak Converter DJP.
"""
import openpyxl
import pytest

from hp_dokumen.dokumen.coretax import (
    BARIS_DATA_DETAIL, BARIS_DATA_FAKTUR, BARIS_JUDUL_FAKTUR,
    JUDUL_DETAIL, JUDUL_FAKTUR, susun_faktur, tulis,
)
from hp_dokumen.dokumen.invoice import susun_baris
from hp_dokumen.konfigurasi import Konfigurasi
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet
from buat_contoh import buat_contoh


@pytest.fixture
def bahan(tmp_path):
    cfg = Konfigurasi.muat()
    berkas = buat_contoh(tmp_path / "contoh.xlsx")
    orders = [o for o in baca_order_sheet(berkas, cfg.customer, tahun_bawaan=2026) if o.qty > 0]
    return cfg, orders


def _buat(cfg, orders, tmp_path, perusahaan=None):
    kumpulan = []
    pt = perusahaan or cfg.perusahaan.bawaan()
    for o in orders:
        c = cfg.customer.cari(o.nama_tab)
        k = tentukan_nett(o, c)
        baris = susun_baris(o, k, pecah_per_ukuran=False,
                            akhiran_y=cfg.pengaturan.akhiran_y_untuk_angka)
        kumpulan.append(susun_faktur(o, k, c, pt, cfg.pengaturan, baris, "0010726", None))
    out = tmp_path / "coretax.xlsx"
    hasil = tulis(out, kumpulan, pt)
    return openpyxl.load_workbook(out), hasil, kumpulan


def test_lembar_dan_judul_sesuai_template_resmi(bahan, tmp_path):
    cfg, orders = bahan
    wb, _, _ = _buat(cfg, orders, tmp_path)
    assert wb.sheetnames[:2] == ["Faktur", "DetailFaktur"]
    f = wb["Faktur"]
    assert f.cell(1, 1).value == "NPWP Penjual"
    for i, judul in enumerate(JUDUL_FAKTUR, start=1):
        assert f.cell(BARIS_JUDUL_FAKTUR, i).value == judul
    d = wb["DetailFaktur"]
    for i, judul in enumerate(JUDUL_DETAIL, start=1):
        assert d.cell(1, i).value == judul


def test_data_mulai_di_baris_yang_benar_dan_ditutup_END(bahan, tmp_path):
    cfg, orders = bahan
    wb, hasil, _ = _buat(cfg, orders, tmp_path)
    f, d = wb["Faktur"], wb["DetailFaktur"]
    assert f.cell(BARIS_DATA_FAKTUR, 1).value == 1, "faktur pertama di baris 4"
    assert d.cell(BARIS_DATA_DETAIL, 1).value == 1, "detail pertama di baris 2"
    assert f.cell(BARIS_DATA_FAKTUR + hasil.faktur, 1).value == "END"
    assert d.cell(BARIS_DATA_DETAIL + hasil.detail, 1).value == "END"


def test_nilai_wajib_mengikuti_aturan_djp(bahan, tmp_path):
    cfg, orders = bahan
    wb, _, _ = _buat(cfg, orders, tmp_path)
    f = wb["Faktur"]
    r = BARIS_DATA_FAKTUR
    tanggal = f.cell(r, 2).value
    assert len(tanggal) == 10 and tanggal[2] == "/" and tanggal[5] == "/", "format DD/MM/YYYY"
    assert f.cell(r, 3).value == "Normal", "Jenis Faktur selalu 'Normal'"
    assert f.cell(r, 4).value == "01"
    assert f.cell(r, 13).value == "IDN"
    # NPWP pembeli belum ada -> harus diisi nol, bukan dikosongkan
    assert f.cell(r, 11).value == "0000000000000000"
    assert f.cell(r, 18).value == "000000"


def test_dpp_sama_dengan_harga_kali_qty_dikurangi_diskon(bahan, tmp_path):
    """Aturan DJP: DPP = Harga Satuan x Jumlah - Total Diskon."""
    cfg, orders = bahan
    wb, _, _ = _buat(cfg, orders, tmp_path)
    d = wb["DetailFaktur"]
    diperiksa = 0
    for r in range(BARIS_DATA_DETAIL, d.max_row + 1):
        if d.cell(r, 1).value in (None, "END"):
            continue
        harga, qty, diskon, dpp = (d.cell(r, c).value for c in (6, 7, 8, 9))
        assert abs(harga * qty - diskon - dpp) < 0.01, f"baris {r}"
        assert d.cell(r, 10).value == dpp, "DPP Nilai Lain harus sama dengan DPP"
        diperiksa += 1
    assert diperiksa > 0


def test_ppn_sama_dengan_tarif_kali_dpp(bahan, tmp_path):
    cfg, orders = bahan
    wb, _, _ = _buat(cfg, orders, tmp_path)
    d = wb["DetailFaktur"]
    for r in range(BARIS_DATA_DETAIL, d.max_row + 1):
        if d.cell(r, 1).value in (None, "END"):
            continue
        dpp, tarif, ppn = d.cell(r, 10).value, d.cell(r, 11).value, d.cell(r, 12).value
        assert abs(dpp * tarif / 100 - ppn) < 0.01, f"baris {r}"


def test_dpp_plus_ppn_kembali_ke_nilai_invoice(bahan, tmp_path):
    """Harga order sheet SUDAH termasuk PPN, jadi DPP+PPN harus kembali ke nett."""
    cfg, orders = bahan
    wb, hasil, kumpulan = _buat(cfg, orders, tmp_path)
    d = wb["DetailFaktur"]
    jadi = sum((d.cell(r, 9).value or 0) + (d.cell(r, 12).value or 0)
               for r in range(BARIS_DATA_DETAIL, d.max_row + 1)
               if d.cell(r, 1).value not in (None, "END"))
    harapan = sum(nett for _, _, nett in kumpulan)
    assert abs(jadi - harapan) < 1.0
    assert abs(hasil.selisih_pembulatan) < 1.0


def test_perusahaan_tanpa_ppn_menghasilkan_ppn_nol(bahan, tmp_path):
    from hp_dokumen.konfigurasi import Perusahaan
    cfg, orders = bahan
    mtn = Perusahaan(kode="MTN", nama="CV MTN", nama_resmi="CV. MTN",
                     kenakan_ppn=False, npwp="012345678901234567890",
                     alamat_baris=["Jl. A"], logo="", rekening=["REK"],
                     kota_penerbitan="Jakarta")
    wb, _, _ = _buat(cfg, orders, tmp_path, perusahaan=mtn)
    d = wb["DetailFaktur"]
    for r in range(BARIS_DATA_DETAIL, d.max_row + 1):
        if d.cell(r, 1).value in (None, "END"):
            continue
        assert d.cell(r, 12).value == 0 and d.cell(r, 11).value == 0


def test_npwp_penjual_kosong_diberi_peringatan(bahan, tmp_path):
    cfg, orders = bahan
    _, hasil, _ = _buat(cfg, orders, tmp_path)
    assert any("NPWP" in w for w in hasil.peringatan)
    assert not hasil.siap_unggah
