"""Tes bahwa berkas dokumen benar-benar terbentuk dan angkanya cocok."""
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from buat_contoh import buat_contoh
from hp_dokumen.dokumen.faktur_pajak import buat_faktur_pajak
from hp_dokumen.dokumen.invoice import buat_invoice
from hp_dokumen.dokumen.surat_jalan import buat_packing_list, buat_surat_jalan
from hp_dokumen.konfigurasi import Customer, DaftarCustomer, Pengaturan, Perusahaan
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet


@pytest.fixture(scope="module")
def bahan(tmp_path_factory):
    berkas = buat_contoh(tmp_path_factory.mktemp("d") / "contoh.xlsx")
    daftar = DaftarCustomer([
        Customer("Contoh TOP", "PT Contoh Satu", "Jl. Contoh 1", "", "per_artikel", 30, "", "", ""),
        Customer("Contoh CBD", "PT Contoh Dua", "Jl. Contoh 2", "", "per_ukuran", 30, "", "", ""),
    ])
    orders = baca_order_sheet(berkas, daftar, tahun_bawaan=2026)
    perusahaan = Perusahaan(kode="DPM", nama="CV CONTOH", nama_resmi="CV. CONTOH",
                           kenakan_ppn=True, npwp="01.234.567.8-901.000",
                           alamat_baris=["Jl. A", "Kota"], logo="",
                           rekening=["REKENING:", "BANK X"], kota_penerbitan="Jakarta")
    return orders, daftar, perusahaan, Pengaturan()


def _ada_mengandung(ws, kolom, potongan):
    """Cari baris yang isinya MEMUAT potongan teks (bukan sama persis)."""
    return any(potongan in str(ws.cell(r, kolom).value or "")
               for r in range(1, ws.max_row + 1))


def _baris_data_pertama(ws, baris_judul):
    """Baris data pertama sesudah judul tabel.

    Judul invoice asli bertingkat tiga baris, jadi data TIDAK mulai tepat di
    bawah baris "No." seperti anggapan versi sebelumnya.
    """
    r = baris_judul + 1
    while r <= ws.max_row and not isinstance(ws.cell(r, 1).value, int):
        r += 1
    return r


def _cari_baris(ws, kolom, teks):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, kolom).value == teks:
            return r
    return None


def test_surat_jalan_satu_tabel_per_blok(bahan, tmp_path):
    orders, daftar, perusahaan, _ = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    wb = Workbook()
    buat_surat_jalan(wb.active, o, daftar.cari(o.nama_tab), perusahaan, "001")
    p = tmp_path / "sj.xlsx"
    wb.save(p)

    ws = openpyxl.load_workbook(p).active
    judul = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "No."]
    assert len(judul) == len(o.blok) == 2, "harus ada satu tabel per blok"

    # Judul kolom berbahasa Indonesia, sesuai faktur asli DPM
    assert ws.cell(judul[0], 2).value == "ARTICLE CODE"
    assert ws.cell(judul[0], 3).value == "DESKRIPSI BARANG"
    assert ws.cell(judul[0], 6).value == "WARNA"

    # Faktur asli TIDAK punya baris TOTAL per tabel; jumlahnya di kolom Qty
    assert not [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "TOTAL"]
    assert _ada_mengandung(ws, 1, "BRAND : HAPPY PUMPKIN")
    assert _ada_mengandung(ws, 1, "Diterima dengan baik")

    # label ukuran tiap tabel diambil dari bloknya sendiri, mulai kolom G
    assert ws.cell(judul[0], 7).value == "0-3M"
    assert ws.cell(judul[1], 7).value == "1"


def test_packing_list_punya_kolom_gudang(bahan, tmp_path):
    orders, daftar, perusahaan, _ = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    wb = Workbook()
    buat_packing_list(wb.active, o, daftar.cari(o.nama_tab), perusahaan, "001")
    p = tmp_path / "pl.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p).active
    semua = {ws.cell(r, c).value for r in range(1, ws.max_row + 1) for c in range(1, 20)}
    assert "JUMLAH DIKIRIM" in semua
    assert "NO. KOLI" in semua


def test_invoice_satu_tabel_menerus_dan_total_cocok(bahan, tmp_path):
    orders, daftar, perusahaan, pengaturan = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    k = tentukan_nett(o, c)
    wb = Workbook()
    ringkas = buat_invoice(wb.active, o, k, c, perusahaan, pengaturan, "001")
    p = tmp_path / "inv.xlsx"
    wb.save(p)

    ws = openpyxl.load_workbook(p).active
    judul = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "No."]
    assert len(judul) == 1, "invoice tidak boleh dipecah per tabel"

    # Faktur asli DPM: kolom H = Jumlah SETELAH diskon, kolom G = Nilai
    # Diskon, kolom E x kolom D = nilai kotor. Dibuktikan pada
    # 0310726 MAE BEBE: 18 x 62.900 = 1.132.200, G = 283.050, H = 849.150.
    # Jadi jumlah kolom H harus sama dengan nett order sheet, dan
    # H + G harus sama dengan qty x harga baris per baris.
    r = _baris_data_pertama(ws, judul[0])
    jumlah_h = 0.0
    while isinstance(ws.cell(r, 1).value, int):
        qty = ws.cell(r, 4).value or 0
        harga = ws.cell(r, 5).value or 0
        nilai_diskon = ws.cell(r, 7).value or 0
        h = ws.cell(r, 8).value or 0
        assert abs((h + nilai_diskon) - qty * harga) < 0.01, (
            f"baris {r}: kolom H bukan nilai setelah diskon"
        )
        jumlah_h += h
        r += 1
    assert abs(jumlah_h - k.nett_total) < 0.01
    assert abs(ringkas["nett"] - k.nett_total) < 0.01
    assert ringkas["qty"] == o.qty
    # warna tidak boleh muncul di invoice
    teks = {ws.cell(rr, 3).value for rr in range(1, ws.max_row + 1)}
    assert "Merah" not in teks and "Biru" not in teks


def test_invoice_penutup_lengkap_dan_berurutan(bahan, tmp_path):
    orders, daftar, perusahaan, pengaturan = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    k = tentukan_nett(o, c)
    wb = Workbook()
    buat_invoice(wb.active, o, k, c, perusahaan, pengaturan, "001")
    p = tmp_path / "inv2.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p).active
    # penutup ada SESUDAH tabel data, jadi mulai memindai dari baris terakhir tabel
    r = _baris_data_pertama(ws, _cari_baris(ws, 1, "No."))
    while isinstance(ws.cell(r, 1).value, int):
        r += 1
    abaikan = {"Hormat kami,", "(................................)"}
    label = [ws.cell(rr, 7).value for rr in range(r, ws.max_row + 1)
             if isinstance(ws.cell(rr, 7).value, str) and ws.cell(rr, 7).value not in abaikan]
    assert label == ["Subtotal", "Diskon", "Total", "Uang Muka", "DPP", "PPN 11%", "Total"]


def test_faktur_pajak_dpp_plus_ppn_sama_dengan_total(bahan, tmp_path):
    orders, daftar, perusahaan, pengaturan = bahan
    o = next(x for x in orders if "CBD" in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    k = tentukan_nett(o, c)
    wb = Workbook()
    ringkas = buat_faktur_pajak(wb.active, o, k, c, perusahaan, pengaturan, "REF-001")
    assert abs(ringkas["dpp"] + ringkas["ppn"] - k.nett_total) < 0.01
    assert abs(ringkas["dpp"] * (1 + pengaturan.tarif_ppn) - k.nett_total) < 0.01


def test_dokumen_tetap_jadi_walau_data_customer_kosong(bahan, tmp_path):
    """Rekan kerja tidak boleh kena error hanya karena alamat belum diisi."""
    orders, _, perusahaan, pengaturan = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    k = tentukan_nett(o, None)
    wb = Workbook()
    buat_surat_jalan(wb.active, o, None, perusahaan, "001")
    buat_invoice(wb.create_sheet("inv"), o, k, None, perusahaan, pengaturan, "001")
    buat_faktur_pajak(wb.create_sheet("fp"), o, k, None, perusahaan, pengaturan, "001")
    wb.save(tmp_path / "kosong.xlsx")


# ------------------------------------------------- PPN ikut perusahaan
def _perusahaan(kode, kena_ppn):
    from hp_dokumen.konfigurasi import Perusahaan
    return Perusahaan(
        kode=kode, nama=f"CV {kode}", nama_resmi=f"CV. {kode}",
        kenakan_ppn=kena_ppn, npwp="01.2-3", alamat_baris=["Jl. A"],
        logo="", rekening=["REK:"], kota_penerbitan="Jakarta",
    )


def test_ppn_dikenakan_kalau_lewat_perusahaan_pemungut(bahan, tmp_path):
    orders, daftar, _, pengaturan = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    k = tentukan_nett(o, c)
    wb = Workbook()
    hasil = buat_invoice(wb.active, o, k, c, _perusahaan("DPM", True),
                         pengaturan, "001")
    assert hasil["kena_ppn"] is True
    assert hasil["ppn"] > 0
    assert abs(hasil["dpp"] * (1 + pengaturan.tarif_ppn) - hasil["total_setelah_muka"]) < 0.01


def test_ppn_nol_kalau_lewat_perusahaan_tanpa_ppn(bahan, tmp_path):
    """Order lewat CV. Mutiara Timur Nusantara tidak kena PPN 11%."""
    orders, daftar, _, pengaturan = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    k = tentukan_nett(o, c)
    wb = Workbook()
    hasil = buat_invoice(wb.active, o, k, c, _perusahaan("MTN", False), pengaturan, "001")
    assert hasil["kena_ppn"] is False
    assert hasil["ppn"] == 0
    assert abs(hasil["dpp"] - hasil["total_setelah_muka"]) < 0.01


def test_faktur_pajak_ikut_aturan_perusahaan(bahan):
    orders, daftar, _, pengaturan = bahan
    o = next(x for x in orders if "CBD" in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    k = tentukan_nett(o, c)

    wb = Workbook()
    kena = buat_faktur_pajak(wb.active, o, k, c, _perusahaan("DPM", True), pengaturan, "R1")
    assert kena["ppn"] > 0
    assert abs(kena["dpp"] + kena["ppn"] - k.nett_total) < 0.01

    wb2 = Workbook()
    bebas = buat_faktur_pajak(wb2.active, o, k, c, _perusahaan("MTN", False), pengaturan, "R1")
    assert bebas["ppn"] == 0
    assert abs(bebas["dpp"] - k.nett_total) < 0.01


def test_perusahaan_dipilih_dari_master_customer(tmp_path):
    """Kolom perusahaan_pemroses di customer.csv yang menentukan."""
    from hp_dokumen.konfigurasi import Customer, DaftarPerusahaan
    daftar = DaftarPerusahaan([_perusahaan("DPM", True), _perusahaan("MTN", False)],
                              bawaan="DPM", kota="Jakarta")
    lewat_mtn = Customer("X", "PT X", "a", "", "per_artikel", 30, "", "MTN", "")
    lewat_dpm = Customer("Y", "PT Y", "a", "", "per_artikel", 30, "", "DPM", "")
    belum = Customer("Z", "PT Z", "a", "", "per_artikel", 30, "", "", "")
    assert daftar.untuk(lewat_mtn).kenakan_ppn is False
    assert daftar.untuk(lewat_dpm).kenakan_ppn is True
    assert daftar.untuk(belum).kode == "DPM", "kalau kosong, pakai perusahaan bawaan"
    assert daftar.untuk(None).kode == "DPM"


# ================= format harus sama dengan faktur asli DPM =================
# Patokan diambil dari berkas asli di Drive: 0250726 KATAMAMA TAPOS dan
# 0010726 BABY WISE. Kalau tes di bawah ini gagal, JANGAN diubah begitu saja —
# periksa dulu berkas aslinya, karena divisi mengenali fakturnya dari bentuk ini.

def _invoice_jadi(bahan, tmp_path, nama="TOP"):
    orders, daftar, perusahaan, pengaturan = bahan
    o = next(x for x in orders if nama in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    wb = Workbook()
    buat_invoice(wb.active, o, tentukan_nett(o, c), c, perusahaan, pengaturan, "0250726")
    p = tmp_path / "inv_format.xlsx"
    wb.save(p)
    return openpyxl.load_workbook(p).active


def test_invoice_kop_seperti_faktur_asli(bahan, tmp_path):
    ws = _invoice_jadi(bahan, tmp_path)
    # ruang logo digabung A2:B6, teks perusahaan di kolom C mulai baris 2
    assert "A2:B6" in [str(m) for m in ws.merged_cells.ranges]
    assert ws.cell(2, 3).value, "nama perusahaan harus di kolom C baris 2"
    # "FAKTUR No." di A9, nomornya di sel terpisah (C9), keduanya berhuruf
    # besar dan TANPA garis bawah. Dikunci dari empat faktur asli yang
    # sepakat: 0010726 & 0110826 BABY WISE, 0310726 MAE BEBE, 0420826 YULIS.
    assert ws.cell(9, 1).value == "FAKTUR No.", "baris FAKTUR No. wajib ada"
    nomor = ws.cell(9, 3)
    assert nomor.value, "nomor faktur harus di sel terpisah"
    assert nomor.font.underline != "single", "nomor faktur TIDAK bergaris bawah"
    assert nomor.font.size >= 16, "nomor faktur berhuruf besar seperti aslinya"
    # Baris BRAND ADA di faktur, tepat di baris 10. Pernah dihapus karena
    # satu berkas (0020826 CV. BASA MANDIRI) menyendiri — jangan diulang.
    assert "BRAND" in str(ws.cell(10, 1).value or ""), \
        "faktur asli memuat baris BRAND di A10"


def test_invoice_judul_tabel_tiga_tingkat(bahan, tmp_path):
    ws = _invoice_jadi(bahan, tmp_path)
    assert ws.cell(12, 1).value == "No."
    assert ws.cell(12, 2).value == "ARTICLE CODE"
    assert ws.cell(12, 3).value == "DESKRIPSI BARANG"
    assert ws.cell(12, 4).value == "Qty" and ws.cell(13, 4).value == "PCS"
    assert ws.cell(12, 5).value == "Harga" and ws.cell(13, 5).value == "Satuan"
    assert str(ws.cell(12, 6).value).strip() == "Diskon"
    assert ws.cell(12, 7).value == "Nilai" and ws.cell(13, 7).value == "Diskon"
    assert ws.cell(12, 8).value == "Jumlah"
    assert isinstance(ws.cell(15, 1).value, int), "data mulai baris 15"


def test_invoice_persen_diskon_di_bawah_labelnya(bahan, tmp_path):
    """Persen diskon ada di F13, bukan F14.

    Faktur asli 0010726 BABY WISE dan 0310726 MAE BEBE sama-sama menaruh
    label "Diskon " sendirian di F12 dan persennya di sel gabungan F13:F14.
    Versi lama terbalik: label digabung F12:F13 dan persen jatuh ke F14.
    """
    ws = _invoice_jadi(bahan, tmp_path)
    assert str(ws.cell(12, 6).value).strip() == "Diskon"
    assert ws.cell(13, 6).value not in (None, ""), "persen diskon harus di F13"
    assert "F13:F14" in {str(m) for m in ws.merged_cells.ranges}


def test_invoice_memuat_blok_rekening(bahan, tmp_path):
    ws = _invoice_jadi(bahan, tmp_path)
    # Faktur asli menaruh blok rekening di kolom B (0420826 YULIS B60:C63,
    # 0110826 BABY WISE B23, 0310726 MAE BEBE B24, 0010726 BABY WISE B75).
    assert _ada_mengandung(ws, 2, "PEMBAYARAN DITRANSFER KE REKENING")
    # kalimat pengantar tidak boleh muncul dua kali
    jumlah = sum("PEMBAYARAN DITRANSFER" in str(ws.cell(r, 2).value or "")
                 for r in range(1, ws.max_row + 1))
    assert jumlah == 1, "baris rekening kembar"


def test_surat_jalan_tidak_mencetak_kolom_ukuran_yang_kosong(bahan, tmp_path):
    """Kolom ukuran ke-9 order sheet tidak pernah terisi; jangan dicetak."""
    orders, daftar, perusahaan, _ = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    wb = Workbook()
    buat_surat_jalan(wb.active, o, daftar.cari(o.nama_tab), perusahaan, "001")
    p = tmp_path / "sj_format.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p).active

    for baris_judul in [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "No."]:
        kolom = 7
        while ws.cell(baris_judul, kolom).value not in (None, "", "Qty"):
            label = str(ws.cell(baris_judul, kolom).value)
            terpakai = any(ws.cell(rr, kolom).value
                           for rr in range(baris_judul + 2, baris_judul + 40)
                           if isinstance(ws.cell(rr, 1).value, int))
            assert terpakai, f"kolom ukuran '{label}' tercetak tapi kosong"
            kolom += 1


def test_surat_jalan_tanpa_total_qty_di_kanan_atas(bahan, tmp_path):
    """Angka total qty TIDAK boleh tercetak di kanan atas Surat Jalan.

    Permintaan Yosua 14 September 2026. Memang begitu aslinya juga: di
    0110826 BABY WISE angka itu ada di kolom N, DI LUAR print_area (A1:M31),
    jadi tidak pernah ikut tercetak. Versi lama menaruhnya di dalam area
    cetak sehingga muncul di kertas.
    """
    orders, daftar, perusahaan, _ = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    wb = Workbook()
    buat_surat_jalan(wb.active, o, daftar.cari(o.nama_tab), perusahaan, "0010126")
    p = tmp_path / "sj_tanpa_total.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p).active

    baris_judul = next(r for r in range(1, ws.max_row + 1)
                       if ws.cell(r, 1).value == "No.")
    # Di atas baris judul tabel tidak boleh ada angka sebesar total qty.
    for r in range(1, baris_judul):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            assert v != o.qty, (
                f"total qty {o.qty} masih tercetak di {ws.cell(r, c).coordinate}"
            )


def test_surat_jalan_nomor_besar_tanpa_awalan_no(bahan, tmp_path):
    """Nomor ditulis besar dan polos, seperti 0110826 BABY WISE (huruf 18)."""
    orders, daftar, perusahaan, _ = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    wb = Workbook()
    buat_surat_jalan(wb.active, o, daftar.cari(o.nama_tab), perusahaan, "0010126")
    p = tmp_path / "sj_nomor.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p).active

    sel = next((ws.cell(r, c)
                for r in range(1, 14) for c in range(1, ws.max_column + 1)
                if str(ws.cell(r, c).value or "").strip() == "0010126"), None)
    assert sel is not None, "nomor harus ditulis polos, tanpa awalan 'No. '"
    assert sel.font.size >= 14, "nomor Surat Jalan harus berukuran besar"


def test_surat_jalan_urutan_tanda_tangan(bahan, tmp_path):
    """Penerima dulu, lalu Pengirim, lalu Mengetahui."""
    orders, daftar, perusahaan, _ = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    wb = Workbook()
    buat_surat_jalan(wb.active, o, daftar.cari(o.nama_tab), perusahaan, "001")
    p = tmp_path / "sj_ttd.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p).active

    urut = [str(ws.cell(r, c).value).strip().rstrip(":").strip()
            for r in range(1, ws.max_row + 1)
            for c in range(1, ws.max_column + 1)
            if str(ws.cell(r, c).value or "").strip().rstrip(":").strip()
            in {"Penerima", "Pengirim", "Mengetahui"}]
    assert urut == ["Penerima", "Pengirim", "Mengetahui"], urut


def test_invoice_tanpa_diskon_tidak_mencetak_kolom_diskon(bahan, tmp_path):
    """Order tanpa diskon memakai bentuk 0020826 CV. BASA MANDIRI.

    Kolom Diskon dan baris Subtotal/Diskon tidak dicetak sama sekali —
    kolom berisi nol hanya menimbulkan pertanyaan dari customer.
    """
    orders, daftar, perusahaan, pengaturan = bahan
    o = next(x for x in orders if "TOP" in x.nama_tab)
    c = daftar.cari(o.nama_tab)
    k = tentukan_nett(o, c)
    # buat order ini seolah tanpa diskon: nilai bersih = nilai kotor
    k.nett_total = sum(b.nilai_kotor for blok in o.blok for b in blok.baris)
    for blok in o.blok:
        for b in blok.baris:
            b.nett[k.kolom] = b.nilai_kotor

    wb = Workbook()
    buat_invoice(wb.active, o, k, c, perusahaan, pengaturan, "0010126")
    p = tmp_path / "inv_tanpa_diskon.xlsx"
    wb.save(p)
    ws = openpyxl.load_workbook(p).active

    label = {str(ws.cell(r, 7).value or "").strip()
             for r in range(1, ws.max_row + 1)}
    assert "Total" in label
    assert "Subtotal" not in label, "baris Subtotal tidak ada di faktur tanpa diskon"
    assert "Diskon" not in label, "baris Diskon tidak ada di faktur tanpa diskon"
    judul = {str(ws.cell(12, c).value or "").strip() for c in range(1, 9)}
    assert "Diskon" not in judul, "kolom Diskon tidak dicetak kalau tidak ada diskon"
