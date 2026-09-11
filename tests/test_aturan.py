"""Tes keempat aturan wajib, memakai order sheet CONTOH (data karangan)."""
from pathlib import Path

import pytest
from openpyxl import Workbook

from buat_contoh import buat_contoh
from hp_dokumen.dokumen.invoice import susun_baris
from hp_dokumen.konfigurasi import DaftarCustomer, Customer, Pengaturan
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet, baca_master_harga
from hp_dokumen.rekonsiliasi import periksa_order
from hp_dokumen.ukuran import bagi_rata_nilai, rapikan_label, tulis_untuk_deskripsi

AKAR = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="module")
def contoh(tmp_path_factory):
    return buat_contoh(tmp_path_factory.mktemp("data") / "contoh.xlsx")


@pytest.fixture(scope="module")
def daftar():
    return DaftarCustomer([
        Customer("Contoh TOP", "PT Contoh Satu", "Jl. Contoh 1", "", "per_artikel", 30, "", "", ""),
        Customer("Contoh CBD", "PT Contoh Dua", "Jl. Contoh 2", "", "per_ukuran", 30, "", "", ""),
    ])


@pytest.fixture(scope="module")
def orders(contoh, daftar):
    return baca_order_sheet(contoh, daftar, tahun_bawaan=2026)


def ambil(orders, kata):
    return next(o for o in orders if kata in o.nama_tab)


# ---------------------------------------------------- Aturan 1: semua blok
def test_semua_blok_tertangkap(orders):
    top = ambil(orders, "TOP")
    assert len(top.blok) == 2, "blok kedua tidak boleh hilang"
    assert top.jumlah_baris == 3
    assert top.qty == 5 + 3 + 10


def test_label_ukuran_per_blok_berbeda(orders):
    top = ambil(orders, "TOP")
    assert top.blok[0].label_ukuran[:3] == ["0-3M", "3-6M", "6-12M"]
    assert top.blok[1].label_ukuran[:3] == ["1", "2", "3"]


def test_qty_diambil_dari_kolom_ato(orders):
    """Qty harus dari N..V, bukan D..L."""
    top = ambil(orders, "TOP")
    assert top.blok[0].baris[0].qty_per_ukuran[:3] == [2, 3, 0]


def test_cocok_dengan_baris_total_order_sheet(orders):
    for o in orders:
        assert o.total_sheet.qty == o.qty
        assert abs(o.total_sheet.nilai_kotor - o.nilai_kotor) < 1


# ------------------------------------------- Aturan 2: nett dari kolomnya
def test_cbd_terisi_penuh_dipakai(orders, daftar):
    cbd = ambil(orders, "CBD")
    k = tentukan_nett(cbd, daftar.cari(cbd.nama_tab))
    assert k.kolom == "CBD"
    assert k.cara_bayar == "CBD"
    diharapkan = sum(b.disc_cbd for b in cbd.semua_baris)
    assert abs(k.nett_total - diharapkan) < 0.01


def test_cbd_terisi_sebagian_diabaikan_seluruhnya(orders, daftar):
    """Satu baris terisi dari tiga -> seluruh kolom diabaikan, order jadi TOP."""
    top = ambil(orders, "TOP")
    k = tentukan_nett(top, daftar.cari(top.nama_tab))
    assert k.kolom == "TOP"
    assert k.cara_bayar == "TOP"
    assert 0 < k.terisi_cbd < k.jumlah_baris
    diharapkan = sum(b.total_value for b in top.semua_baris)
    assert abs(k.nett_total - diharapkan) < 0.01
    # nilai kolom AD yang terisi sebagian TIDAK boleh ikut terhitung
    sebagian = sum(b.disc_cbd or 0 for b in top.semua_baris)
    assert abs(k.nett_total - sebagian) > 1


def test_override_manual_dilaporkan(orders):
    top = ambil(orders, "TOP")
    paksa = Customer("Contoh TOP", "PT X", "", "", "per_artikel", 30, "CBD", "", "")
    k = tentukan_nett(top, paksa)
    assert k.dioverride is True
    assert k.kolom == "CBD"


# ------------------------------- Aturan 3 & 4: label ukuran & deskripsi
def test_label_angka_dirapikan():
    assert rapikan_label(2.0) == "2"
    assert rapikan_label("7-8Y") == "7-8Y"
    assert rapikan_label(None) is None
    assert rapikan_label("   ") is None


def test_akhiran_y_hanya_untuk_angka_polos():
    assert tulis_untuk_deskripsi("2", True) == "2Y"
    assert tulis_untuk_deskripsi("2", False) == "2"
    for tetap in ("0-3M", "7-8Y", "S", "M", "L", "XL", "NB"):
        assert tulis_untuk_deskripsi(tetap, True) == tetap


def test_invoice_per_artikel_menggabungkan_warna(orders, daftar):
    top = ambil(orders, "TOP")
    k = tentukan_nett(top, daftar.cari(top.nama_tab))
    baris = susun_baris(top, k, pecah_per_ukuran=False, akhiran_y=True)
    kode = [x.kode for x in baris]
    assert kode.count("AA.1") == 1, "dua warna harus jadi satu baris invoice"
    assert sum(x.qty for x in baris) == top.qty
    assert abs(sum(x.nett for x in baris) - k.nett_total) < 0.01


def test_invoice_per_ukuran_memakai_label_blok_asal(orders, daftar):
    cbd = ambil(orders, "CBD")
    k = tentukan_nett(cbd, daftar.cari(cbd.nama_tab))
    baris = susun_baris(cbd, k, pecah_per_ukuran=True, akhiran_y=True)
    desk = {x.deskripsi for x in baris}
    assert "Contoh Set Big Size - S" in desk
    assert "Contoh Set Big Size - M" in desk
    assert "Contoh Set Big Size - L" in desk
    # total tetap sama persis walau dipecah
    assert sum(x.qty for x in baris) == cbd.qty
    assert abs(sum(x.nett for x in baris) - k.nett_total) < 0.005


def test_pembagian_nilai_per_ukuran_tidak_menimbulkan_selisih():
    for total in (100.0, 1000.01, 892688.0, 42939750.0, 1.0):
        for bobot in ([1, 1, 1], [3, 5, 7, 11], [1], [2, 2]):
            bagian = bagi_rata_nilai(total, bobot)
            assert abs(sum(bagian) - total) < 0.005
            assert len(bagian) == len(bobot)
    assert bagi_rata_nilai(100.0, [0, 0]) == [0.0, 0.0]


# ------------------------------------------------------- rekonsiliasi
def test_rekonsiliasi_lolos_untuk_data_bersih(orders, daftar, contoh):
    master = baca_master_harga(contoh)
    for o in orders:
        c = daftar.cari(o.nama_tab)
        h = periksa_order(o, tentukan_nett(o, c), Pengaturan(), master, c)
        assert h.lolos, [p.nama for p in h.yang_gagal]


def test_rekonsiliasi_menangkap_selisih(contoh, daftar, tmp_path):
    """Kalau baris TOTAL order sheet diubah, pencocokan harus GAGAL."""
    import openpyxl

    rusak = tmp_path / "rusak.xlsx"
    wb = openpyxl.load_workbook(contoh)
    ws = wb["PO 06 Januari - Contoh CBD"]
    for r in range(1, ws.max_row + 1):
        nilai = ws.cell(r, 23).value
        if isinstance(nilai, (int, float)) and not ws.cell(r, 1).value:
            ws.cell(r, 23, int(nilai) + 99)   # rusakkan baris TOTAL milik sheet
            break
    else:
        pytest.fail("baris TOTAL contoh tidak ketemu")
    wb.save(rusak)

    orders = baca_order_sheet(rusak, daftar, tahun_bawaan=2026)
    o = next(x for x in orders if "CBD" in x.nama_tab)
    h = periksa_order(o, tentukan_nett(o, daftar.cari(o.nama_tab)), Pengaturan())
    assert not h.lolos
    assert any("qty" in p.nama.lower() for p in h.yang_gagal)


def test_peringatan_saat_cbd_terisi_sebagian(orders, daftar):
    top = ambil(orders, "TOP")
    c = daftar.cari(top.nama_tab)
    h = periksa_order(top, tentukan_nett(top, c), Pengaturan(), None, c)
    assert any("terisi" in w and "CBD" in w for w in h.peringatan)
