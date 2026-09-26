"""Hasil sapuan harus sampai ke Drive — dua jalur, dua-duanya dikunci di sini.

Permintaan Yosua 20 September 2026: "atur sendiri agar hasil sapuan dapat
diupload ke drive". Akun layanan tidak bisa mengunggah berkas baru
(`storageQuotaExceeded`, CLAUDE.md bagian 14 dan 30), jadi jalurnya:

  1. tab BOT_REKAP di sheet OTOMATISASI — lewat Sheets API ke berkas yang
     SUDAH ADA, jadi lepas dari batasan kuota;
  2. berkas laporan ditaruh di `folder_draf/_LAPORAN` supaya ikut naik lewat
     Google Drive for Desktop, jalur yang sama dengan dokumen.
"""
from pathlib import Path

import pytest

from hp_dokumen.sapu import bot, rekap
from hp_dokumen.sapu.laporan_sapu import tulis
from hp_dokumen.sapu.tulis_sheet import TAB_REKAP, AWALAN_BOT


def _rekaman():
    return [
        {"sumber": "OS Sept", "tab": "PO A", "customer": "A", "qty": 60,
         "kotor": 3917000.0, "nett": 3009431.0, "cara_bayar": "CBD", "siap": True},
        {"sumber": "OS Sept", "tab": "PO B", "customer": "B", "qty": 10,
         "kotor": 500000.0, "nett": 400000.0, "cara_bayar": "TOP", "siap": True},
        {"sumber": "OS Feb", "tab": "PO C", "customer": "C", "qty": 99,
         "kotor": 999999.0, "nett": 888888.0, "siap": False,
         "keterangan": "angka belum cocok: TOTAL ATO"},
    ]


# ------------------------------------------------------- angka rekap
def test_rekap_menjumlah_hanya_po_yang_dokumennya_terbit():
    baris = rekap.baris_per_sheet(_rekaman())
    total = baris[-1]
    assert total[0] == "TOTAL"
    assert total[1] == 3, "jumlah PO harus menghitung yang belum terbit juga"
    assert total[2] == 2, "hanya 2 PO yang dokumennya terbit"
    assert total[3] == 70, "qty PO yang gagal tidak boleh ikut dijumlah"
    assert total[5] == 3409431


def test_angka_po_yang_belum_terbit_DIKOSONGKAN_bukan_ditampilkan():
    """Angka itu justru yang belum cocok dengan baris TOTAL order sheet.

    Kalau ikut ditampilkan, orang memakainya sebagai angka resmi — padahal
    itu persis yang sedang bermasalah.
    """
    feb = [b for b in rekap.baris_per_sheet(_rekaman()) if b[0] == "OS Feb"][0]
    assert feb[3] == "" and feb[4] == "" and feb[5] == "", f"angka bocor: {feb}"
    assert "belum cocok" in feb[6]

    per_po = [b for b in rekap.baris_per_po(_rekaman()) if b[1] == "PO C"][0]
    assert per_po[4] == "" and per_po[7] == ""
    assert per_po[8] == rekap.TIDAK_TERBIT


# --------------------------------------------- laporan ikut ke Drive
def test_laporan_ditaruh_sefolder_dengan_dokumen(tmp_path):
    """Kalau laporan ditaruh di folder proyek, ia tidak pernah sampai ke Drive."""
    p = bot.Pengaturan(
        folder=[], folder_laporan_id="", nama_folder_laporan="X",
        sheet_otomatisasi_id="", buat_draf=True, pantau_perubahan=True,
        pantau_rumus=True, hanya_hari=0,
        berkas_kredensial=tmp_path / "k.json",
        berkas_kondisi=tmp_path / "kondisi.json",
        folder_draf=tmp_path / "DOKUMEN OTOMATIS",
    )
    assert bot.folder_laporan(p).parent == p.folder_draf


def test_laporan_memuat_lembar_rekap(tmp_path):
    berkas = tulis(tmp_path / "lap.xlsx", [], [("a", "b", 1)], [], [],
                   rekaman=_rekaman())
    from openpyxl import load_workbook
    wb = load_workbook(berkas)
    assert "REKAP" in wb.sheetnames and "REKAP PER PO" in wb.sheetnames
    isi = [c.value for row in wb["REKAP"].iter_rows() for c in row]
    assert "TOTAL" in isi


# ------------------------------------------------- pemangkasan aman
def test_pangkas_hanya_menyentuh_berkas_laporan(tmp_path):
    """Folder `_LAPORAN` ada di dalam folder Drive yang disinkronkan.

    Kalau polanya dilonggarkan, berkas orang lain di situ ikut terhapus —
    dan hilangnya sampai ke Drive.
    """
    f = tmp_path / "_LAPORAN"
    f.mkdir()
    for n in range(5):
        (f / f"LAPORAN_SAPUAN_2026092{n}_0600.xlsx").write_text("x")
    lain = [f / "Catatan Yosua.xlsx", f / "LAPORAN_SAPUAN_manual.txt",
            f / "INVOICE_PO_A.xlsx"]
    for x in lain:
        x.write_text("jangan dihapus")

    bot.pangkas_laporan(f, 2)

    tersisa = sorted(x.name for x in f.glob("LAPORAN_SAPUAN_20*.xlsx"))
    assert tersisa == ["LAPORAN_SAPUAN_20260923_0600.xlsx",
                       "LAPORAN_SAPUAN_20260924_0600.xlsx"]
    for x in lain:
        assert x.exists(), f"{x.name} ikut terhapus"


def test_pangkas_nol_berarti_jangan_hapus_apa_pun(tmp_path):
    f = tmp_path / "_LAPORAN"
    f.mkdir()
    for n in range(4):
        (f / f"LAPORAN_SAPUAN_2026092{n}_0600.xlsx").write_text("x")
    assert bot.pangkas_laporan(f, 0) == []
    assert len(list(f.glob("*.xlsx"))) == 4


# ------------------------------------------------------ tab BOT_REKAP
def test_tab_rekap_memakai_awalan_bot():
    """Bot hanya boleh menulis ke tab berawalan BOT_ (CLAUDE.md bagian 14)."""
    assert TAB_REKAP.startswith(AWALAN_BOT)
