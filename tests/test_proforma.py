"""Tes Proforma Invoice dan kedua gaya faktur.

Dua hal yang dikunci di sini, keduanya dari permintaan Yosua 17 Sep 2026:

1. Faktur per ukuran (Haritsa, Katamama) memakai template yang BERBEDA dari
   faktur per artikel. Sebelumnya program memakai satu template untuk semua,
   dan itulah yang membuat faktur Katamama terlihat salah.
2. Proforma memuat level VARIAN di kolom KETERANGAN, dan hanya diterbitkan
   untuk customer yang fakturnya dipecah per ukuran.
"""
import pytest
from openpyxl import Workbook

from buat_contoh import buat_contoh
from hp_dokumen.dokumen.invoice import (
    GAYA_PER_ARTIKEL, GAYA_PER_UKURAN, buat_invoice,
)
from hp_dokumen.dokumen.proforma import JUDUL_KOLOM, buat_proforma, keterangan
from hp_dokumen.konfigurasi import Customer, DaftarCustomer, Pengaturan, Perusahaan
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet


@pytest.fixture(scope="module")
def bahan(tmp_path_factory):
    berkas = buat_contoh(tmp_path_factory.mktemp("d") / "contoh.xlsx")
    daftar = DaftarCustomer([
        Customer("Contoh TOP", "PT Contoh Satu", "Jl. Contoh 1", "", "per_artikel", 30, "", "", ""),
        Customer("Contoh CBD", "PT Contoh Dua", "Jl. Contoh 2", "", "per_ukuran", 30, "", "", ""),
    ])
    orders = baca_order_sheet(berkas, daftar, tahun_bawaan=2026)
    perusahaan = Perusahaan(
        kode="DPM", nama="CV CONTOH", nama_resmi="CV. CONTOH", kenakan_ppn=True,
        npwp="01.234.567.8-901.000",
        alamat_baris=["Jl. A No. 1, Kelurahan Contoh", "Kota Contoh - 11460",
                      "Phone. 021", "Email : a@b.c"],
        logo="", rekening=["REKENING:", "BANK X"], kota_penerbitan="Jakarta",
    )
    return orders, daftar, perusahaan, Pengaturan()


def _proforma(bahan):
    orders, daftar, pt, peng = bahan
    o = next(x for x in orders if daftar.cari(x.nama_tab)
             and daftar.cari(x.nama_tab).pecah_per_ukuran)
    ws = Workbook().active
    ringkas = buat_proforma(ws, o, tentukan_nett(o, daftar.cari(o.nama_tab)),
                            daftar.cari(o.nama_tab), pt, peng, "0010826")
    return ws, ringkas, o


def test_kolom_proforma_sesuai_foto(bahan):
    """Sembilan kolom, urutannya persis seperti foto dari Yosua."""
    ws, _, _ = _proforma(bahan)
    diharapkan = [t for _, t in JUDUL_KOLOM]
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "NO":
            nyata = [ws.cell(r, c).value for c in range(1, 10)]
            assert nyata == diharapkan
            return
    pytest.fail("baris judul tabel proforma tidak ketemu")


def test_keterangan_proforma_memuat_varian(bahan):
    """Permintaan inti Yosua: level varian ikut di kolom KETERANGAN."""
    ws, _, _ = _proforma(bahan)
    isi = [str(ws.cell(r, 3).value or "") for r in range(1, ws.max_row + 1)]
    assert any(" - " in x and x.strip() for x in isi), (
        "tidak ada satu pun baris KETERANGAN yang memuat varian"
    )


def test_keterangan_tanpa_warna_tidak_menggantung():
    """Warna kosong tidak boleh meninggalkan tanda hubung menggantung."""
    class B:
        deskripsi, warna = "Kaos Uk. M", ""
    assert keterangan(B()) == "Kaos Uk. M"
    B.warna = "Merah"
    assert keterangan(B()) == "Kaos Uk. M - Merah"


def test_penutup_proforma_konsisten(bahan):
    """Sub Total - Diskon harus sama dengan Grand Total, dan Grand Total = nett."""
    ws, ringkas, _ = _proforma(bahan)
    nilai = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 6).value
        if isinstance(label, str) and label.strip():
            nilai[label.strip()] = ws.cell(r, 9).value
    assert abs(nilai["Sub Total"] - ringkas["kotor"]) < 0.5
    assert abs(nilai["Grand Total"] - ringkas["nett"]) < 0.5
    assert abs(nilai["Sub Total"] - nilai["Diskon"] - nilai["Grand Total"]) < 0.5


def test_proforma_hanya_untuk_customer_per_ukuran(bahan, tmp_path):
    """Customer per artikel TIDAK boleh ikut dibuatkan proforma.

    Yosua hanya meminta Haritsa dan Katamama. Menerbitkannya untuk semua
    customer berarti mengirim dokumen yang tidak pernah diminta.
    """
    from hp_dokumen.berkas_dokumen import buat_berkas

    orders, daftar, pt, peng = bahan

    class Cfg:
        customer = daftar
        pengaturan = peng

        class perusahaan:
            @staticmethod
            def untuk(_):
                return pt

    for o in orders:
        cust = daftar.cari(o.nama_tab)
        if cust is None:
            continue
        folder = tmp_path / o.nama_tab.replace(" ", "_")
        berkas = buat_berkas(Cfg(), o, tentukan_nett(o, cust), folder)
        punya = any("PROFORMA" in p.name for p in berkas)
        assert punya == cust.pecah_per_ukuran, (
            f"{o.nama_tab}: proforma={punya}, per_ukuran={cust.pecah_per_ukuran}"
        )


def test_dua_gaya_faktur_benar_benar_berbeda(bahan):
    """Kop faktur per ukuran memakai huruf 16/12, per artikel 18/11.

    Dibongkar dari TIGA faktur asli: 0110826 BABY WISE (per artikel),
    0160826 HARITSA dan 0400826 KATAMAMA (keduanya per ukuran). Haritsa dan
    Katamama sepakat melawan Baby Wise di setiap ukuran huruf.
    """
    orders, daftar, pt, peng = bahan
    hasil = {}
    for kunci, per_ukuran in (("artikel", False), ("ukuran", True)):
        o = next(x for x in orders if daftar.cari(x.nama_tab)
                 and daftar.cari(x.nama_tab).pecah_per_ukuran == per_ukuran)
        cust = daftar.cari(o.nama_tab)
        ws = Workbook().active
        buat_invoice(ws, o, tentukan_nett(o, cust), cust, pt, peng, "0010826")
        hasil[kunci] = ws

    a, u = hasil["artikel"], hasil["ukuran"]
    assert a["C2"].font.size == GAYA_PER_ARTIKEL.ukuran_nama_perusahaan == 18
    assert u["C2"].font.size == GAYA_PER_UKURAN.ukuran_nama_perusahaan == 16
    assert a["A9"].font.size == 18 and u["A9"].font.size == 16
    assert a["A10"].font.size == 18 and u["A10"].font.size == 16
    assert a["C3"].font.bold is False and u["C3"].font.bold is True
    assert a.row_dimensions[15].height == 26.25
    assert u.row_dimensions[15].height == 31.5


def test_alamat_perusahaan_dilipat_tanpa_menelan_baris_kontak():
    """Melipat alamat tidak boleh menggabungkan Phone/Email ke dalam alamat.

    Kalau tertelan, nomor telepon bisa muncul di tengah baris alamat.
    """
    from hp_dokumen.dokumen.gaya import _alamat_perusahaan

    baris = ["Jl. A No. 1, Kelurahan Contoh", "Kota Contoh - 11460",
             "Phone. 021", "Wa:(+62) 0", "Email : a@b.c"]
    hasil = _alamat_perusahaan(baris, 44)
    assert hasil[-3:] == ["Phone. 021", "Wa:(+62) 0", "Email : a@b.c"]
    assert all(len(x) <= 44 for x in hasil[:-3])
    # tanpa batas, tidak ada yang diubah
    assert _alamat_perusahaan(baris, None) == baris
