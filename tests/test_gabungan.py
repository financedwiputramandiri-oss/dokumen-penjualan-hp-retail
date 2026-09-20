"""Satu berkas Excel: Invoice + Surat Jalan + Master Harga, dengan rumus.

Permintaan Yosua 20 September 2026, sesudah ia membetulkan sendiri berkas
keluaran program di Excel dan mengirimkannya kembali. Yang dikunci di sini
adalah hal-hal yang kalau salah TIDAK kelihatan dari angka totalnya:
ukuran logo, garis kepala dokumen, dan rumus yang menunjuk ke lembar lain.
"""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from buat_contoh import buat_contoh
from hp_dokumen.berkas_dokumen import buat_berkas
from hp_dokumen.dokumen import gaya
from hp_dokumen.dokumen import rumus as rms
from hp_dokumen.konfigurasi import Konfigurasi, Customer, DaftarCustomer
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet, master_harga_dari_buku


def _daftar(kode_pt: str) -> DaftarCustomer:
    # Nama & alamat SENGAJA dikosongkan: itu keadaan sebenarnya untuk
    # sebagian besar customer, dan Yosua sendiri yang akan mengisinya.
    return DaftarCustomer([
        Customer("Contoh TOP", "", "", "", "per_artikel", 30, "", kode_pt, ""),
    ])


@pytest.fixture(scope="module")
def cfg():
    return Konfigurasi.muat()


def _siapkan(tmp_path, cfg, kode_pt):
    daftar = _daftar(kode_pt)
    sumber = buat_contoh(tmp_path / "contoh.xlsx")
    order = next(o for o in baca_order_sheet(sumber, daftar, tahun_bawaan=2026)
                 if "TOP" in o.nama_tab)
    master = master_harga_dari_buku(load_workbook(sumber))
    cfg.customer = daftar
    cust = daftar.cari(order.nama_tab)
    berkas = buat_berkas(cfg, order, tentukan_nett(order, cust),
                         tmp_path / "keluar", "001", master=master)
    utama = next(p for p in berkas if p.name.startswith("INVOICE_SURAT_JALAN"))
    return load_workbook(utama), master


@pytest.fixture(params=["DPM", "MTN"])
def buku(request, tmp_path, cfg):
    asli = cfg.customer
    try:
        wb, master = _siapkan(tmp_path, cfg, request.param)
    finally:
        cfg.customer = asli
    return wb, master, request.param


# ------------------------------------------------------- satu berkas
def test_invoice_dan_surat_jalan_satu_berkas_lembar_berbeda(buku):
    wb, _, _ = buku
    assert wb.sheetnames[:2] == [rms.TAB_INVOICE, rms.TAB_SURAT_JALAN]
    assert rms.TAB_MASTER in wb.sheetnames


def test_lembar_master_harga_memuat_harga_saat_dokumen_dibuat(buku):
    """Harga retail berubah sepanjang tahun; faktur lama harus bawa buktinya."""
    wb, master, _ = buku
    ws = wb[rms.TAB_MASTER]
    kode = {ws.cell(r, 1).value for r in range(rms.BARIS_DATA_MASTER, ws.max_row + 1)}
    assert set(master) <= kode, "ada artikel master yang tidak ikut tersalin"
    assert ws.cell(rms.BARIS_TARIF, rms.KOLOM_TARIF).value is not None, "tarif PPN"


# ------------------------------------------------------------- logo
def test_logo_sama_persis_di_invoice_dan_surat_jalan(buku):
    """Ini keluhan Yosua: logo berbeda ukuran antar dokumen.

    Sebabnya lebar kolom B menyempit mengikuti kode artikel yang pendek,
    sehingga logonya ikut mengecil di dokumen yang kodenya lebih pendek.
    """
    wb, _, _ = buku
    ukuran = []
    for nama in (rms.TAB_INVOICE, rms.TAB_SURAT_JALAN):
        gambar = wb[nama]._images
        assert gambar, f"logo hilang di lembar {nama}"
        a = gambar[0].anchor
        ukuran.append((round(a.ext.cx / 9525), round(a.ext.cy / 9525)))
    assert ukuran[0] == ukuran[1], f"logo beda ukuran: {ukuran}"


def test_logo_seukuran_berkas_asli_bukan_86_piksel_lagi(buku):
    """Diukur dari berkas asli: DPM +-119x120 px, MTN +-111x126 px.

    Versi lama memakai tinggi tetap 86 px untuk semua perusahaan, jadi
    logonya jauh lebih kecil daripada dokumen yang dipakai divisi.
    """
    wb, _, kode = buku
    a = wb[rms.TAB_INVOICE]._images[0].anchor
    lebar, tinggi = a.ext.cx / 9525, a.ext.cy / 9525
    assert lebar == pytest.approx(gaya.LEBAR_LOGO[kode], abs=1)
    assert tinggi > 100, f"logo masih kekecilan: {lebar:.0f}x{tinggi:.0f}"


def test_logo_tidak_pernah_gepeng(buku):
    """Tingginya mengikuti bentuk asli gambar, bukan angka tetap."""
    wb, _, _ = buku
    img = wb[rms.TAB_INVOICE]._images[0]
    a = img.anchor
    rasio_asli = img.width / img.height
    rasio_pasang = (a.ext.cx / 9525) / (a.ext.cy / 9525)
    assert rasio_pasang == pytest.approx(rasio_asli, rel=0.02)


# ------------------------------------------------------------ rumus
def _sel(ws, kolom, awal=1):
    for r in range(awal, ws.max_row + 1):
        nilai = ws.cell(r, kolom).value
        if isinstance(nilai, str) and nilai.startswith("="):
            return r, nilai
    return None, None


def test_harga_satuan_diambil_lewat_vlookup_ke_master(buku):
    wb, _, _ = buku
    _, isi = _sel(wb[rms.TAB_INVOICE], 5)
    assert isi and "VLOOKUP" in isi and rms.TAB_MASTER in isi


def test_jumlah_per_baris_berupa_rumus_qty_kali_harga(buku):
    wb, _, _ = buku
    _, isi = _sel(wb[rms.TAB_INVOICE], 8)
    assert isi and "*" in isi, f"kolom Jumlah bukan rumus: {isi!r}"


def test_penutup_memakai_rumus_bukan_angka_mati(buku):
    wb, _, _ = buku
    ws = wb[rms.TAB_INVOICE]
    semua = [str(ws.cell(r, 8).value or "") for r in range(1, ws.max_row + 1)]
    assert any("SUMPRODUCT" in x for x in semua), "Subtotal bukan rumus"
    assert any(x.startswith("=SUM(") for x in semua), "Diskon bukan rumus"


def test_nilai_diskon_per_baris_TETAP_angka_bukan_rumus(buku):
    """Sengaja, dan jangan diubah tanpa membaca dokumen/rumus.py.

    Nilai bersih di order sheet TIDAK dihitung dari persentase — ia dibaca
    apa adanya dari kolom nilai bersih. Menghitungnya ulang lewat
    `qty x harga x persen` meleset Rp1-2 dari baris TOTAL order sheet, dan
    kecocokan sampai rupiah terakhir itulah yang dijaga seluruh program ini.
    """
    wb, _, _ = buku
    ws = wb[rms.TAB_INVOICE]
    baris_rumus, _ = _sel(ws, 8)
    diskon = ws.cell(baris_rumus, 7).value
    assert not (isinstance(diskon, str) and diskon.startswith("=")), (
        "Nilai Diskon jadi rumus — total faktur tidak lagi sama persis "
        "dengan order sheet"
    )


def test_qty_surat_jalan_berupa_penjumlahan_kolom_ukuran(buku):
    """Berkas asli MTN menulisnya `=SUM(E21:J21)`."""
    wb, _, _ = buku
    ws = wb[rms.TAB_SURAT_JALAN]
    ketemu = any(
        isinstance(ws.cell(r, c).value, str)
        and str(ws.cell(r, c).value).startswith("=SUM(")
        for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)
    )
    assert ketemu, "kolom Qty Surat Jalan tidak memakai rumus"


# -------------------------------------------- data customer dikosongkan
def test_customer_yang_belum_diketahui_dibiarkan_KOSONG(buku):
    """Permintaan Yosua 20 Sep 2026: biar dia sendiri yang mengisi.

    Tulisan "(alamat belum diisi)" ikut tercetak ke dokumen yang dikirim ke
    customer — lebih buruk daripada sel kosong yang tinggal diketik.
    """
    wb, _, _ = buku
    for nama in (rms.TAB_INVOICE, rms.TAB_SURAT_JALAN):
        ws = wb[nama]
        teks = " ".join(
            str(ws.cell(r, c).value or "")
            for r in range(1, min(ws.max_row, 20) + 1)
            for c in range(1, min(ws.max_column, 12) + 1)
        )
        assert "belum diisi" not in teks, f"{nama} masih memuat teks penampung"


# ------------------------------------------------------- garis MTN
def test_kepala_dokumen_mtn_bergaris_seperti_berkas_asli(tmp_path, cfg):
    """Berkas asli MTN mengotaki blok CUSTOMER (A11:C15) dan tiap labelnya.

    Versi pertama melewatkannya sehingga kepala dokumen MTN tampil polos,
    dan Yosua menambahkan garisnya sendiri di Excel.
    """
    asli = cfg.customer
    try:
        wb, _ = _siapkan(tmp_path, cfg, "MTN")
    finally:
        cfg.customer = asli
    for nama in (rms.TAB_INVOICE, rms.TAB_SURAT_JALAN):
        ws = wb[nama]
        for r in range(11, 16):
            b = ws.cell(r, 1).border
            assert b.left.style, f"{nama} A{r} tidak bergaris kiri"
        assert ws.cell(11, 1).value == "CUSTOMER"
        assert ws.cell(11, 1).alignment.horizontal == "center"
