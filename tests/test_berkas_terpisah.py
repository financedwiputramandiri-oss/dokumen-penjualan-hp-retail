"""Tiap dokumen berkas SENDIRI, dengan rumus dan lembar Master Harga.

Permintaan Yosua 26 September 2026: *"MULAI SEKARANG DAN SETERUSNYA BUATLAH
FILE PROFORMA INVOICE, INVOICE DAN SURAT JALAN SECARA TERPISAH"*. Ini
membatalkan penggabungan Invoice + Surat Jalan yang ia minta 20 September;
tes pertama di bawah menjaga supaya tidak digabung lagi tanpa sengaja.

Sisanya tetap mengunci hal-hal yang kalau salah TIDAK kelihatan dari angka
totalnya: ukuran logo, garis kepala dokumen, dan rumus antar lembar.
"""
import pytest
from openpyxl import load_workbook

from buat_contoh import buat_contoh
from hp_dokumen.berkas_dokumen import buat_berkas, nama_aman
from hp_dokumen.dokumen import gaya
from hp_dokumen.dokumen import rumus as rms
from hp_dokumen.konfigurasi import Konfigurasi, Customer, DaftarCustomer
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet, master_harga_dari_buku


def _daftar(kode_pt: str, per_ukuran: bool = False) -> DaftarCustomer:
    # Nama & alamat SENGAJA dikosongkan: itu keadaan sebenarnya untuk
    # sebagian besar customer, dan Yosua sendiri yang akan mengisinya.
    bentuk = "per_ukuran" if per_ukuran else "per_artikel"
    return DaftarCustomer([
        Customer("Contoh TOP", "", "", "", bentuk, 30, "", kode_pt, ""),
    ])


@pytest.fixture(scope="module")
def cfg():
    return Konfigurasi.muat()


def _siapkan(tmp_path, cfg, kode_pt, per_ukuran=False):
    daftar = _daftar(kode_pt, per_ukuran)
    sumber = buat_contoh(tmp_path / "contoh.xlsx")
    order = next(o for o in baca_order_sheet(sumber, daftar, tahun_bawaan=2026)
                 if "TOP" in o.nama_tab)
    master = master_harga_dari_buku(load_workbook(sumber))
    asli = cfg.customer
    try:
        cfg.customer = daftar
        cust = daftar.cari(order.nama_tab)
        berkas = buat_berkas(cfg, order, tentukan_nett(order, cust),
                             tmp_path / "keluar", "001", master=master)
    finally:
        cfg.customer = asli
    akhiran = "_" + nama_aman(order.nama_tab) + ".xlsx"
    return ({p.name[:-len(akhiran)]: p for p in berkas
             if p.name.endswith(akhiran)}, master)


@pytest.fixture(params=["DPM", "MTN"])
def buku(request, tmp_path, cfg):
    berkas, master = _siapkan(tmp_path, cfg, request.param)
    return berkas, master, request.param


def _wb(berkas, nama):
    assert nama in berkas, f"berkas {nama} tidak dibuat: {sorted(berkas)}"
    return load_workbook(berkas[nama])


# --------------------------------------------------- berkas TERPISAH
def test_invoice_dan_surat_jalan_berkas_TERPISAH(buku):
    """Jangan pernah digabung lagi — permintaan Yosua 26 September 2026."""
    berkas, _, _ = buku
    assert "INVOICE" in berkas and "SURAT_JALAN" in berkas
    assert berkas["INVOICE"] != berkas["SURAT_JALAN"]
    assert not any(n.startswith("INVOICE_SURAT_JALAN") for n in berkas), (
        "Invoice dan Surat Jalan tergabung lagi dalam satu berkas"
    )
    assert rms.TAB_SURAT_JALAN not in _wb(berkas, "INVOICE").sheetnames


def test_proforma_berkas_sendiri(tmp_path, cfg):
    berkas, _ = _siapkan(tmp_path, cfg, "DPM", per_ukuran=True)
    assert "PROFORMA" in berkas
    for lain in ("INVOICE", "SURAT_JALAN"):
        assert berkas["PROFORMA"] != berkas[lain]


def test_master_harga_HANYA_di_berkas_invoice(buku):
    """Surat Jalan tidak memakai rumus harga, jadi tidak perlu daftar harga.

    Menyertakannya cuma membocorkan seluruh harga retail Happy Pumpkin ke
    satu customer tanpa guna apa pun.
    """
    berkas, _, _ = buku
    assert rms.TAB_MASTER in _wb(berkas, "INVOICE").sheetnames
    for nama in ("SURAT_JALAN", "PACKING_LIST", "FAKTUR_PAJAK"):
        if nama in berkas:
            assert rms.TAB_MASTER not in _wb(berkas, nama).sheetnames, nama


def test_lembar_master_harga_memuat_harga_saat_dokumen_dibuat(buku):
    """Harga retail berubah sepanjang tahun; faktur lama harus bawa buktinya."""
    berkas, master, _ = buku
    ws = _wb(berkas, "INVOICE")[rms.TAB_MASTER]
    kode = {ws.cell(r, 1).value for r in range(rms.BARIS_DATA_MASTER, ws.max_row + 1)}
    assert set(master) <= kode, "ada artikel master yang tidak ikut tersalin"
    assert ws.cell(rms.BARIS_TARIF, rms.KOLOM_TARIF).value is not None, "tarif PPN"


# ------------------------------------------------------------- logo
def _ukuran_logo(ws):
    assert ws._images, "logo hilang"
    a = ws._images[0].anchor
    return round(a.ext.cx / 9525), round(a.ext.cy / 9525)


def test_logo_sama_persis_di_invoice_dan_surat_jalan(buku):
    """Keluhan Yosua: logo berbeda ukuran antar dokumen.

    Sebabnya lebar kolom B menyempit mengikuti kode artikel yang pendek,
    sehingga logonya ikut mengecil di dokumen yang kodenya lebih pendek.
    """
    berkas, _, _ = buku
    inv = _ukuran_logo(_wb(berkas, "INVOICE")[rms.TAB_INVOICE])
    wb_sj = _wb(berkas, "SURAT_JALAN")
    sj = _ukuran_logo(wb_sj[wb_sj.sheetnames[0]])
    assert inv == sj, f"logo beda ukuran: invoice {inv}, surat jalan {sj}"


def test_logo_seukuran_berkas_asli_bukan_86_piksel_lagi(buku):
    """Diukur dari berkas asli: DPM +-119x120 px, MTN +-111x126 px."""
    berkas, _, kode = buku
    lebar, tinggi = _ukuran_logo(_wb(berkas, "INVOICE")[rms.TAB_INVOICE])
    assert lebar == pytest.approx(gaya.LEBAR_LOGO[kode], abs=1)
    assert tinggi > 100, f"logo masih kekecilan: {lebar}x{tinggi}"


def test_logo_tidak_pernah_gepeng(buku):
    """Tingginya mengikuti bentuk asli gambar, bukan angka tetap."""
    berkas, _, _ = buku
    img = _wb(berkas, "INVOICE")[rms.TAB_INVOICE]._images[0]
    a = img.anchor
    rasio_pasang = (a.ext.cx / 9525) / (a.ext.cy / 9525)
    assert rasio_pasang == pytest.approx(img.width / img.height, rel=0.02)


# ------------------------------------------------------------ rumus
def _sel(ws, kolom, awal=1):
    for r in range(awal, ws.max_row + 1):
        nilai = ws.cell(r, kolom).value
        if isinstance(nilai, str) and nilai.startswith("="):
            return r, nilai
    return None, None


def test_harga_satuan_diambil_lewat_vlookup_ke_master(buku):
    berkas, _, _ = buku
    _, isi = _sel(_wb(berkas, "INVOICE")[rms.TAB_INVOICE], 5)
    assert isi and "VLOOKUP" in isi and rms.TAB_MASTER in isi


def test_jumlah_per_baris_berupa_rumus_qty_kali_harga(buku):
    berkas, _, _ = buku
    _, isi = _sel(_wb(berkas, "INVOICE")[rms.TAB_INVOICE], 8)
    assert isi and "*" in isi, f"kolom Jumlah bukan rumus: {isi!r}"


def test_penutup_memakai_rumus_bukan_angka_mati(buku):
    berkas, _, _ = buku
    ws = _wb(berkas, "INVOICE")[rms.TAB_INVOICE]
    semua = [str(ws.cell(r, 8).value or "") for r in range(1, ws.max_row + 1)]
    assert any("SUMPRODUCT" in x for x in semua), "Subtotal bukan rumus"
    assert any(x.startswith("=SUM(") for x in semua), "Diskon bukan rumus"


def test_nilai_diskon_per_baris_TETAP_angka_bukan_rumus(buku):
    """Sengaja, dan jangan diubah tanpa membaca dokumen/rumus.py.

    Nilai bersih di order sheet TIDAK dihitung dari persentase — ia dibaca
    apa adanya dari kolom nilai bersih. Menghitungnya ulang lewat
    `qty x harga x persen` meleset Rp1-2 dari baris TOTAL order sheet, dan
    kecocokan sampai rupiah terakhir itulah yang dijaga seluruh program ini.
    """
    berkas, _, _ = buku
    ws = _wb(berkas, "INVOICE")[rms.TAB_INVOICE]
    baris_rumus, _ = _sel(ws, 8)
    diskon = ws.cell(baris_rumus, 7).value
    assert not (isinstance(diskon, str) and diskon.startswith("=")), (
        "Nilai Diskon jadi rumus — total faktur tidak lagi sama persis "
        "dengan order sheet"
    )


def test_qty_surat_jalan_berupa_penjumlahan_kolom_ukuran(buku):
    """Berkas asli MTN menulisnya `=SUM(E21:J21)`."""
    berkas, _, _ = buku
    wb = _wb(berkas, "SURAT_JALAN")
    ws = wb[wb.sheetnames[0]]
    ketemu = any(
        isinstance(ws.cell(r, c).value, str)
        and str(ws.cell(r, c).value).startswith("=SUM(")
        for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)
    )
    assert ketemu, "kolom Qty Surat Jalan tidak memakai rumus"


# -------------------------------------------- data customer dikosongkan
def test_customer_yang_belum_diketahui_dibiarkan_KOSONG(buku):
    """Permintaan Yosua 20 Sep 2026: biar dia sendiri yang mengisi.

    Tulisan "(alamat belum diisi)" ikut tercetak ke dokumen yang dikirim ke
    customer — lebih buruk daripada sel kosong yang tinggal diketik.
    """
    berkas, _, _ = buku
    for nama in ("INVOICE", "SURAT_JALAN"):
        wb = _wb(berkas, nama)
        ws = wb[wb.sheetnames[0]]
        teks = " ".join(
            str(ws.cell(r, c).value or "")
            for r in range(1, min(ws.max_row, 20) + 1)
            for c in range(1, min(ws.max_column, 12) + 1)
        )
        assert "belum diisi" not in teks, f"{nama} masih memuat teks penampung"


# ------------------------------------------------------- garis MTN
def test_kepala_dokumen_mtn_bergaris_seperti_berkas_asli(tmp_path, cfg):
    """Berkas asli MTN mengotaki blok CUSTOMER (A11:C15) dan tiap labelnya."""
    berkas, _ = _siapkan(tmp_path, cfg, "MTN")
    for nama in ("INVOICE", "SURAT_JALAN"):
        wb = _wb(berkas, nama)
        ws = wb[wb.sheetnames[0]]
        for r in range(11, 16):
            assert ws.cell(r, 1).border.left.style, f"{nama} A{r} tidak bergaris kiri"
        assert ws.cell(11, 1).value == "CUSTOMER"
        assert ws.cell(11, 1).alignment.horizontal == "center"
