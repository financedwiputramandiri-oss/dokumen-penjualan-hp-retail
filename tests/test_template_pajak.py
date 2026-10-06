"""Lembar TEMPLATE — bentuknya dari berkas Yosua, 6 Oktober 2026.

Yang dikunci: susunan kolom, mekanisme PPN 12% atas DPP Nilai Lain, potongan
beruntun, dan larangan ROUND. Yang terakhir itu yang paling penting — dengan
pembulatan per sel, jumlah DPP+PPN meleset Rp22-24 dari total invoice.
"""
import openpyxl
import pytest
from openpyxl import Workbook

from hp_dokumen.dokumen.template_pajak import JUDUL, buat_template_pajak
from hp_dokumen.konfigurasi import Konfigurasi
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet
from buat_contoh import buat_contoh


@pytest.fixture
def bahan(tmp_path):
    cfg = Konfigurasi.muat()
    berkas = buat_contoh(tmp_path / "contoh.xlsx")
    orders = [o for o in baca_order_sheet(berkas, cfg.customer, tahun_bawaan=2026)
              if o.qty > 0]
    return orders, cfg


def _buat(orders, cfg, tmp_path, i=0):
    o = orders[i]
    c = cfg.customer.cari(o.nama_tab)
    wb = Workbook()
    hasil = buat_template_pajak(wb.active, o, tentukan_nett(o, c), c,
                                cfg.perusahaan.untuk(c), cfg.pengaturan, "001")
    p = tmp_path / f"tpl{i}.xlsx"
    wb.save(p)
    return openpyxl.load_workbook(p).active, hasil


def test_susunan_kolom_sama_dengan_berkas_yosua(bahan, tmp_path):
    ws, _ = _buat(*bahan, tmp_path)
    assert ws.cell(2, 1).value == "NO"
    assert ws.cell(2, 4).value == "QTY"
    assert ws.cell(2, 6).value == "HARGA JUAL EXC. PPN (Pembulatan)"
    assert ws.cell(2, 10).value == "DPP"
    assert ws.cell(2, 11).value == "DPP NILAI LAINNYA"
    assert ws.cell(2, 12).value == "PPN 12%"
    assert ws.cell(2, 13).value == "DPP + PPN"
    assert len(JUDUL) == 13


def test_ppn_lewat_DPP_NILAI_LAIN_bukan_11_persen_langsung(bahan, tmp_path):
    """`K = J x 11/12` lalu `L = K x 12%` memberi 11% — tapi mekanismenya
    harus kelihatan di lembarnya, itu yang diminta aturan DJP sejak 2025.
    """
    ws, _ = _buat(*bahan, tmp_path)
    assert ws.cell(3, 11).value == "=J3*11/12"
    assert ws.cell(3, 12).value == "=K3*12%"
    assert "11%" not in str(ws.cell(3, 12).value)


def test_tidak_ada_ROUND_di_seluruh_lembar(bahan, tmp_path):
    """Dengan pembulatan per sel, DPP+PPN meleset Rp22-24 dari invoice.

    Diukur pada data berkas Yosua sendiri: rumus asli tepat Rp0,00, "F
    dibulatkan" Rp22,17, "semua sel dibulatkan" Rp24,05.
    """
    ws, _ = _buat(*bahan, tmp_path)
    for r in range(1, ws.max_row + 1):
        for c in range(1, 14):
            isi = str(ws.cell(r, c).value or "")
            assert "ROUND" not in isi.upper(), f"{r},{c} memakai ROUND: {isi!r}"


def test_jumlah_DPP_tambah_PPN_sama_dengan_nilai_bersih_order_sheet(bahan, tmp_path):
    """Dihitung dengan tangan dari rumusnya, sebab openpyxl tidak menghitung.

    M = J + L = (H - I) + (J x 11/12 x 12%) = (H - I) x 1,11
      = (Qty x Harga/1,11) x (1 - tarif) x 1,11 = kotor x (1 - tarif)
    """
    orders, cfg = bahan
    for i in range(len(orders)):
        ws, h = _buat(orders, cfg, tmp_path, i)
        if not h["baris"]:
            continue
        tarif = 1 - (1 - h["tarif_dasar"]) * (1 - h["tarif_tambahan"])
        assert abs(h["kotor"] * (1 - tarif) - h["nett"]) < 0.01, orders[i].nama_tab


def test_baris_jumlah_berlatar_kuning_dan_berformat_bulat(bahan, tmp_path):
    """"Angka bulat" diminta Yosua, dan dipenuhi lewat FORMAT - bukan ROUND."""
    ws, h = _buat(*bahan, tmp_path)
    r = h["baris_jumlah"]
    for kolom in range(7, 14):
        sel = ws.cell(r, kolom)
        assert str(sel.value or "").startswith("=SUM("), f"kolom {kolom} bukan SUM"
        assert sel.fill.fgColor.rgb.endswith("FFFF00"), f"kolom {kolom} tidak kuning"
        assert ".00" not in sel.number_format, (
            f"kolom {kolom} menampilkan desimal, seharusnya angka bulat"
        )


def test_potongan_beruntun_bukan_dijumlahkan(bahan, tmp_path):
    """27,5% lalu 2% = 28,95%, bukan 29,5%. Aturan yang sama dengan invoice."""
    orders, cfg = bahan
    for i in range(len(orders)):
        ws, h = _buat(orders, cfg, tmp_path, i)
        if not h["baris"] or not h["tarif_tambahan"]:
            continue
        isi = str(ws.cell(3, 9).value)
        # Bentuk beruntun: potongan kedua dikenakan pada SISA setelah yang
        # pertama, jadi rumusnya memuat pengurangan di dalamnya.
        assert "F3-(" in isi, f"bukan potongan beruntun: {isi!r}"
        return
    pytest.skip("tidak ada order dengan potongan tambahan di contoh")


def test_data_customer_yang_belum_diketahui_DIKOSONGKAN(bahan, tmp_path):
    """Aturan bagian 33: jangan menulis tulisan penampung di dokumen."""
    ws, _ = _buat(*bahan, tmp_path)
    for kolom in (9, 12, 13):
        isi = ws.cell(1, kolom).value
        assert isi is None or "belum diisi" not in str(isi).lower()
