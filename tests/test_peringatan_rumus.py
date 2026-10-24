"""Peringatan untuk rumus order sheet yang salah — permintaan Yosua 20 Okt 2026.

Cacat rumus TIDAK PERNAH kelihatan dari angkanya, jadi tes di sini membuat
order sheet kecil yang rumusnya memang salah lalu memastikan peringatannya
muncul — dan yang tidak salah TIDAK memunculkan peringatan.
"""
import openpyxl
import pytest

from hp_dokumen.periksa_rumus import baris_tercakup, periksa, rentang_rumus
from hp_dokumen.pemindai import pindai_tab
from hp_dokumen.tata_letak import kenali

UKURAN = ["NB", "S", "M", "L", "XL", "2", "3", "4", "9"]


def _tab(wb, nama, rumus_qty_sampai="V", rumus_total=None, baris=3):
    """Satu tab PO 9 kolom ukuran, N..V, TOTAL di W.

    `rumus_qty_sampai` kolom terakhir yang ikut dijumlah rumus qty.
    """
    ws = wb.create_sheet(nama)
    ws.cell(1, 1, "ARTICLE CODE")
    ws.cell(1, 2, "PRODUCT NAME")
    ws.cell(1, 3, "COLOUR")
    ws.cell(1, 4, "ORIGINAL PO")
    ws.cell(1, 14, "AVAILABLE TO ORDER")
    ws.cell(2, 13, "TOTAL")
    ws.cell(2, 23, "TOTAL")
    for i, u in enumerate(UKURAN):
        ws.cell(2, 4 + i, u)
        ws.cell(2, 14 + i, u)
    ws.cell(1, 24, "PRICE W/ VAT")
    ws.cell(1, 26, "TOTAL ATO")
    ws.cell(1, 27, "LOSSES")
    ws.cell(1, 28, "DISC")
    ws.cell(1, 29, "TOTAL VALUE")
    ws.cell(1, 30, "DISCOUNT CBD + 1,5%")
    return ws


def _isi(ws, r, kode, qty_per_kolom, harga=10000):
    ws.cell(r, 1, kode)
    ws.cell(r, 2, "Barang Uji")
    ws.cell(r, 3, "Merah")
    for kolom, q in qty_per_kolom.items():
        ws.cell(r, kolom, q)
    ws.cell(r, 24, harga)


def _buku(rumus_qty_sampai=22, qty_kolom_akhir=0):
    """Dua buku: satu berisi NILAI, satu berisi RUMUS, isinya sepadan."""
    nilai = openpyxl.Workbook()
    nilai.remove(nilai.active)
    rumus = openpyxl.Workbook()
    rumus.remove(rumus.active)
    for wb in (nilai, rumus):
        _tab(wb, "PO 01 Jan - Uji")
    wn, wr = nilai["PO 01 Jan - Uji"], rumus["PO 01 Jan - Uji"]
    for i, r in enumerate(range(3, 6)):
        q = {14: 5}
        if qty_kolom_akhir:
            q[22] = qty_kolom_akhir          # kolom V, ukuran ke-9
        _isi(wn, r, f"KODE.{i}", q)
        _isi(wr, r, f"KODE.{i}", q)
        total = 5 + (qty_kolom_akhir if rumus_qty_sampai >= 22 else 0)
        wn.cell(r, 23, total)
        wn.cell(r, 26, total * 10000)
        wn.cell(r, 29, total * 10000)
        wr.cell(r, 23, f"=SUM(N{r}:{chr(64+rumus_qty_sampai-26)}{r})"
                if rumus_qty_sampai > 26 else
                f"=SUM(N{r}:{openpyxl.utils.get_column_letter(rumus_qty_sampai)}{r})")
    return wn, wr


def _blok(wn):
    blok, _, _ = pindai_tab(wn)
    return blok


def test_rumus_qty_pendek_yang_menelan_qty_diperingatkan_GENTING():
    wn, wr = _buku(rumus_qty_sampai=21, qty_kolom_akhir=3)   # berhenti di U
    pesan = periksa(wn, wr, _blok(wn))
    genting = [p for p in pesan if "GENTING" in p and "rumus qty" in p]
    assert genting, pesan
    assert "9 pcs" in genting[0]          # 3 baris x 3 pcs
    assert "kolom V" in genting[0]


def test_rumus_qty_pendek_tapi_kolomnya_kosong_hanya_POTENSI():
    wn, wr = _buku(rumus_qty_sampai=21, qty_kolom_akhir=0)
    pesan = periksa(wn, wr, _blok(wn))
    assert not [p for p in pesan if "GENTING" in p], pesan
    ringan = [p for p in pesan if p.startswith("RUMUS KURANG TEPAT")]
    assert ringan, pesan
    assert "HILANG DIAM-DIAM" in ringan[0]


def test_rumus_qty_yang_benar_tidak_memunculkan_peringatan():
    wn, wr = _buku(rumus_qty_sampai=22, qty_kolom_akhir=3)
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "rumus qty" in p]
    assert pesan == []


def test_baris_total_yang_melewatkan_baris_diperingatkan():
    wn, wr = _buku(rumus_qty_sampai=22)
    wn.cell(10, 23, 10)
    wr.cell(10, 23, "=SUM(W3:W4)")        # baris 5 terlewat
    pesan = periksa(wn, wr, _blok(wn))
    total = [p for p in pesan if "baris TOTAL W10" in p]
    assert total, pesan
    assert "GENTING" in total[0]


def test_baris_total_bersel_satuan_tidak_dianggap_melewatkan():
    """`=sum(W3,W4,W5)` menjumlah semuanya walau tanpa rentang.

    Yang hanya membaca rentang akan memberi alarm palsu — dan alarm palsu
    membuat orang berhenti membaca peringatannya (CLAUDE.md bagian 14).
    """
    wn, wr = _buku(rumus_qty_sampai=22)
    wn.cell(10, 23, 15)
    wr.cell(10, 23, "=sum(W3,W4,W5)")
    assert [p for p in periksa(wn, wr, _blok(wn)) if "baris TOTAL" in p] == []


def test_galat_di_kolom_nilai_bersih_baris_data_diperingatkan_GENTING():
    wn, wr = _buku(rumus_qty_sampai=22)
    wn.cell(4, 29, "#REF!")
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "galat" in p]
    assert pesan and "GENTING" in pesan[0], pesan
    assert "AC4" in pesan[0]


def test_galat_di_baris_TOTAL_tidak_disebut_genting():
    """Dokumennya tetap benar — nilai bersih dibaca PER BARIS, bukan dari
    baris TOTAL. Menyebutnya genting akan membuat 266 tab berteriak percuma."""
    wn, wr = _buku(rumus_qty_sampai=22)
    wn.cell(10, 29, "#REF!")
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "galat" in p]
    assert pesan, "galat di baris TOTAL tetap harus disebut"
    assert "GENTING" not in pesan[0]


def test_tanpa_lembar_rumus_tidak_meledak():
    wn, _ = _buku(rumus_qty_sampai=21, qty_kolom_akhir=3)
    assert isinstance(periksa(wn, None, _blok(wn)), list)


@pytest.mark.parametrize("rumus,kolom,harap", [
    ("=SUM(W3:W5)", 23, {3, 4, 5}),
    ("=sum(W3,W7,W11:W12)", 23, {3, 7, 11, 12}),
    ("=SUM(N3:V3)", 23, set()),
])
def test_baris_tercakup_membaca_rentang_dan_sel_satuan(rumus, kolom, harap):
    assert baris_tercakup(rumus, kolom) == harap


def test_rentang_rumus_membaca_koordinat_dengan_benar():
    assert rentang_rumus("=SUM(N3:V3)") == [(14, 3, 22, 3)]
    assert rentang_rumus("=SUM($W$3:$W$9)") == [(23, 3, 23, 9)]


# --------------------------------------------------- harga & diskon
def _pasang_diskon(wn, wr, persen_rumus, persen_disc, harga_rumus=True):
    """Isi kolom DISC (AB=28) dan TOTAL VALUE (AC=29) tiga baris."""
    for r in range(3, 6):
        kotor = 50000
        wn.cell(r, 26, kotor)
        wr.cell(r, 26, f"=X{r}*W{r}")
        wn.cell(r, 28, persen_disc)
        wr.cell(r, 28, persen_disc)
        wn.cell(r, 29, kotor * (1 - persen_rumus))
        wr.cell(r, 29, f"=Z{r}-(Z{r}*{persen_rumus*100:g}%)")
        if harga_rumus:
            wr.cell(r, 24, f"=VLOOKUP(A{r},'Harga Retail'!$A$1:$C$99,3,0)")


def test_kolom_disc_beda_dari_persen_di_rumus_diperingatkan_GENTING():
    """Persen di kolom DISC tidak menunjuk rumusnya, jadi keduanya bisa
    berbeda tanpa tanda apa pun. Kasus nyata: Maret 2026 Katamama Cikaret,
    DISC tertulis 25% tapi rumus AC memakai 22%."""
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, persen_rumus=0.22, persen_disc=0.25)
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "kolom DISC" in p]
    assert pesan, periksa(wn, wr, _blok(wn))
    assert "GENTING" in pesan[0]
    assert "25" in pesan[0] and "22" in pesan[0]


def test_kolom_disc_berisi_TEKS_persen_tetap_diperiksa():
    """Kolom DISC kadang teks ("25%"), kadang angka (0.25).

    Versi pertama pemeriksa ini hanya membaca angka, jadi tab ber-DISC teks
    tidak pernah diperiksa dan hasilnya terlihat bersih padahal
    pemeriksaannya memang tidak jalan.
    """
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, persen_rumus=0.22, persen_disc="25%")
    assert [p for p in periksa(wn, wr, _blok(wn)) if "kolom DISC" in p]


def test_kolom_disc_yang_cocok_tidak_memunculkan_peringatan():
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, persen_rumus=0.25, persen_disc=0.25)
    assert [p for p in periksa(wn, wr, _blok(wn)) if "kolom DISC" in p] == []


def test_persen_diskon_tidak_seragam_dalam_satu_tab_diperingatkan():
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, persen_rumus=0.25, persen_disc=0.25)
    wr.cell(5, 29, "=Z5-(Z5*20%)")      # satu baris beda sendiri
    wn.cell(5, 29, 40000)
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "tidak seragam" in p]
    assert pesan and "GENTING" in pesan[0]


def test_harga_diketik_tangan_diperingatkan():
    """Harga yang diketik tidak ikut berubah kalau Harga Retail diperbarui —
    melanggar KETETAPAN bagian 41."""
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, 0.25, 0.25, harga_rumus=False)
    for r in range(3, 6):
        wn.cell(r, 24, 51000)           # angka, bukan rumus
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "DIKETIK" in p]
    assert pesan, periksa(wn, wr, _blok(wn))
    assert "Harga Retail" in pesan[0]


def test_harga_VLOOKUP_tidak_memunculkan_peringatan():
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, 0.25, 0.25, harga_rumus=True)
    assert [p for p in periksa(wn, wr, _blok(wn)) if "DIKETIK" in p] == []


@pytest.mark.parametrize("rumus,harap", [
    ("=Z3-(Z3*25%)", 0.25),
    ("=Z3-(Z3*1,5%)", 0.015),
    ("=Z3*0.75", 0.25),
    ("=Z3", None),
])
def test_persen_dalam_rumus_terbaca(rumus, harap):
    from hp_dokumen.periksa_rumus import persen_dalam_rumus
    assert persen_dalam_rumus(rumus) == harap


# --------------------------------------------------- tanda % hilang
def test_tanda_persen_hilang_di_rumus_diskon_diperingatkan_GENTING():
    """`=AC3-(AC3*1.5)` bukan `*1.5%` — hasilnya AC x (-0,5), jadi NEGATIF.

    Ditemukan 21 Oktober 2026 di 249 sel pada 16 tab. Satu di antaranya
    sudah membuat invoice kurang menagih Rp999.033 dan pencocokan tetap
    LULUS, sebab order sheet dan program sama-sama membaca angka yang
    sama salahnya.
    """
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, persen_rumus=0.25, persen_disc=0.25)
    for r in range(3, 6):
        wr.cell(r, 30, f"=AC{r}-(AC{r}*1.5)")   # % HILANG
        wn.cell(r, 30, wn.cell(r, 29).value * -0.5)
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "TANDA % HILANG" in p]
    assert pesan, periksa(wn, wr, _blok(wn))
    assert "GENTING" in pesan[0]
    assert "NEGATIF" in pesan[0]


def test_rumus_diskon_memakai_pengali_pecahan_TIDAK_diperingatkan():
    """`=AC3*0,985` memang bentuk yang benar — jangan dianggap cacat."""
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, persen_rumus=0.25, persen_disc=0.25)
    for r in range(3, 6):
        wr.cell(r, 30, f"=AC{r}*0.985")
        wn.cell(r, 30, wn.cell(r, 29).value * 0.985)
    assert [p for p in periksa(wn, wr, _blok(wn)) if "TANDA % HILANG" in p] == []


def test_nilai_bersih_negatif_selalu_diperingatkan():
    """Apa pun sebabnya, nilai bersih negatif tidak pernah benar."""
    wn, wr = _buku(rumus_qty_sampai=22)
    _pasang_diskon(wn, wr, persen_rumus=0.25, persen_disc=0.25)
    for r in range(3, 6):
        wn.cell(r, 30, -100000)      # negatif, tanpa rumus bercacat
    pesan = [p for p in periksa(wn, wr, _blok(wn)) if "NEGATIF" in p]
    assert pesan and "GENTING" in pesan[0]
