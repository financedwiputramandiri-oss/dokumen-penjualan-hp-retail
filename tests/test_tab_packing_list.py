"""Tab PACKING LIST di order sheet: hanya qty, tanpa harga dan tanpa nilai.

Dibuktikan dari tab `Packing List - Dunia Bayi` pada order sheet Oktober 2026:
16 kolom, satu blok berjudul "AVAILABLE TO ORDER (QTY)", tanpa PRICE W/ VAT.
Catatan lama mengira tab semacam ini berjudul ORIGINAL PO (Packing List
Haritsa), jadi dukungannya tidak pernah benar-benar jalan — pemindai menolak
tabnya dengan "susunan kolomnya tidak dikenali".
"""
from openpyxl import Workbook

from hp_dokumen.konfigurasi import pecah_nama_tab
from hp_dokumen.pemindai import pindai_tab


def _tab_packing_list():
    wb = Workbook()
    ws = wb.active
    ws.title = "Packing List - Dunia Bayi"
    ws.cell(1, 1, "ARTICLE \nCODE")
    ws.cell(1, 2, "PRODUCT NAME")
    ws.cell(1, 3, "COLOUR")
    ws.cell(1, 4, "AVAILABLE TO ORDER (QTY)")
    for i, label in enumerate(["1", "2", "3", "4", "5", "6"]):
        ws.cell(2, 4 + i, label)
    ws.cell(2, 13, "TOTAL")
    ws.cell(3, 1, "73071")
    ws.cell(3, 2, "Coco Tennis Dress Small Size")
    ws.cell(3, 3, "Lemon")
    ws.cell(3, 4, 3)
    ws.cell(3, 5, 6)
    ws.cell(3, 6, 2)
    ws.cell(3, 13, 11)
    ws.cell(4, 13, 11)      # baris TOTAL milik sheet
    return ws


def test_tab_packing_list_terbaca_walau_tanpa_kolom_harga():
    blok, total, _ = pindai_tab(_tab_packing_list(), pakai_kolom_ori=True)
    assert len(blok) == 1, "tab Packing List harus terbaca, bukan dilewati"
    baris = blok[0].baris
    assert len(baris) == 1
    assert baris[0].qty == 11
    assert baris[0].qty_per_ukuran[:3] == [3, 6, 2]
    assert blok[0].label_ukuran[:3] == ["1", "2", "3"]
    # Baris TOTAL hanya memuat qty: tidak ada nilai kotor maupun TOTAL VALUE.
    assert total.qty == 11
    assert total.nilai_kotor is None
    assert total.total_value is None


def test_tab_packing_list_ditandai_hanya_qty():
    blok, _, _ = pindai_tab(_tab_packing_list(), pakai_kolom_ori=True)
    assert blok[0].tata.hanya_qty is True, (
        "tanpa penanda ini, Invoice dan Faktur Pajak ikut terbit dari tab "
        "yang tidak punya harga sama sekali — dokumen penagihan bernilai nol"
    )


def test_awalan_packing_list_dibuang_saat_mencocokkan_customer():
    # Tanpa ini kop dokumennya terbit TANPA nama dan alamat customer, dan
    # itu tidak kelihatan sama sekali dari membaca nilai sel.
    _, sisa = pecah_nama_tab("Packing List - Dunia Bayi")
    assert sisa.strip() == "Dunia Bayi"
    # Tab PO biasa tidak boleh ikut berubah.
    _, sisa_po = pecah_nama_tab("PO 08 Okt - Dunia Bayi")
    assert sisa_po.strip() == "Dunia Bayi"
