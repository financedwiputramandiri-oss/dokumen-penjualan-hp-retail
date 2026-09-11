"""Tes alarm bot penyapu: perubahan apa yang harus berbunyi, apa yang tidak."""
from pathlib import Path

import openpyxl
import pytest

from buat_contoh import buat_contoh
from hp_dokumen.konfigurasi import Customer, DaftarCustomer
from hp_dokumen.nilai_bersih import tentukan_nett
from hp_dokumen.pemindai import baca_order_sheet
from hp_dokumen.sapu.kondisi import Kondisi, SidikPO, sidik_dari_order
from hp_dokumen.sapu.pantau import GENTING, KABAR, PERHATIAN, bandingkan


@pytest.fixture
def daftar():
    return DaftarCustomer([
        Customer("Contoh TOP", "PT A", "x", "", "per_artikel", 30, "", "", ""),
        Customer("Contoh CBD", "PT B", "y", "", "per_ukuran", 30, "", "", ""),
    ])


def _sidik(berkas: Path, daftar, tab_mengandung="TOP", rumus=None) -> SidikPO:
    orders = baca_order_sheet(berkas, daftar, tahun_bawaan=2026)
    o = next(x for x in orders if tab_mengandung in x.nama_tab)
    k = tentukan_nett(o, daftar.cari(o.nama_tab))
    return sidik_dari_order("sheet1", "Order Sheet Contoh", o.nama_tab, o, k, rumus)


@pytest.fixture
def berkas(tmp_path):
    return buat_contoh(tmp_path / "contoh.xlsx")


# ------------------------------------------------------------ tidak berbunyi
def test_tidak_ada_perubahan_tidak_berbunyi(berkas, daftar):
    a = _sidik(berkas, daftar)
    b = _sidik(berkas, daftar)
    assert bandingkan(a, b) == []


def test_po_baru_hanya_kabar(berkas, daftar):
    b = _sidik(berkas, daftar)
    hasil = bandingkan(None, b)
    assert len(hasil) == 1
    assert hasil[0].tingkat == KABAR
    assert not hasil[0].genting


def test_ato_baru_terisi_bukan_masalah(berkas, daftar):
    """PO yang ATO-nya baru diisi itu yang ditunggu, bukan alarm."""
    b = _sidik(berkas, daftar)
    kosong = SidikPO(b.id_sheet, b.nama_sheet, b.tab, ato_terisi=False,
                     baris=0, qty=0, kotor=0, nett=0, cara_bayar="TOP")
    hasil = bandingkan(kosong, b)
    assert len(hasil) == 1
    assert hasil[0].tingkat == PERHATIAN
    assert not hasil[0].genting


# ------------------------------------------------------------ harus berbunyi
def test_qty_berubah_berbunyi_genting(berkas, daftar, tmp_path):
    a = _sidik(berkas, daftar)
    wb = openpyxl.load_workbook(berkas)
    ws = wb["PO 05 Januari - Contoh TOP"]
    ws.cell(3, 15, (ws.cell(3, 15).value or 0) + 7)   # tambah qty di kolom ATO
    ubah = tmp_path / "ubah.xlsx"
    wb.save(ubah)
    b = _sidik(ubah, daftar)

    hasil = bandingkan(a, b)
    jenis = {x.jenis for x in hasil}
    assert "Qty berubah" in jenis
    assert any(x.genting for x in hasil)


def test_susunan_qty_berubah_walau_total_sama(berkas, daftar, tmp_path):
    """Total tetap, tapi qty dipindah antar ukuran -> tetap harus berbunyi."""
    a = _sidik(berkas, daftar)
    wb = openpyxl.load_workbook(berkas)
    ws = wb["PO 05 Januari - Contoh TOP"]
    ws.cell(3, 15, (ws.cell(3, 15).value or 0) + 1)
    ws.cell(3, 16, (ws.cell(3, 16).value or 0) - 1)
    ubah = tmp_path / "geser.xlsx"
    wb.save(ubah)
    b = _sidik(ubah, daftar)

    assert b.qty == a.qty, "totalnya memang harus tetap sama"
    hasil = bandingkan(a, b)
    assert "Susunan qty berubah" in {x.jenis for x in hasil}
    assert any(x.genting for x in hasil)


def test_nilai_bersih_berubah_berbunyi(berkas, daftar, tmp_path):
    a = _sidik(berkas, daftar)
    wb = openpyxl.load_workbook(berkas)
    ws = wb["PO 05 Januari - Contoh TOP"]
    ws.cell(3, 29, (ws.cell(3, 29).value or 0) - 50000)   # ubah TOTAL VALUE
    ubah = tmp_path / "nilai.xlsx"
    wb.save(ubah)
    b = _sidik(ubah, daftar)

    jenis = {x.jenis for x in bandingkan(a, b)}
    assert "Nilai bersih berubah" in jenis


def test_rumus_berubah_berbunyi_walau_angka_sama(berkas, daftar):
    """Rumus diubah tapi hasilnya masih sama -> tetap harus dilaporkan."""
    a = _sidik(berkas, daftar, rumus=[["=A1*2", "=B1"], ["=C1"]])
    b = _sidik(berkas, daftar, rumus=[["=A1*3", "=B1"], ["=C1"]])
    assert a.qty == b.qty and a.nett == b.nett
    hasil = bandingkan(a, b)
    assert "Rumus berubah" in {x.jenis for x in hasil}
    assert any(x.genting for x in hasil)


def test_ato_dikosongkan_berbunyi(berkas, daftar):
    a = _sidik(berkas, daftar)
    kosong = SidikPO(a.id_sheet, a.nama_sheet, a.tab, ato_terisi=False,
                     baris=0, qty=0, kotor=0, nett=0, cara_bayar="TOP")
    hasil = bandingkan(a, kosong)
    assert "ATO dikosongkan" in {x.jenis for x in hasil}
    assert any(x.genting for x in hasil)


def test_cara_bayar_berubah_berbunyi(berkas, daftar):
    a = _sidik(berkas, daftar)
    b = _sidik(berkas, daftar)
    b.cara_bayar = "CBD"
    hasil = bandingkan(a, b)
    assert "Cara bayar berubah" in {x.jenis for x in hasil}


# ------------------------------------------------------------ penyimpanan
def test_kondisi_tersimpan_dan_terbaca(berkas, daftar, tmp_path):
    a = _sidik(berkas, daftar)
    berkas_kondisi = tmp_path / "kondisi.json"
    k = Kondisi()
    k.pasang(a)
    k.simpan(berkas_kondisi)

    k2 = Kondisi.muat(berkas_kondisi)
    kembali = k2.ambil(a.kunci)
    assert kembali is not None
    assert kembali.qty == a.qty
    assert kembali.sidik_qty == a.sidik_qty
    assert bandingkan(kembali, a) == []


def test_kondisi_hilang_tidak_bikin_error(tmp_path):
    assert Kondisi.muat(tmp_path / "belum_ada.json").po == {}


def test_nilai_nol_karena_gagal_baca_bukan_alarm_palsu(berkas, daftar):
    """Kalau rumus tidak ikut terbaca, nilainya jadi nol. Itu masalah baca,
    bukan angka yang diubah orang — tidak boleh dilaporkan sebagai GENTING."""
    a = _sidik(berkas, daftar)
    b = _sidik(berkas, daftar)
    b.kotor = 0.0
    b.nett = 0.0
    hasil = bandingkan(a, b)
    assert [x.jenis for x in hasil] == ["Nilai tidak terbaca"]
    assert not any(x.genting for x in hasil)


def test_qty_tetap_dibandingkan_walau_nilai_tidak_terbaca(berkas, daftar, tmp_path):
    """Nilai gagal terbaca tidak boleh menutupi perubahan qty yang sungguhan."""
    a = _sidik(berkas, daftar)
    wb = openpyxl.load_workbook(berkas)
    ws = wb["PO 05 Januari - Contoh TOP"]
    ws.cell(3, 15, (ws.cell(3, 15).value or 0) + 1)
    ws.cell(3, 16, (ws.cell(3, 16).value or 0) - 1)
    ubah = tmp_path / "geser2.xlsx"
    wb.save(ubah)
    b = _sidik(ubah, daftar)
    b.kotor = 0.0
    b.nett = 0.0

    jenis = {x.jenis for x in bandingkan(a, b)}
    assert "Susunan qty berubah" in jenis


# ------------------------------------------- pengaman menulis ke sheet Yosua
class _SambunganPalsu:
    """Sambungan tiruan, untuk menguji tanpa menyentuh Google beneran."""

    def __init__(self, judul):
        self._judul = list(judul)
        self.ditulis = {}
        self.dikosongkan = []
        self.tab_dibuat = []
        sambungan = self

        class _Values:
            def clear(self, spreadsheetId, range, body):
                sambungan.dikosongkan.append(range)
                return self

            def update(self, spreadsheetId, range, valueInputOption, body):
                sambungan.ditulis[range] = body["values"]
                return self

            def execute(self):
                return {}

        class _Spreadsheets:
            def values(self_inner):
                return _Values()

            def batchUpdate(self_inner, spreadsheetId, body):
                for r in body["requests"]:
                    judul = r["addSheet"]["properties"]["title"]
                    sambungan.tab_dibuat.append(judul)
                    sambungan._judul.append(judul)
                return _Values()

        class _Sheets:
            def spreadsheets(self_inner):
                return _Spreadsheets()

        self.sheets = _Sheets()

    def nama_tab(self, id_sheet):
        return list(self._judul)


def test_bot_menolak_menulis_ke_tab_buatan_manusia():
    """Tab PENGATURAN, MASTER_CUSTOMER, dan lainnya tidak boleh disentuh bot."""
    from hp_dokumen.sapu.tulis_sheet import PenulisSheet

    s = _SambunganPalsu(["PENGATURAN", "MASTER_CUSTOMER", "INVOICE"])
    p = PenulisSheet(s, "id")
    p.muat_daftar_tab()
    for terlarang in ("PENGATURAN", "MASTER_CUSTOMER", "INVOICE", "SURAT_JALAN"):
        with pytest.raises(ValueError, match="hanya boleh menulis"):
            p.tulis_tab(terlarang, [["x"]])
    assert s.ditulis == {}
    assert s.dikosongkan == []
    assert s.tab_dibuat == []


def test_bot_membuat_dan_mengisi_tab_sendiri():
    from hp_dokumen.sapu.tulis_sheet import PenulisSheet, TAB_STATUS

    s = _SambunganPalsu(["PENGATURAN", "MASTER_CUSTOMER"])
    p = PenulisSheet(s, "id")
    p.muat_daftar_tab()
    p.tulis_tab("BOT_UJI", [["a", "b"], [1, 2]])

    assert s.tab_dibuat == ["BOT_UJI"]
    assert s.ditulis["'BOT_UJI'!A1"] == [["a", "b"], [1, 2]]
    # tab milik Yosua tidak tersentuh sama sekali
    assert all("PENGATURAN" not in k and "MASTER_CUSTOMER" not in k for k in s.ditulis)


def test_tab_bot_yang_sudah_ada_diisi_ulang_bukan_digandakan():
    from hp_dokumen.sapu.tulis_sheet import PenulisSheet

    s = _SambunganPalsu(["BOT_DAFTAR_PO"])
    p = PenulisSheet(s, "id")
    p.muat_daftar_tab()
    p.tulis_tab("BOT_DAFTAR_PO", [["baru"]])
    assert s.tab_dibuat == []
    assert s.dikosongkan == ["'BOT_DAFTAR_PO'"]


def test_daftar_po_memuat_kolom_yang_dibutuhkan():
    from hp_dokumen.sapu.tulis_sheet import PenulisSheet, TAB_DAFTAR

    s = _SambunganPalsu([])
    p = PenulisSheet(s, "id")
    p.muat_daftar_tab()
    p.tulis_daftar_po([{"sumber": "Agustus 2026", "tab": "PO 25 Agustus - Panda & Bear",
                        "customer": "Panda & Bear", "qty": 120, "nett": 5126640.0,
                        "cara_bayar": "TOP", "perusahaan": "CV DWI PUTRA MANDIRI",
                        "kena_ppn": True, "siap": True}])
    isi = s.ditulis[f"'{TAB_DAFTAR}'!A1"]
    assert "CARA BAYAR" in isi[0] and "PPN" in isi[0] and "DRAF SIAP?" in isi[0]
    assert isi[1][6] == 120 and isi[1][9] == 5126640
    assert isi[1][12] == "SIAP"
