# -*- coding: utf-8 -*-
"""Menyusun spreadsheet ALUR KERJA sesuai penjelasan Yosua 12 September 2026.

Menghasilkan dua bentuk:
  keluaran/ALUR_KERJA_OTOMATISASI.xlsx  - tujuh tab, rapi, untuk dibuka di Excel
  keluaran/ALUR_KERJA_OTOMATISASI.csv   - satu lembar, untuk diunggah ke Drive
"""
import csv, io, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from hp_dokumen.konfigurasi import Konfigurasi   # noqa: E402

JUDUL = Font(bold=True, size=14, color="FFFFFF")
ISI_JUDUL = PatternFill("solid", fgColor="1F3864")
KEPALA = Font(bold=True, size=11, color="FFFFFF")
ISI_KEPALA = PatternFill("solid", fgColor="4472C4")
TEBAL = Font(bold=True)
GARIS = Border(*[Side(style="thin", color="BFBFBF")] * 4)
BUNGKUS = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
wb.remove(wb.active)


def lembar(nama, judul, kolom, baris, lebar):
    ws = wb.create_sheet(nama)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(kolom))
    s = ws.cell(1, 1, judul)
    s.font, s.fill = JUDUL, ISI_JUDUL
    s.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28
    for i, j in enumerate(kolom, start=1):
        c = ws.cell(3, i, j)
        c.font, c.fill, c.alignment, c.border = KEPALA, ISI_KEPALA, BUNGKUS, GARIS
    for r, isi in enumerate(baris, start=4):
        for i, nilai in enumerate(isi, start=1):
            c = ws.cell(r, i, nilai)
            c.alignment, c.border = BUNGKUS, GARIS
            if i == 1:
                c.font = TEBAL
    for i, w_ in enumerate(lebar, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w_
    ws.freeze_panes = "A4"
    return ws


# ------------------------------------------------------------------ 1
lembar(
    "ALUR UTAMA",
    "ALUR KERJA SISTEM — SESUAI PENJELASAN YOSUA 12 SEPTEMBER 2026",
    ["Urutan", "Yang terjadi", "Keterangan"],
    [
        ("1. Sumber",
         "Spreadsheet (External) Order Sheet di Google Drive.",
         "Satu berkas per bulan. Ada arsip 2025 dan 2026. Sistem harus bisa "
         "membaca yang mana pun, lama maupun baru."),
        ("2. Masuk ke tab PO",
         "Tiap tab bernama pola 'PO <tanggal> - <customer>' diperlakukan sebagai "
         "satu pesanan.",
         "Satu tab bisa berisi beberapa tabel bertumpuk dengan sistem ukuran "
         "berbeda. SEMUA tabel ikut dibaca, bukan hanya yang pertama."),
        ("3. Ambil bagian AVAILABLE TO ORDER",
         "Hanya kolom ATO yang dipakai sebagai dasar dokumen. Kolom ORIGINAL PO "
         "TIDAK dipakai.",
         "ATO = barang yang benar-benar bisa dikirim. Baris dengan ATO nol "
         "dilewati."),
        ("4. Ambil nilai bersih",
         "Nilai bersih diambil langsung dari kolomnya sendiri (TOTAL VALUE, "
         "CBD, atau COD), tidak pernah dihitung dari persentase.",
         "Kolom yang terisi hanya sebagian diabaikan seluruhnya dan pesanan itu "
         "diperlakukan sebagai TOP."),
        ("5. Keluarkan dokumen",
         "INVOICE, SURAT JALAN, dan FAKTUR PAJAK — mengikuti karakteristik "
         "masing-masing customer.",
         "Lihat tab KARAKTERISTIK CUSTOMER untuk aturan tiap toko."),
        ("6. Ikut berubah terus",
         "Setiap (External) Order Sheet berubah, dokumen ikut diperbarui "
         "memakai data PALING BARU.",
         "ATO terisi = pesanan dianggap final untuk pertama kali, TAPI belum "
         "final selamanya. Kalau ada revisi, dokumen wajib dibetulkan lagi. "
         "Lihat tab ATURAN SELALU UPDATE."),
    ],
    [26, 62, 72],
)

# ------------------------------------------------------------------ 2
lembar(
    "ATURAN SELALU UPDATE",
    "ATURAN 'SELALU IKUT DATA TERBARU' — PENJELASAN YOSUA",
    ["Keadaan", "Yang harus dilakukan sistem"],
    [
        ("ATO baru pertama kali terisi",
         "Pesanan dianggap final. Draf pertama Invoice, Surat Jalan, dan "
         "Faktur Pajak dikeluarkan."),
        ("ATO direvisi setelah draf keluar",
         "Dokumen WAJIB dibuat ulang mengikuti angka terbaru. Draf lama tidak "
         "boleh dipakai. Ini ditegaskan Yosua: 'tidak sepenuhnya final, kalau "
         "ada revisi harus diperbaiki lagi dan disesuaikan dengan data paling "
         "terbaru'."),
        ("Qty berubah",
         "Surat Jalan, Invoice, dan Faktur Pajak ikut berubah semua, karena "
         "ketiganya bersumber dari angka yang sama."),
        ("Harga atau nilai bersih berubah",
         "Invoice dan Faktur Pajak berubah. Surat Jalan tidak, karena Surat "
         "Jalan tidak memuat nilai rupiah."),
        ("Baris artikel bertambah atau berkurang",
         "Seluruh dokumen dibuat ulang. Nomor urut baris ikut menyesuaikan."),
        ("Cara bayar berubah (TOP / CBD / COD)",
         "Nilai bersih ikut berubah, jadi Invoice dan Faktur Pajak dibuat "
         "ulang."),
        ("Rumus diubah walau angkanya belum berubah",
         "Dianggap perubahan dan dilaporkan, karena angkanya bisa berubah "
         "sewaktu-waktu setelah itu."),
        ("Pengaman",
         "Sebelum dokumen dibuat ulang, angka program dicocokkan dulu dengan "
         "baris TOTAL milik order sheet sendiri. Kalau TIDAK COCOK, dokumen "
         "TIDAK dikeluarkan dan masalahnya dilaporkan."),
        ("Siapa yang memantau",
         "Bot penyapu memeriksa seluruh order sheet tiap 12 jam dan "
         "membunyikan alarm GENTING kalau pesanan yang ATO-nya sudah terisi "
         "ternyata diubah."),
    ],
    [34, 96],
)

# ------------------------------------------------------------------ 3
cfg = Konfigurasi.muat()
baris_cust = []
for c in cfg.customer.semua:
    fmt = ("per artikel + ukuran" if c.format_invoice == "per_ukuran" else "per artikel")
    baris_cust.append((
        c.kunci,
        c.nama_di_dokumen or "BELUM ADA",
        "ada" if c.alamat else "BELUM ADA",
        c.npwp or "BELUM ADA",
        fmt,
        f"{c.termin_hari} hari",
        c.perusahaan_pemroses or "belum ditentukan",
    ))

lembar(
    "KARAKTERISTIK CUSTOMER",
    "KARAKTERISTIK TIAP CUSTOMER — MENENTUKAN BENTUK DOKUMENNYA",
    ["Customer", "Nama di dokumen", "Alamat", "NPWP", "Format invoice",
     "Termin", "Perusahaan pemroses"],
    baris_cust,
    [22, 30, 12, 12, 20, 12, 20],
)

# ------------------------------------------------------------------ 4
lembar(
    "ISI TIAP DOKUMEN",
    "BENTUK KETIGA DOKUMEN YANG DIMINTA",
    ["Dokumen", "Susunannya", "Hal khusus"],
    [
        ("INVOICE",
         "Kolom: No. | ARTICLE CODE | DESKRIPSI BARANG | Qty (PCS) | Harga "
         "(Satuan) | Diskon (%) | Nilai (Diskon) | Jumlah. Penutup: Subtotal, "
         "Diskon, Total, Uang Muka, DPP, PPN, Total.",
         "SATU tabel menerus untuk seluruh PO, tidak dipecah per blok. Tanpa "
         "warna. Baris diskon CBD/COD TIDAK dicetak. Haritsa dan Katamama "
         "dipecah sampai ukuran, customer lain cukup per artikel."),
        ("SURAT JALAN",
         "Kolom: No. | ARTICLE CODE | PRODUCT NAME | COLOUR | <kolom ukuran> | "
         "TOTAL.",
         "SATU tabel per blok, bertumpuk dalam satu dokumen. Tiap tabel punya "
         "label ukurannya sendiri dan nomor urut mulai dari 1 lagi. Di bawah "
         "semua tabel ada TOTAL SELURUH PO. Tanda tangan: Pengirim / Penerima "
         "/ Mengetahui."),
        ("FAKTUR PAJAK",
         "Data siap ketik ke Coretax: nama dan NPWP pembeli, alamat, nomor "
         "referensi, tanggal, DPP, PPN, total, ditambah rincian per artikel.",
         "Bukan untuk dicetak. Harga di order sheet SUDAH termasuk PPN, jadi "
         "DPP dihitung mundur: DPP = Total / 1,11. Tarif 11% masih menunggu "
         "kepastian."),
        ("PACKING LIST",
         "Sama seperti Surat Jalan, ditambah kolom JUMLAH DIKIRIM dan NO. KOLI "
         "yang diisi gudang.",
         "CATATAN: Yosua menyebut tiga dokumen saja (Invoice, Surat Jalan, "
         "Faktur Pajak). Packing List tetap dibuat program karena sudah ada, "
         "tapi PERLU DIPASTIKAN apakah memang masih dipakai."),
    ],
    [18, 62, 76],
)

# ------------------------------------------------------------------ 5
lembar(
    "SUMBER ANGKA",
    "DARI KOLOM MANA ANGKANYA DIAMBIL",
    ["Yang dibutuhkan", "Diambil dari", "Catatan penting"],
    [
        ("Kode artikel & nama barang", "Kolom ARTICLE CODE dan PRODUCT NAME",
         "Nama barang terbukti selalu sama persis dengan master Harga Retail, "
         "diperiksa pada seluruh 1.194 baris."),
        ("Qty per ukuran", "Kolom AVAILABLE TO ORDER",
         "BUKAN dari ORIGINAL PO. Label ukurannya mengikuti baris judul milik "
         "blok itu sendiri."),
        ("Harga satuan", "Kolom PRICE W/ VAT",
         "Sudah termasuk PPN."),
        ("Nilai sebelum diskon", "Kolom TOTAL ATO (VALUE)",
         "qty x harga, sudah diperiksa cocok untuk 1.194 dari 1.194 baris."),
        ("Nilai bersih", "Kolom TOTAL VALUE / CBD / COD",
         "Diambil langsung, TIDAK PERNAH dihitung dari persen. Judul kolomnya "
         "yang menentukan cara bayar dan tarifnya."),
        ("Cara bayar (TOP/CBD/COD)", "Judul kolom nilai bersih yang terisi penuh",
         "Kolom yang terisi sebagian diabaikan seluruhnya, pesanan jadi TOP."),
        ("Letak semua kolom di atas", "Dicari dari TULISAN di baris judul",
         "Susunan kolom berubah tiap periode: Januari 2025 punya 3 kolom "
         "ukuran, Feb-Mei 6, Jul-Sep 2025 8, Okt 2025 sampai kini 9. Huruf "
         "kolom tetap TIDAK boleh dipakai."),
    ],
    [28, 34, 76],
)

# ------------------------------------------------------------------ 6
lembar(
    "SUDAH PASTI",
    "YANG SUDAH PASTI — SUDAH DIJAWAB YOSUA",
    ["Hal", "Keputusan", "Tanggal"],
    [
        ("Sumber data", "Bagian AVAILABLE TO ORDER pada tab PO di (External) "
         "Order Sheet.", "12 Sep 2026"),
        ("Dokumen yang diminta", "Invoice, Surat Jalan, dan Faktur Pajak, "
         "mengikuti karakteristik masing-masing customer.", "12 Sep 2026"),
        ("Kapan pesanan dianggap final", "Saat ATO sudah terisi. Tapi TIDAK "
         "final selamanya — kalau direvisi, dokumen wajib disesuaikan ke data "
         "terbaru.", "12 Sep 2026"),
        ("Sifat sistem", "Harus selalu ikut berubah setiap (External) Order "
         "Sheet berubah.", "12 Sep 2026"),
        ("Cakupan", "Semua order sheet, lama maupun baru, kapan pun "
         "dibutuhkan.", "11 Sep 2026"),
        ("Akhiran Y pada ukuran", "Tetap dipakai (2 menjadi 2Y).", "11 Sep 2026"),
        ("PPN 11%", "Hanya untuk pesanan lewat CV. Dwi Putra Mandiri, tidak "
         "berlaku untuk CV. Mutiara Timur Nusantara.", "11 Sep 2026"),
        ("Termin 30 hari", "Tidak berlaku untuk semua customer, diambil dari "
         "riwayat order sheet.", "11 Sep 2026"),
    ],
    [28, 84, 14],
)

# ------------------------------------------------------------------ 7
lembar(
    "MENUNGGU KONFIRMASI",
    "YANG MASIH MENUNGGU KEPASTIAN DARI YOSUA",
    ["Hal", "Pertanyaannya", "Sementara program memakai", "Risiko kalau tebakan salah"],
    [
        ("Pengiriman bertahap",
         "Satu PO dikirim sekali atau beberapa kali? Kalau bertahap, satu PO "
         "jadi berapa Surat Jalan, dan Invoice terbit per pengiriman atau "
         "sekali untuk seluruh PO?",
         "1 PO = 1 Surat Jalan = 1 Invoice",
         "BESAR. Kalau kenyataannya bertahap, rancangan sekarang salah, bukan "
         "sekadar kurang. Petunjuknya sudah ada: tab Packing List Haritsa dulu "
         "berisi 333 pcs dari PO yang 1.053 pcs."),
        ("Perusahaan pemroses (DPM / MTN)",
         "Melekat ke customer selamanya, atau bisa berbeda tiap pesanan?",
         "Melekat ke customer, diisi di config/customer.csv",
         "BESAR. Kalau ternyata per pesanan, PPN 11% bisa salah kena atau "
         "salah tidak kena."),
        ("Urutan penerbitan dokumen",
         "Surat Jalan dulu baru Invoice, atau bersamaan? Faktur Pajak dibuat "
         "saat kirim atau saat tagih?",
         "Ketiganya dibuat sekaligus",
         "SEDANG. Tidak membuat angka salah, tapi bisa tidak sesuai prosedur "
         "kantor."),
        ("Rumus nomor dokumen",
         "Siapa pemegang nomor urutnya? DPM dan MTN satu deret atau "
         "masing-masing? Surat Jalan dan Invoice nomornya sama atau beda?",
         "Nomor DIKOSONGKAN, diisi manual saat cetak",
         "KECIL selama diisi manual. Program sengaja tidak menebak nomor."),
        ("Packing List",
         "Masih dipakai atau tidak? Yosua menyebut tiga dokumen saja.",
         "Tetap dibuat",
         "KECIL. Hanya berkas tambahan yang mungkin tidak terpakai."),
        ("Tarif PPN",
         "Apakah 11% masih sesuai aturan yang berlaku sekarang?",
         "11%",
         "BESAR untuk pelaporan pajak. Perlu dipastikan ke konsultan pajak."),
    ],
    [26, 54, 34, 56],
)

Path("keluaran").mkdir(exist_ok=True)
wb.save("keluaran/ALUR_KERJA_OTOMATISASI.xlsx")

out = io.StringIO()
w = csv.writer(out, lineterminator="\n")
w.writerow(["ALUR KERJA SISTEM OTOMATISASI DOKUMEN PENJUALAN - HAPPY PUMPKIN"])
w.writerow(["CV Dwi Putra Mandiri - sesuai penjelasan Yosua 12 September 2026"])
for nama in wb.sheetnames:
    ws = wb[nama]
    w.writerow([])
    w.writerow(["### " + ws.cell(1, 1).value])
    w.writerow([ws.cell(3, c).value for c in range(1, ws.max_column + 1)])
    for r in range(4, ws.max_row + 1):
        isi = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(isi):
            w.writerow([("" if v is None else str(v)) for v in isi])
Path("keluaran/ALUR_KERJA_OTOMATISASI.csv").write_text(out.getvalue(), encoding="utf-8")
print("tab:", wb.sheetnames)
print("panjang csv:", len(out.getvalue()))
