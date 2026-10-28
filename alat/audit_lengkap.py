"""Audit MENYELURUH rantai uang order sheet: harga, diskon, CBD, COD.

Permintaan Yosua 21 Oktober 2026: *"auditlah seluruh ordersheet dan periksa
seluruh cell apakah ada kesalahan formula atau hitungan harga, diskon dan
diskon cbd ataupun cod"*.

Memperluas alat/audit_harga_diskon.py dengan tiga pemeriksaan yang belum
pernah dijalankan, dan ketiganya adalah cacat yang TIDAK kelihatan dari
angkanya:

  1. REFERENSI MELESET BARIS — `=Z3-(Z4*25%)` pada baris 3. Rumus per baris
     yang menunjuk baris LAIN memberi angka yang masuk akal tapi salah,
     dan tidak ada tanda apa pun di hasilnya.
  2. CBD/COD vs ANGKA, bukan cuma vs judul — AD harus = AC x (1 - tarif).
     Audit sebelumnya hanya membandingkan tarif di RUMUS dengan judul
     kolom; kalau rumusnya menunjuk sel yang salah, keduanya tetap cocok.
  3. CBD/COD menunjuk kolom SUMBER yang salah — AD seharusnya diturunkan
     dari AC (nilai bersih), bukan dari Z (nilai kotor).

Ditambah pemeriksaan kelengkapan kolom CBD/COD, sebab kolom yang terisi
SEBAGIAN membuat Aturan 2 memperlakukan order sebagai TOP — dan itu
mengubah total tanpa ada yang salah secara teknis.
"""
from __future__ import annotations

import csv
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from hp_dokumen import tata_letak as tl
from hp_dokumen.pemindai import adalah_tab_harga, cari_di_master, master_harga_dari_buku
from hp_dokumen.periksa_rumus import persen_dalam_rumus, persen_sel

STAT = Counter()
# sel yang disebut rumus, contoh Z3 / $AC$12
_SEL_RUMUS = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def _n(v):
    return float(v) if isinstance(v, (int, float)) else None


def _rumus(ws, r, c):
    v = ws.cell(r, c).value
    return v if isinstance(v, str) and v.startswith("=") else None


def baris_disebut(rumus: str) -> set[int]:
    """Nomor baris yang disebut rumus ini (tanpa bagian nama lembar)."""
    bersih = re.sub(r"'[^']*'!", "", rumus)
    bersih = re.sub(r"\b[A-Za-z ]+!", "", bersih)
    return {int(m[1]) for m in _SEL_RUMUS.findall(bersih)}


def kolom_disebut(rumus: str) -> set[str]:
    bersih = re.sub(r"'[^']*'!", "", rumus)
    bersih = re.sub(r"\b[A-Za-z ]+!", "", bersih)
    return {m[0] for m in _SEL_RUMUS.findall(bersih)}


def audit_tab(berkas, ws, wv, master):
    t = []

    def lapor(tingkat, jenis, sel, isi, pesan):
        t.append([tingkat, berkas, ws.title, jenis, sel, str(isi)[:70], pesan])

    mr = min(ws.max_row or 1, 2000)
    mc = min(ws.max_column or 1, 60)
    judul = [r for r in range(1, mr + 1)
             if tl.seragam(ws.cell(r, 1).value) == "ARTICLE CODE"]
    if not judul:
        return t
    tata0 = tl.kenali(ws, judul[0], mc)
    if not tata0.lengkap() or tata0.hanya_qty:
        return t
    STAT["tab"] += 1

    k_top = tata0.kolom_top()
    kol_nett = [k for k in tata0.nett if k.jenis in ("CBD", "COD")]
    batas = judul + [mr + 1]

    meleset = []            # referensi salah baris
    cbd_cod_salah = defaultdict(list)
    terisi = defaultdict(int)
    n_baris = 0
    persen_ac = Counter()
    persen_ab = Counter()
    harga_diketik = 0

    for i, br in enumerate(judul):
        tata = tl.kenali(ws, br, mc)
        if not tata.lengkap():
            continue
        kt = tata.kolom_top()
        for r in range(br + 2, batas[i + 1]):
            kode = ws.cell(r, tata.kol_kode).value
            if not kode or tl.seragam(kode) in ("ARTICLE CODE", "TOTAL"):
                continue
            qty = _n(wv.cell(r, tata.kol_ato_total).value) if tata.kol_ato_total else None
            if not qty:
                continue
            n_baris += 1
            STAT["baris"] += 1

            # ---- 1. referensi MELESET BARIS -----------------------------
            # Rumus per baris hanya boleh menyebut barisnya sendiri.
            for c in [tata.kol_harga, tata.kol_ato_value, tata.kol_ato_total] + \
                     [k.kolom for k in tata.nett]:
                if not c:
                    continue
                f = _rumus(ws, r, c)
                if not f or "VLOOKUP" in f.upper():
                    continue
                lain = {x for x in baris_disebut(f) if x != r and x > 2}
                if lain:
                    meleset.append((f"{L(c)}{r}", f, sorted(lain)[:3]))

            harga = _n(wv.cell(r, tata.kol_harga).value) if tata.kol_harga else None
            kotor = _n(wv.cell(r, tata.kol_ato_value).value) if tata.kol_ato_value else None
            nett = _n(wv.cell(r, kt.kolom).value) if kt else None

            if tata.kol_harga and not _rumus(ws, r, tata.kol_harga) and harga:
                harga_diketik += 1
            else:
                STAT["harga VLOOKUP"] += 1

            # ---- harga vs tab Harga Retail -----------------------------
            if master and harga:
                m = cari_di_master(master, str(kode))
                if m is None:
                    lapor("GENTING", "KODE TAK TERDAFTAR", f"{L(tata.kol_kode)}{r}",
                          kode, "tidak ada di tab Harga Retail")
                else:
                    hm = m[1] if isinstance(m, tuple) else getattr(m, "harga", 0)
                    if hm and abs(harga - hm) > 0.5:
                        lapor("GENTING", "HARGA BEDA MASTER", f"{L(tata.kol_harga)}{r}",
                              f"{harga:,.0f}", f"Harga Retail menulis {hm:,.0f}")

            # ---- nilai kotor --------------------------------------------
            if harga is not None and kotor is not None and abs(qty * harga - kotor) > 1:
                lapor("GENTING", "NILAI KOTOR SALAH", f"{L(tata.kol_ato_value)}{r}",
                      f"{kotor:,.0f}",
                      f"qty {qty:g} x harga {harga:,.0f} = {qty*harga:,.0f}")

            # ---- persen diskon -------------------------------------------
            if kt:
                p = persen_dalam_rumus(_rumus(ws, r, kt.kolom))
                if p is not None:
                    persen_ac[round(p, 6)] += 1
                    STAT["AC berumus"] += 1
                    if kotor and nett is not None and abs(kotor * (1 - p) - nett) > 1:
                        lapor("GENTING", "NILAI BERSIH TIDAK COCOK RUMUS",
                              f"{L(kt.kolom)}{r}", f"{nett:,.2f}",
                              f"kotor {kotor:,.0f} x (1-{p:.4%}) = {kotor*(1-p):,.2f}")
            if tata.kol_disc:
                d = persen_sel(wv.cell(r, tata.kol_disc).value)
                if d is not None:
                    persen_ab[round(d, 6)] += 1
                    STAT["DISC terbaca"] += 1

            # ---- 2 & 3. CBD / COD ----------------------------------------
            for k in tata.nett:
                if k.jenis not in ("CBD", "COD"):
                    continue
                nilai = _n(wv.cell(r, k.kolom).value)
                if nilai:
                    terisi[(k.kolom, k.judul)] += 1
                f = _rumus(ws, r, k.kolom)
                if not f:
                    continue
                STAT["CBD/COD berumus"] += 1
                # sumbernya harus kolom nilai bersih (AC), bukan kolom kotor
                if kt:
                    kols = kolom_disebut(f)
                    if kols and L(kt.kolom) not in kols:
                        cbd_cod_salah[("SUMBER", k.kolom, k.judul)].append(
                            (f"{L(k.kolom)}{r}", f))
                # angkanya harus = AC x (1 - tarif judul)
                if k.tarif is not None and nett is not None and nilai is not None:
                    harap = nett * (1 - k.tarif)
                    if abs(harap - nilai) > 1:
                        cbd_cod_salah[("ANGKA", k.kolom, k.judul)].append(
                            (f"{L(k.kolom)}{r}", f"{nilai:,.2f} vs {harap:,.2f}"))

    if not n_baris:
        return t

    # ---- ringkasan per tab -------------------------------------------
    if meleset:
        contoh = "; ".join(f"{s} -> {f} (menyebut baris {l})"
                           for s, f, l in meleset[:3])
        lapor("GENTING", "REFERENSI MELESET BARIS", meleset[0][0],
              f"{len(meleset)} sel", contoh)

    if harga_diketik:
        lapor("POTENSI", "HARGA DIKETIK TANGAN", "", f"{harga_diketik} baris",
              "bukan VLOOKUP ke tab Harga Retail, jadi tidak ikut berubah "
              "kalau Harga Retail diperbarui")

    if len(persen_ac) > 1:
        lapor("GENTING", "DISKON TIDAK SERAGAM", "", "",
              "persen di rumus nilai bersih berbeda-beda: "
              + ", ".join(f"{p:.4%} ({n} baris)" for p, n in sorted(persen_ac.items())))

    if persen_ac and persen_ab:
        pac = max(persen_ac, key=persen_ac.get)
        pab = max(persen_ab, key=persen_ab.get)
        if abs(pac - pab) > 0.0005:
            lapor("GENTING", "KOLOM DISC TIDAK COCOK RUMUS", "",
                  f"DISC tertulis {pab:.4%}",
                  f"rumus nilai bersih memakai {pac:.4%}. Yang dipakai "
                  "menghitung uang adalah yang di rumus")

    for (jenis, kol, jud), daftar in cbd_cod_salah.items():
        nama = "CBD/COD SUMBER SALAH" if jenis == "SUMBER" else "CBD/COD ANGKA SALAH"
        lapor("GENTING", nama, L(kol), jud,
              f"{len(daftar)} baris — contoh: "
              + "; ".join(f"{s} {d}" for s, d in daftar[:2]))

    # kelengkapan kolom CBD/COD — Aturan 2
    for k in kol_nett:
        n = terisi.get((k.kolom, k.judul), 0)
        if 0 < n < n_baris:
            lapor("POTENSI", "CBD/COD TERISI SEBAGIAN", L(k.kolom), k.judul,
                  f"terisi {n} dari {n_baris} baris. Aturan 2 MENGABAIKANNYA "
                  "seluruhnya dan memperlakukan order sebagai TOP")
    return t


def main(folder):
    semua = []
    for p in sorted(Path(folder).glob("*.xlsx")):
        wb = openpyxl.load_workbook(p)
        wv = openpyxl.load_workbook(p, data_only=True)
        master = master_harga_dari_buku(wv)
        n = 0
        for j in wb.sheetnames:
            if adalah_tab_harga(j):
                continue
            tt = audit_tab(p.stem, wb[j], wv[j], master)
            semua += tt
            n += len(tt)
        g = sum(1 for x in semua if x[1] == p.stem and x[0] == "GENTING")
        print(f"{p.stem:<28} master {len(master):>3}  temuan {n:>4}  GENTING {g:>3}")

    Path("keluaran").mkdir(exist_ok=True)
    with open("keluaran/AUDIT_LENGKAP_2026.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["TINGKAT", "ORDER SHEET", "TAB", "JENIS", "SEL", "ISI",
                    "KETERANGAN"])
        w.writerows(semua)

    print("\nYANG BENAR-BENAR DIPERIKSA:")
    for k in ("tab", "baris", "harga VLOOKUP", "AC berumus", "DISC terbaca",
              "CBD/COD berumus"):
        print(f"  {k:<18} {STAT[k]:>7}")
    print(f"\nTOTAL {len(semua)} temuan")
    for (tk, j), n in Counter((x[0], x[3]) for x in semua).most_common():
        print(f"  {tk:<8} {j:<30} {n}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "data/audit_yosua")
