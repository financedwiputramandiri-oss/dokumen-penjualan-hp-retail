"""Susun hasil audit jadi satu berkas Excel yang enak dibaca divisi."""
from __future__ import annotations

import csv
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

JUDUL = Font(bold=True, color="FFFFFF", size=11)
BG = PatternFill("solid", fgColor="31506E")
MERAH = PatternFill("solid", fgColor="F8CBAD")
KUNING = PatternFill("solid", fgColor="FFE699")
HIJAU = PatternFill("solid", fgColor="C6E0B4")


def lembar(wb, nama, judul, baris, lebar):
    ws = wb.create_sheet(nama)
    ws.append(judul)
    for c in range(1, len(judul) + 1):
        ws.cell(1, c).font = JUDUL
        ws.cell(1, c).fill = BG
        ws.cell(1, c).alignment = Alignment(horizontal="center", vertical="center")
    for b in baris:
        ws.append(b)
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    temuan = list(csv.DictReader(open("keluaran/AUDIT_TEMUAN.csv")))
    tab = list(csv.DictReader(open("keluaran/AUDIT_TAB.csv")))
    periksa = Path("/tmp/periksa2.txt").read_text() if Path("/tmp/periksa2.txt").exists() else ""

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ---- RINGKASAN ----
    ws = wb.create_sheet("RINGKASAN")
    ws.append(["AUDIT ORDER SHEET 2026 - 20 Oktober 2026"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Order sheet diperiksa", len({t["ORDER SHEET"] for t in tab})])
    ws.append(["Tab diperiksa", len(tab)])
    ws.append(["Baris data (ATO terisi)", sum(int(t["BARIS DATA"]) for t in tab)])
    ws.append([])
    ws.append(["TINGKAT", "JENIS", "JUMLAH", "ARTINYA"])
    for c in range(1, 5):
        ws.cell(7, c).font = JUDUL
        ws.cell(7, c).fill = BG
    arti = {
        ("GENTING", "RUMUS QTY PENDEK"):
            "qty di kolom ukuran terakhir TIDAK ikut terhitung - angka PO salah",
        ("GENTING", "BARIS TOTAL KURANG"):
            "baris TOTAL order sheet tidak menjumlah seluruh barisnya",
        ("GENTING", "GALAT EXCEL"):
            "#REF! di baris TOTAL kolom nilai bersih - totalnya tidak terbaca",
        ("GENTING", "TATA LETAK"):
            "susunan kolom tab tidak dikenali program",
        ("POTENSI", "RUMUS QTY PENDEK"):
            "rumus berhenti sebelum kolom terakhir, tapi kolom sisanya masih "
            "kosong - akan salah begitu Sales mengisinya",
        ("POTENSI", "GALAT EXCEL"):
            "#REF! di kolom AA (LOSSES) - tidak dipakai dokumen mana pun",
    }
    from collections import Counter
    for (tk, j), n in Counter((t["TINGKAT"], t["JENIS"]) for t in temuan).most_common():
        ws.append([tk, j, n, arti.get((tk, j), "")])
        ws.cell(ws.max_row, 1).fill = MERAH if tk == "GENTING" else KUNING
    for kol, w in zip("ABCD", (12, 24, 10, 78)):
        ws.column_dimensions[kol].width = w

    # ---- PER TAB ----
    lembar(wb, "PER TAB",
           ["ORDER SHEET", "TAB", "BLOK", "BARIS DATA", "SEL GALAT", "CATATAN",
            "DIKENALI"],
           [[t["ORDER SHEET"], t["TAB"], int(t["BLOK"]), int(t["BARIS DATA"]),
             int(t["SEL GALAT"]), t["CATATAN"], t["DIKENALI"]] for t in tab],
           (28, 36, 7, 12, 11, 16, 11))

    # ---- GENTING ----
    g = [t for t in temuan if t["TINGKAT"] == "GENTING"]
    ws = lembar(wb, "GENTING",
                ["ORDER SHEET", "TAB", "JENIS", "SEL", "ISI SEL", "KETERANGAN"],
                [[t["ORDER SHEET"], t["TAB"], t["JENIS"], t["SEL"], t["ISI"],
                  t["KETERANGAN"]] for t in g],
                (28, 36, 22, 9, 40, 70))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 3).fill = MERAH

    # ---- POTENSI, diringkas per tab supaya tidak 12.000 baris ----
    pot = [t for t in temuan if t["TINGKAT"] == "POTENSI"]
    ring = {}
    for t in pot:
        k = (t["ORDER SHEET"], t["TAB"], t["JENIS"])
        d = ring.setdefault(k, {"n": 0, "contoh": t["SEL"], "ket": t["KETERANGAN"]})
        d["n"] += 1
    lembar(wb, "POTENSI (ringkas)",
           ["ORDER SHEET", "TAB", "JENIS", "JUMLAH SEL", "CONTOH SEL",
            "KETERANGAN"],
           [[o, tb, j, d["n"], d["contoh"], d["ket"]]
            for (o, tb, j), d in sorted(ring.items(), key=lambda x: -x[1]["n"])],
           (28, 36, 22, 12, 12, 70))

    # ---- PENCOCOKAN ----
    baris = []
    blok = ""
    for ln in periksa.splitlines():
        if ln.startswith("### "):
            blok = ln[4:].strip()
        elif "COCOK" in ln or "GAGAL" in ln:
            m = re.match(r"^(.*?)\s{2,}(\d+)\s+(\d+)\s+([\d.]+)\s+Rp([\d.]+)\s+"
                         r"(\w+)\s+Rp([\d.]+)\s+(COCOK|GAGAL)", ln)
            if m:
                baris.append([blok] + list(m.groups()))
    ws = lembar(wb, "PENCOCOKAN",
                ["ORDER SHEET", "TAB", "BLOK", "BARIS", "QTY", "SEBELUM DISKON",
                 "BAYAR", "NILAI BERSIH", "HASIL"],
                baris, (28, 36, 7, 8, 10, 18, 8, 18, 9))
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 9).fill = MERAH if ws.cell(r, 9).value == "GAGAL" else HIJAU

    Path("keluaran").mkdir(exist_ok=True)
    out = "keluaran/AUDIT_ORDER_SHEET_2026.xlsx"
    wb.save(out)
    print("tersimpan:", out, "| lembar:", wb.sheetnames)


if __name__ == "__main__":
    main()
