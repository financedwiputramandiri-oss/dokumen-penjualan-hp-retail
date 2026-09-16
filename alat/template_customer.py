# -*- coding: utf-8 -*-
"""Menyusun lembar isian nama & alamat customer untuk diisi Yosua.

Gunanya supaya pembetulan nama customer tidak perlu mengetik langsung ke
`config/customer.csv`. Berkas CSV gampang rusak kalau nama customer mengandung
koma (`PT. ABC, Tbk`) — di Excel hal itu ditangani sendiri.

Jalankan:  python3 alat/template_customer.py
Hasil:     keluaran/TEMPLATE_NAMA_CUSTOMER.xlsx
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AKAR = Path(__file__).resolve().parent.parent
BERKAS_CUSTOMER = AKAR / "config" / "customer.csv"
BERKAS_DATABASE = AKAR / "keluaran" / "DATABASE_CUSTOMER.xlsx"
KELUARAN = AKAR / "keluaran" / "TEMPLATE_NAMA_CUSTOMER.xlsx"

JUDUL = Font(bold=True, size=14, color="FFFFFF")
LATAR_JUDUL = PatternFill("solid", fgColor="1F3864")
KEPALA = Font(bold=True, size=11, color="FFFFFF")
LATAR_KEPALA = PatternFill("solid", fgColor="4472C4")
LATAR_ISI = PatternFill("solid", fgColor="FFF2CC")      # kuning: kolom yang diisi orang
LATAR_SUDAH = PatternFill("solid", fgColor="E2EFDA")    # hijau: sudah terisi
GARIS = Border(*[Side(style="thin", color="BFBFBF")] * 4)
BUNGKUS = Alignment(wrap_text=True, vertical="top")

# Kolom yang benar-benar perlu diisi orang. Sisanya di customer.csv dibiarkan.
KOLOM = [
    ("kunci_tab", "KUNCI TAB\n(jangan diubah)", 24, False),
    ("nama_di_dokumen", "NAMA DI DOKUMEN\n(yang tercetak di faktur)", 38, True),
    ("alamat", "ALAMAT LENGKAP", 46, True),
    ("npwp", "NPWP", 24, True),
    ("id_tku", "NITKU / ID TKU\n(22 digit, dari Coretax)", 26, True),
    ("format_invoice", "FORMAT INVOICE\nper_artikel / per_ukuran", 22, True),
    ("perusahaan_pemroses", "DIPROSES LEWAT\nDPM / MTN", 18, True),
    ("termin_hari", "TERMIN\n(hari)", 12, True),
]


def baca_customer() -> list[dict]:
    with BERKAS_CUSTOMER.open(encoding="utf-8-sig", newline="") as f:
        return [r for r in csv.DictReader(f) if (r.get("kunci_tab") or "").strip()]


def baca_database() -> dict[str, dict]:
    """Ambil jumlah PO dan periode aktif dari DATABASE_CUSTOMER, kalau ada.

    Berkas itu hasil `jalankan.py telusuri` dan tidak selalu ada, jadi
    ketiadaannya tidak boleh menggagalkan pembuatan template.
    """
    if not BERKAS_DATABASE.exists():
        return {}
    ws = load_workbook(BERKAS_DATABASE, read_only=True)["MASTER_CUSTOMER"]
    baris = list(ws.iter_rows(values_only=True))
    judul = next((i for i, r in enumerate(baris) if r and r[0] == "NAMA CUSTOMER"), None)
    if judul is None:
        return {}
    k = {n: i for i, n in enumerate(baris[judul]) if n}
    hasil = {}
    for r in baris[judul + 1:]:
        if not r or not r[0]:
            continue
        hasil[str(r[0]).strip().lower()] = {
            "po": r[k.get("JML PO", 1)],
            "periode": r[k.get("PERIODE AKTIF", 2)],
        }
    return hasil


def _sel(ws, baris, kolom, nilai, *, latar=None, tebal=False):
    s = ws.cell(baris, kolom, nilai)
    s.border = GARIS
    s.alignment = BUNGKUS
    if latar:
        s.fill = latar
    if tebal:
        s.font = Font(bold=True)
    return s


def lembar_isian(wb, customer, database):
    ws = wb.create_sheet("ISI DI SINI")
    lebar_total = len(KOLOM) + 2

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=lebar_total)
    s = ws.cell(1, 1, "ISIAN NAMA & ALAMAT CUSTOMER")
    s.font, s.fill = JUDUL, LATAR_JUDUL

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=lebar_total)
    ws.cell(2, 1, "Isi kolom KUNING saja. Kolom hijau berarti sudah terisi — "
                  "boleh dibetulkan kalau ada yang salah. Jangan mengubah kolom KUNCI TAB.")

    for i, (_, judul, lebar, _) in enumerate(KOLOM, 1):
        _sel(ws, 4, i, judul, latar=LATAR_KEPALA).font = KEPALA
        ws.column_dimensions[get_column_letter(i)].width = lebar
    for i, judul in enumerate(("JML PO", "PERIODE AKTIF"), len(KOLOM) + 1):
        _sel(ws, 4, i, judul, latar=LATAR_KEPALA).font = KEPALA
        ws.column_dimensions[get_column_letter(i)].width = 26
    ws.row_dimensions[4].height = 34

    for r, row in enumerate(customer, 5):
        for i, (kunci, _, _, diisi) in enumerate(KOLOM, 1):
            nilai = (row.get(kunci) or "").strip()
            latar = LATAR_SUDAH if nilai else (LATAR_ISI if diisi else None)
            _sel(ws, r, i, nilai, latar=latar)
        d = database.get(row["kunci_tab"].strip().lower(), {})
        _sel(ws, r, len(KOLOM) + 1, d.get("po"))
        _sel(ws, r, len(KOLOM) + 2, d.get("periode"))

    ws.freeze_panes = "A5"
    return ws


def lembar_cara_pakai(wb, jumlah):
    ws = wb.create_sheet("CARA PAKAI", 0)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 95

    ws.merge_cells("A1:B1")
    s = ws.cell(1, 1, "CARA MEMBETULKAN NAMA CUSTOMER")
    s.font, s.fill = JUDUL, LATAR_JUDUL

    langkah = [
        ("1", "Buka tab ISI DI SINI."),
        ("2", "Isi kolom yang berwarna KUNING. Yang paling penting dua: "
              "NAMA DI DOKUMEN dan ALAMAT LENGKAP."),
        ("3", "NAMA DI DOKUMEN adalah nama yang akan tercetak di faktur dan surat "
              "jalan, contoh: PT. MAMA PAPA JUARA. Boleh berbeda dari nama tab."),
        ("4", "JANGAN mengubah kolom KUNCI TAB. Itu yang dipakai program untuk "
              "mencocokkan ke nama tab di order sheet."),
        ("5", "Kolom berwarna HIJAU sudah terisi. Boleh dibetulkan kalau salah."),
        ("6", "Simpan berkasnya, lalu kirimkan kembali untuk dimasukkan ke sistem."),
        ("", ""),
        ("", "CATATAN PENTING"),
        ("", "Setelah nama customer diperbarui, dokumen yang sudah pernah dibuat "
             "TIDAK ikut berubah dengan sendirinya. Bot hanya membuat ulang dokumen "
             "kalau ORDER SHEET-nya yang berubah."),
        ("", "Supaya seluruh dokumen ikut diperbarui, hapus berkas "
             "data/kondisi_sapu.json lalu jalankan sapuan sekali lagi. "
             "Menghapus berkas itu aman — isinya hanya catatan sapuan terakhir."),
        ("", ""),
        ("", f"Jumlah customer di lembar ini: {jumlah}"),
    ]
    for r, (no, teks) in enumerate(langkah, 3):
        ws.cell(r, 1, no).font = Font(bold=True)
        c = ws.cell(r, 2, teks)
        c.alignment = BUNGKUS
        if teks in ("CATATAN PENTING",):
            c.font = Font(bold=True, color="C00000")
        ws.row_dimensions[r].height = 30
    return ws


def main() -> int:
    if not BERKAS_CUSTOMER.exists():
        print(f"Tidak ketemu: {BERKAS_CUSTOMER}")
        return 1
    customer = baca_customer()
    wb = Workbook()
    wb.remove(wb.active)
    lembar_isian(wb, customer, baca_database())
    lembar_cara_pakai(wb, len(customer))
    KELUARAN.parent.mkdir(parents=True, exist_ok=True)
    wb.save(KELUARAN)
    belum = sum(1 for r in customer if not (r.get("nama_di_dokumen") or "").strip())
    print(f"{KELUARAN.relative_to(AKAR)} — {len(customer)} customer, "
          f"{belum} belum ada nama di dokumen.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
