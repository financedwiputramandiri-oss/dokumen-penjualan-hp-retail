"""Laporan audit menyeluruh harga, diskon, CBD & COD — 21 Oktober 2026."""
import csv, re
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

JUD = Font(bold=True, color="FFFFFF", size=11)
BG  = PatternFill("solid", fgColor="31506E")
MERAH = PatternFill("solid", fgColor="F8CBAD")
KUNING = PatternFill("solid", fgColor="FFF2CC")
HIJAU = PatternFill("solid", fgColor="C6E0B4")
tp = Side(style="thin", color="999999")
KOTAK = Border(left=tp, right=tp, top=tp, bottom=tp)
RP = '_-"Rp"* #,##0_-;\\-"Rp"* #,##0_-;_-"Rp"* "-"_-;_-@_-'


def lembar(wb, nama, judul, baris, lebar, warna_kol=None):
    ws = wb.create_sheet(nama)
    ws.append(judul)
    for c in range(1, len(judul) + 1):
        s = ws.cell(1, c); s.font, s.fill, s.border = JUD, BG, KOTAK
        s.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    for b in baris:
        ws.append(b)
        r = ws.max_row
        for c in range(1, len(judul) + 1):
            ws.cell(r, c).border = KOTAK
            ws.cell(r, c).alignment = Alignment(vertical="center", wrap_text=True)
        if warna_kol:
            warna_kol(ws, r)
    for i, w in enumerate(lebar, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    temuan = list(csv.DictReader(open("keluaran/AUDIT_LENGKAP_2026.csv")))
    persen = list(csv.DictReader(open("keluaran/PERSEN_HILANG.csv")))
    dampak = list(csv.DictReader(open("keluaran/DAMPAK_PERSEN_HILANG.csv")))

    wb = openpyxl.Workbook(); wb.remove(wb.active)

    # ---------- RINGKASAN ----------
    ws = wb.create_sheet("RINGKASAN")
    ws["A1"] = "AUDIT MENYELURUH ORDER SHEET 2026 — HARGA, DISKON, CBD & COD"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "21 Oktober 2026 · 11 order sheet (Januari–Oktober) · 239 tab · 11.045 baris"
    for b in [
        [], ["LINGKUP", ""],
        ["Order sheet diperiksa", 11],
        ["Tab berdata", 239],
        ["Baris diperiksa", 11045],
        ["Sel berumus diperiksa", "seluruhnya"],
        [],
        ["RANTAI YANG DIPERIKSA", "HASIL"],
        ["Harga = VLOOKUP ke tab Harga Retail", "11.045 dari 11.045 baris — BERSIH"],
        ["Harga diketik tangan", "0 — BERSIH"],
        ["Kode artikel tidak ada di Harga Retail", "0 — BERSIH"],
        ["Harga beda dari Harga Retail", "0 — BERSIH"],
        ["qty x harga = nilai kotor", "0 selisih — BERSIH"],
        ["Rumus menunjuk baris lain (meleset)", "0 — BERSIH"],
        ["CBD/COD menunjuk kolom sumber salah", "0 — BERSIH"],
        ["Nilai bersih cocok dengan persen di rumusnya", "0 selisih — BERSIH"],
        [],
        ["TEMUAN", "JUMLAH"],
        ["GENTING — tanda % HILANG di rumus diskon", "16 tab, 249 sel, 200 negatif"],
        ["GENTING — rumus qty berhenti terlalu awal", "2 tab"],
        ["GENTING — baris TOTAL tidak lengkap", "2 tab"],
        ["GENTING — kolom DISC beda dari rumus", "1 tab"],
        ["POTENSI — kolom CBD/COD terisi sebagian", "56 tab"],
        [],
        ["DOKUMEN YANG BENAR-BENAR SALAH", "1"],
        ["PO 11 September - Katamama (Tapos)", "kurang menagih Rp999.034"],
    ]:
        ws.append(b)
    for r in range(1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "")
        if a in ("LINGKUP", "RANTAI YANG DIPERIKSA", "TEMUAN",
                 "DOKUMEN YANG BENAR-BENAR SALAH"):
            for c in (1, 2):
                ws.cell(r, c).font, ws.cell(r, c).fill = JUD, BG
        if a.startswith("GENTING"):
            ws.cell(r, 1).fill = MERAH
        if a.startswith("POTENSI"):
            ws.cell(r, 1).fill = KUNING
        if "BERSIH" in str(ws.cell(r, 2).value or ""):
            ws.cell(r, 2).fill = HIJAU
    ws.cell(ws.max_row, 1).fill = MERAH
    ws.cell(ws.max_row, 2).fill = MERAH
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 40

    # ---------- TANDA % HILANG ----------
    def w_persen(ws_, r):
        ws_.cell(r, 5).fill = MERAH
    lembar(wb, "TANDA % HILANG",
           ["ORDER SHEET", "TAB", "KOLOM", "JUMLAH SEL", "HASIL NEGATIF",
            "CONTOH SEL", "RUMUS SALAH", "SEHARUSNYA", "NILAI SEKARANG"],
           [[x["ORDER SHEET"], x["TAB"], x["KOLOM"], int(x["JUMLAH SEL"]),
             int(x["HASIL NEGATIF"]), x["CONTOH SEL"], x["RUMUS"],
             x["RUMUS"].replace(")", "%)") if x["RUMUS"] else "", x["NILAI"]]
            for x in persen],
           (26, 34, 22, 11, 13, 11, 22, 23, 15), w_persen)

    # ---------- DAMPAK KE DOKUMEN ----------
    def w_dampak(ws_, r):
        salah = ws_.cell(r, 6).value == "YA"
        ws_.cell(r, 6).fill = MERAH if salah else HIJAU
    lembar(wb, "DAMPAK KE DOKUMEN",
           ["ORDER SHEET", "TAB", "CARA BAYAR DIPILIH", "KOLOM", "NILAI BERSIH",
            "DOKUMEN IKUT SALAH?"],
           [[x["ORDER SHEET"], x["TAB"], x["CARA BAYAR DIPILIH"], x["KOLOM"],
             float(x["NETT"]), x["ADA NILAI NEGATIF"]] for x in dampak],
           (26, 34, 18, 12, 18, 20), w_dampak)
    wsd = wb["DAMPAK KE DOKUMEN"]
    for r in range(2, wsd.max_row + 1):
        wsd.cell(r, 5).number_format = RP

    # ---------- GENTING LAIN ----------
    lain = [x for x in temuan if x["TINGKAT"] == "GENTING"
            and x["JENIS"] != "CBD/COD ANGKA SALAH"]
    lembar(wb, "GENTING LAIN",
           ["ORDER SHEET", "TAB", "JENIS", "SEL", "ISI", "KETERANGAN"],
           [[x["ORDER SHEET"], x["TAB"], x["JENIS"], x["SEL"], x["ISI"],
             x["KETERANGAN"]] for x in lain],
           (26, 34, 26, 10, 26, 70))

    # ---------- CBD/COD TERISI SEBAGIAN ----------
    sbg = [x for x in temuan if x["JENIS"] == "CBD/COD TERISI SEBAGIAN"]
    lembar(wb, "CBD-COD SEBAGIAN",
           ["ORDER SHEET", "TAB", "KOLOM", "JUDUL KOLOM", "KETERANGAN"],
           [[x["ORDER SHEET"], x["TAB"], x["SEL"], x["ISI"], x["KETERANGAN"]]
            for x in sbg],
           (26, 34, 8, 24, 64))

    out = "keluaran/AUDIT_MENYELURUH_2026.xlsx"
    wb.save(out)
    print("tersimpan:", out)
    for n in wb.sheetnames:
        print(f"  {n:<22} {wb[n].max_row - 1} baris")


if __name__ == "__main__":
    main()
